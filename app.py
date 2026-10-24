# =========================================================
# RNG PRIVATE LOGIN AND PERMANENT SUPABASE STORAGE
# =========================================================
import copy
import json
import streamlit as st
from supabase import create_client

st.set_page_config(page_title="My Dashboard", page_icon="D", layout="wide", initial_sidebar_state="collapsed")

# =========================================================
# PERSIAN FONT: B NAZANIN (LOCAL DEPLOYMENT ASSET)
# =========================================================
# Keep B-NAZANIN.TTF beside app.py on the hosting service.
# Only deploy it if your font license permits web use and embedding.
import base64 as rng_base64
from pathlib import Path as RNGPath

RNG_FONT_FILE = RNGPath(__file__).resolve().parent / "B-NAZANIN.TTF"

if RNG_FONT_FILE.is_file():
    rng_font_base64 = rng_base64.b64encode(RNG_FONT_FILE.read_bytes()).decode("ascii")
    st.markdown(
        """<style>
        @font-face {
            font-family: 'RNG B Nazanin Persian';
            src: url('data:font/ttf;base64,""" + rng_font_base64 + """') format('truetype');
            font-style: normal;
            font-weight: 400 900;
            font-display: swap;
            unicode-range: U+0600-06FF, U+0750-077F, U+08A0-08FF,
                           U+FB50-FDFF, U+FE70-FEFF;
        }
        /* Arabic/Persian script uses Nazanin; Latin keeps the existing UI fonts. */
        .stApp, .stApp input, .stApp textarea, .stApp button,
        .stApp select, .stApp label,
        .stApp [data-testid="stMarkdownContainer"] {
            font-family: 'RNG B Nazanin Persian', 'Source Sans Pro', sans-serif !important;
        }
        </style>""",
        unsafe_allow_html=True
    )
else:
    st.warning("B-NAZANIN.TTF not found beside app.py; Persian font unavailable.")


def rng_pdf_font_name():
    """Register B Nazanin for PDF Persian text; do not need Windows fonts."""
    if not RNG_FONT_FILE.is_file():
        raise FileNotFoundError(
            "B-NAZANIN.TTF not found. Upload it beside app.py (if licensed)."
        )
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    name = "RNGBNazanin"
    if name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(name, str(RNG_FONT_FILE)))
    return name


def rng_pdf_runs(visual_text):
    """B Nazanin lacks Latin letters; use Helvetica for Latin runs."""
    from reportlab.pdfbase import pdfmetrics
    persian_font = rng_pdf_font_name()
    available = pdfmetrics.getFont(persian_font).face.charWidths
    runs = []
    for ch in str(visual_text):
        code = ord(ch)
        persian = (0x0600 <= code <= 0x08FF or 0xFB50 <= code <= 0xFEFF)
        font = persian_font if code in available and (persian or not ch.isascii()) else "Helvetica"
        # Digits and punctuation can stay with B Nazanin where supported.
        if (ch.isdigit() or ch in " .,:;/-+()[]%") and code in available:
            font = persian_font
        if runs and runs[-1][0] == font:
            runs[-1] = (font, runs[-1][1] + ch)
        else:
            runs.append((font, ch))
    return runs


def rng_pdf_width(visual_text, size):
    from reportlab.pdfbase import pdfmetrics
    return sum(
        pdfmetrics.stringWidth(part, name, size)
        for name, part in rng_pdf_runs(visual_text)
    )


def rng_pdf_draw(pdf, visual_text, anchor_x, baseline, size, align="left"):
    """Draw shaped visual-order Persian/English without missing Latin letters."""
    from reportlab.pdfbase import pdfmetrics
    runs = rng_pdf_runs(visual_text)
    total_width = sum(pdfmetrics.stringWidth(s, name, size) for name, s in runs)
    x = anchor_x - (total_width / 2 if align == "center" else total_width if align == "right" else 0)
    for name, part in runs:
        pdf.setFont(name, size)
        pdf.drawString(x, baseline, part)
        x += pdfmetrics.stringWidth(part, name, size)


# Persistent record groups. UI/navigation flags remain session-local.
RNG_PERSISTENT_KEYS = (
    "sticky_notes", "tasks", "development_tasks",
    "meeting_cards", "warehouse_cards", "warehouse_card_ids",
    "warehouse_forms", "mission_report_cards",
    "mission_report_card_ids", "mission_report_forms",
    "bom_cards", "bom_card_ids", "bom_forms",
    "financial_projects",
)
RNG_DYNAMIC_PREFIXES = ("meeting_data_",)


def rng_is_persistent(key):
    return key in RNG_PERSISTENT_KEYS or any(
        key.startswith(prefix) for prefix in RNG_DYNAMIC_PREFIXES
    )


def rng_snapshot():
    return {
        key: copy.deepcopy(value)
        for key, value in st.session_state.items()
        if rng_is_persistent(key)
    }


def rng_clear_private_state():
    for key in list(st.session_state.keys()):
        if rng_is_persistent(key) or key.startswith("rng_"):
            del st.session_state[key]
    # Remove pre-existing page widget values after logging out.
    for key in list(st.session_state.keys()):
        if key not in ("page",):
            del st.session_state[key]


def rng_fail(message):
    st.error(message)
    st.stop()


try:
    settings = st.secrets["supabase"]
    rng_url = str(settings["url"]).strip()
    rng_key = str(settings["key"]).strip()
    # Owner UID is an additional access-control gate on ALL Streamlit pages.
    rng_owner_uid = str(settings["owner_uid"]).strip().lower()
    if not rng_url.startswith("https://") or not rng_key or not rng_owner_uid:
        raise ValueError("Missing or invalid Supabase configuration")
except (KeyError, ValueError, TypeError) as exc:
    rng_fail(
        "Database setup is incomplete. In Streamlit Cloud → App settings → "
        "Secrets, set [supabase] url, key and owner_uid. "
        "Use your Supabase Authentication → Users UID for owner_uid."
    )

if "rng_client" not in st.session_state:
    try:
        st.session_state.rng_client = create_client(rng_url, rng_key)
    except Exception:
        rng_fail("Unable to initialize Supabase. Check your Streamlit Secrets.")

rng_client = st.session_state.rng_client

# Verify current session with Supabase before rendering any private content.
rng_user = None
if st.session_state.get("rng_logged_in"):
    try:
        response = rng_client.auth.get_user()
        rng_user = response.user
        if rng_user is None or str(rng_user.id).lower() != rng_owner_uid:
            rng_user = None
    except Exception:
        rng_user = None

if rng_user is None:
    if st.session_state.get("rng_logged_in"):
        try:
            rng_client.auth.sign_out()
        except Exception:
            pass
        rng_clear_private_state()

    st.markdown("## My Dashboard")
    st.caption("Private workspace · Sign in to continue")
    with st.form("rng_login_form"):
        email = st.text_input("Email", autocomplete="email")
        password = st.text_input("Password", type="password")
        submit = st.form_submit_button("Sign in", use_container_width=True)
    if submit:
        try:
            signed_in = rng_client.auth.sign_in_with_password({
                "email": email.strip(), "password": password
            })
            signed_user = signed_in.user
            if signed_user is None or str(signed_user.id).lower() != rng_owner_uid:
                rng_client.auth.sign_out()
                st.error("This account is not authorized to open the Dashboard.")
            else:
                st.session_state.rng_logged_in = True
                st.rerun()
        except Exception:
            st.error("Sign in failed. Check your email and password.")
    st.stop()

# Fetch stored records once per authenticated browser session.
# A failed read must stop the app (never replace server data with blank defaults).
if not st.session_state.get("rng_data_loaded"):
    try:
        rows = (
            rng_client.table("dashboard_data")
            .select("section,data")
            .execute()
            .data
        )
        for row in rows:
            section = row.get("section", "")
            if rng_is_persistent(section):
                st.session_state[section] = copy.deepcopy(row["data"])
        st.session_state.rng_data_loaded = True
        st.session_state.rng_saved_snapshot = rng_snapshot()
    except Exception as exc:
        rng_fail("Could not load your records from Supabase. No changes were made. "
                 "Check the database connection and RLS policy. " + str(exc))


def rng_save():
    """Save only changed sections, including nested forms and task lists."""
    before = st.session_state.get("rng_saved_snapshot", {})
    after = rng_snapshot()
    changed = {
        section: value for section, value in after.items()
        if section not in before or before[section] != value
    }
    if not changed:
        return
    for section, value in changed.items():
        # One row per section prevents editing notes from overwriting BOM data.
        rng_client.table("dashboard_data").upsert(
            {"section": section, "data": value},
            on_conflict="section"
        ).execute()
        # Update the baseline only after a successful database write.
        before[section] = copy.deepcopy(value)
    st.session_state.rng_saved_snapshot = before


def rng_rerun():
    try:
        rng_save()
    except Exception as exc:
        rng_fail("The change could not be saved. Please retry. Details: " + str(exc))
    st.rerun()


# =========================================================
# NATIVE PDF SHARE BUTTON (ANDROID / SUPPORTED BROWSERS)
# =========================================================
def rng_pdf_share_button(pdf_bytes, filename, *, disabled=False):
    """Share the actual PDF file, not an unauthenticated public URL.

    The Web Share API is called directly from the user's click. The
    component keeps the PDF private in browser memory; no upload or
    public Supabase storage bucket is needed.
    """
    import base64
    import json
    import streamlit.components.v1 as components

    if disabled:
        components.html(
            '<button type="button" disabled style="width:100%;height:42px;'
            'border-radius:9px;border:1px solid #ddd;opacity:.5;">'
            'Share PDF</button>', height=46,
        )
        return

    encoded_pdf = json.dumps(base64.b64encode(pdf_bytes).decode('ascii'))
    safe_filename = json.dumps(str(filename), ensure_ascii=True)
    # JSON encoding prevents stored form titles/file names becoming executable HTML.
    html = """<!doctype html>
<html><head><meta name="viewport" content="width=device-width,initial-scale=1"></head>
<body style="margin:0;padding:0;font-family:Arial,sans-serif;">
<button id="rng-share" type="button" style="width:100%;height:42px;cursor:pointer;
 border-radius:9px;border:1px solid #d9e1ec;background:#fff;color:#172333;
 font:600 14px Arial,sans-serif;">&#x1F4E4; Share PDF</button>
<div id="rng-message" role="status" style="display:none;font-size:11px;
 color:#9a3412;padding:3px 2px;overflow-wrap:anywhere;"></div>
<script>
const pdfBase64 = __PDF_DATA__;
const pdfFilename = __PDF_FILENAME__;
const shareButton = document.getElementById('rng-share');
const msg = document.getElementById('rng-message');
function showError(message) {
  msg.style.display = 'block';
  msg.textContent = message;
  shareButton.disabled = false;
}
shareButton.addEventListener('click', async () => {
  msg.style.display = 'none';
  // Streamlit hosts this component in an iframe. Where allowed, invoking
  // the parent's navigator avoids the iframe's restricted web-share policy.
  let shareNavigator = window.navigator;
  try {
    if (window.parent && window.parent !== window &&
        typeof window.parent.navigator.share === 'function') {
      shareNavigator = window.parent.navigator;
    }
  } catch (_) { /* Cross-origin embedding: use this frame's navigator. */ }
  if (typeof shareNavigator.share !== 'function') {
    showError('Sharing is unavailable here. Open the Dashboard in Chrome on Android, or use Download PDF.');
    return;
  }
  try {
    const binary = atob(pdfBase64);
    const bytes = new Uint8Array(binary.length);
    for (let i = 0; i < binary.length; i++) bytes[i] = binary.charCodeAt(i);
    const file = new File([bytes], pdfFilename, {type:'application/pdf'});
    if (typeof shareNavigator.canShare === 'function' &&
        !shareNavigator.canShare({files:[file]})) {
      showError('This browser cannot share PDF files directly. Try Chrome on Android, or use Download PDF.');
      return;
    }
    // Call inside the click callback; file sharing requires user activation.
    await shareNavigator.share({files:[file], title:pdfFilename});
  } catch (error) {
    if (error && error.name === 'AbortError') return; // User closed share sheet.
    showError('Native sharing was blocked or unsupported. Try Chrome on Android, or use Download PDF.');
  }
});
</script></body></html>"""
    html = html.replace('__PDF_DATA__', encoded_pdf).replace('__PDF_FILENAME__', safe_filename)
    components.html(html, height=70, scrolling=False)

import pandas as pd

# =========================================================
# SESSION STATE
# =========================================================

if "page" not in st.session_state:
    st.session_state.page = "Dashboard"

if "sticky_notes" not in st.session_state:
    st.session_state.sticky_notes = []

if "editing_note" not in st.session_state:
    st.session_state.editing_note = None


# =========================================================
# CUSTOM CSS
# =========================================================

st.markdown("""
<style>

.stApp {
    background-color: #f5f7fa;
}

.block-container {
    padding-top: 2.5rem;
    padding-bottom: 3rem;
    max-width: 1400px;
}

[data-testid="stSidebar"], [data-testid="stSidebarCollapsedControl"] {
    display: none;
}

.dashboard-title {
    font-size: 38px;
    font-weight: 700;
    color: #172333;
    margin-bottom: 4px;
}

.dashboard-subtitle {
    font-size: 16px;
    color: #697586;
    margin-bottom: 25px;
}

.section-title {
    font-size: 21px;
    font-weight: 650;
    color: #172333;
    margin-top: 10px;
    margin-bottom: 16px;
}

.nav-card {
    background: white;
    border: 1px solid #e3e7ed;
    border-radius: 14px;
    padding: 22px;
    min-height: 110px;
    box-shadow: 0 2px 8px rgba(20, 30, 45, 0.05);
}

.nav-card-title {
    font-size: 17px;
    font-weight: 650;
    color: #172333;
    margin-bottom: 7px;
}

.nav-card-text {
    font-size: 13px;
    color: #7b8794;
}

.note-box {
    background: #fffdf0;
    border: 1px solid #eee5b8;
    border-radius: 12px;
    padding: 16px 18px;
    margin-bottom: 10px;
    color: #444;
}

.add-note-box {
    background: white;
    border: 1px solid #dfe4ea;
    border-radius: 14px;
    padding: 18px;
    margin-top: 15px;
    box-shadow: 0 2px 8px rgba(20, 30, 45, 0.04);
}

.stButton > button {
    border-radius: 9px;
    min-height: 42px;
}

@media (max-width: 768px) {

    .block-container {
        padding: 1.2rem 1rem 2rem 1rem;
    }

    .dashboard-title {
        font-size: 30px;
    }

    .dashboard-subtitle {
        font-size: 14px;
        margin-bottom: 20px;
    }

    .nav-card {
        padding: 16px;
    }
}

</style>
""", unsafe_allow_html=True)


# =========================================================
# PAGE NAVIGATION
# =========================================================

pages = [
    "Dashboard", "To Do List", "Meeting Minutes", "Warehouse",
    "Mission Report", "Rah Nia Development", "BOM", "Employer Report", "RNG Financial",
]


def rng_logout():
    try:
        rng_save()
        rng_client.auth.sign_out()
    except Exception as exc:
        st.error("Could not finish logout: " + str(exc))
        st.stop()
    rng_clear_private_state()
    st.rerun()


def rng_delete_warning(message):
    with st.container(key="rng_delete_confirmation"):
        st.warning(message)


def rng_task_actions(tasks_key, task_index, editing_key, prefix=""):
    """Shared task actions; deletion requires an explicit second click."""
    pending_key = f"{tasks_key}_pending_delete"
    edit_col, delete_col = st.columns(2)
    with edit_col:
        if st.button("Edit", key=f"{prefix}edit_button_{task_index}", use_container_width=True):
            st.session_state[pending_key] = None
            st.session_state[editing_key] = task_index
            rng_rerun()
    with delete_col:
        if st.button("Delete", key=f"{prefix}delete_button_{task_index}", use_container_width=True):
            st.session_state[pending_key] = task_index
            rng_rerun()
    if st.session_state.get(pending_key) == task_index:
        rng_delete_warning("Delete this task and all its subtasks? This cannot be undone.")
        confirm_col, cancel_col = st.columns(2)
        with confirm_col:
            if st.button("Yes, delete", key=f"{prefix}confirm_delete_{task_index}", type="primary", use_container_width=True):
                st.session_state[tasks_key].pop(task_index)
                st.session_state[editing_key] = None
                st.session_state[pending_key] = None
                # Task widgets use list positions. Discard editor drafts before
                # those positions are reused by the remaining tasks.
                draft_prefixes = tuple(f"{prefix}{name}_" for name in (
                    "edit_title", "edit_subtask", "edit_new_subtask"
                ))
                for key in list(st.session_state):
                    if key.startswith(draft_prefixes):
                        del st.session_state[key]
                rng_rerun()
        with cancel_col:
            if st.button("Cancel", key=f"{prefix}cancel_delete_{task_index}", use_container_width=True):
                st.session_state[pending_key] = None
                rng_rerun()


# A compact, horizontally scrollable route to every workspace.
if st.session_state.page != "Dashboard":
    st.markdown("""<style>
    .st-key-rng_delete_confirmation [data-testid="stAlertContainer"] { background: #fff4d6 !important; border: 1px solid #e4cb8f !important; border-radius: 10px; }
    .st-key-rng_delete_confirmation [data-testid="stAlertContainer"], .st-key-rng_delete_confirmation [data-testid="stMarkdownContainer"],
    .st-key-rng_delete_confirmation p, .st-key-rng_delete_confirmation strong { color: #61420b !important; }
    .st-key-rng_section_nav { background: #ffffffdf; border: 1px solid #dce5f0; border-radius: 16px; padding: 10px; margin-bottom: 12px; }
    .st-key-rng_section_nav [data-testid="stHorizontalBlock"] { flex-wrap: nowrap !important; overflow-x: auto; gap: 6px; padding-bottom: 4px; scrollbar-width: thin; scrollbar-color: #a6bbd8 transparent; }
    .st-key-rng_section_nav [data-testid="stColumn"] { flex: 0 0 auto !important; width: auto !important; min-width: max-content !important; }
    .st-key-rng_section_nav button { white-space: nowrap; min-height: 42px; padding: 7px 14px; background: transparent; border: 1px solid transparent; border-radius: 10px; color: #536b89; }
    .st-key-rng_section_nav button p { font-family: Arial, Tahoma, sans-serif !important; font-size: 13px; }
    .st-key-rng_section_nav button:hover { background: #edf3fb; color: #214b7f; border-color: #d6e2f2; }
    .st-key-rng_section_nav button[kind="primary"] { background: #203f68; color: #fff; border-color: #203f68; }
    .st-key-rng_section_nav button:focus-visible { outline: 3px solid #8aaff0; outline-offset: 2px; }
    [class*="st-key-rng_todo_card_"] [data-testid="stHorizontalBlock"],
    [class*="st-key-rng_development_card_"] [data-testid="stHorizontalBlock"] { flex-wrap: nowrap !important; gap: 8px; }
    [class*="st-key-rng_todo_card_"] [data-testid="stHorizontalBlock"] > [data-testid="stColumn"],
    [class*="st-key-rng_development_card_"] [data-testid="stHorizontalBlock"] > [data-testid="stColumn"] { min-width: 0 !important; flex: 1 1 0 !important; }
    </style>""", unsafe_allow_html=True)
    with st.container(key="rng_section_nav"):
        for column, destination in zip(st.columns(len(pages) + 1), pages + ["Log out"]):
            with column:
                key = ("rng_back_to_dashboard" if destination == "Dashboard" else
                       "rng_logout_button" if destination == "Log out" else f"rng_section_nav_{destination}")
                if st.button(destination, key=key, type="primary" if destination == st.session_state.page else "secondary"):
                    if destination == "Log out":
                        rng_logout()
                    st.session_state.page = destination
                    # A confirmation applies only while the task card is open.
                    st.session_state.tasks_pending_delete = None
                    st.session_state.development_tasks_pending_delete = None
                    rng_rerun()

page = st.session_state.page


# =========================================================
# DASHBOARD
# =========================================================




if page == "Dashboard":

    import base64
    from pathlib import Path
    from html import escape

    dashboard_folder = Path(__file__).resolve().parent
    portrait_path = dashboard_folder / "My Pic.png"
    portrait_url = (
        "data:image/png;base64," + base64.b64encode(portrait_path.read_bytes()).decode("ascii")
        if portrait_path.is_file() else ""
    )

    from datetime import datetime
    from zoneinfo import ZoneInfo

    import jdatetime

    # Both calendars use the same instant and Tehran's local day.
    today = datetime.now(ZoneInfo("Asia/Tehran"))
    today_label = today.strftime("%A, %d %B %Y")
    solar_date = jdatetime.date.fromgregorian(date=today.date())
    solar_months = (
        "فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور",
        "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند",
    )
    solar_weekdays = ("دوشنبه", "سه‌شنبه", "چهارشنبه", "پنجشنبه", "جمعه", "شنبه", "یکشنبه")
    solar_label = (
        f"{solar_weekdays[today.weekday()]}، {solar_date.day} "
        f"{solar_months[solar_date.month - 1]} {solar_date.year}"
    ).translate(str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹"))
    st.markdown('<div id="rng-dashboard-marker"></div>', unsafe_allow_html=True)
    scene_css = """
    <style>
    [data-testid="stMain"]:has(#rng-dashboard-marker) {
        position: relative; isolation: isolate; container: rng-main / inline-size;
        background:
            radial-gradient(ellipse at 24% 36%, #45688c 0, #233f61aa 30%, transparent 62%),
            radial-gradient(ellipse at 100% 0%, #24496e 0, transparent 55%),
            linear-gradient(135deg, #0c192e, #162e4c 55%, #0e2038);
        background-attachment: local;
    }
    [data-testid="stMain"]:has(#rng-dashboard-marker)::after {
        content: ""; position: absolute; inset: 0; min-height: 100vh; z-index: 0; pointer-events: none;
        background:
            radial-gradient(ellipse at 12% 50%, transparent 0 38%, #a4c9ff12 38.1% 38.2%, transparent 38.3% 49%, #a4c9ff09 49.1% 49.2%, transparent 49.3%),
            linear-gradient(120deg, transparent 0 62%, #adcfff06 62.1% 62.3%, transparent 62.4%),
            repeating-linear-gradient(90deg, #c8ddff04 0 1px, transparent 1px 80px);
        -webkit-mask-image: linear-gradient(#000 0%, #0009 70%, transparent);
        mask-image: linear-gradient(#000 0%, #0009 70%, transparent);
    }
    [data-testid="stMain"]:has(#rng-dashboard-marker)::before {
        content: ""; position: absolute; z-index: 1; pointer-events: none;
        top: 0; left: 0; width: 51%; height: min(100vh, 1080px);
        mask-composite: intersect;
        filter: contrast(1.05) drop-shadow(-12px 0 30px #a7cbff20);
        background-image: url("__PORTRAIT_URL__"); background-repeat: no-repeat;
        background-size: auto min(calc(100% - 100px), 68cqw); background-position: -40px bottom;
        -webkit-mask-image: linear-gradient(90deg, #000 0%, #000 55%, transparent 100%), linear-gradient(180deg, #000 90%, transparent 100%);
        mask-image: linear-gradient(90deg, #000 0%, #000 55%, transparent 100%), linear-gradient(180deg, #000 90%, transparent 100%);
    }
    .stApp:has(#rng-dashboard-marker) [data-testid="stHeader"] { background: #0d1b30eb; border-bottom: 1px solid #c6ddff0c; }
    .stApp:has(#rng-dashboard-marker) [data-testid="stHeader"] button,
    .stApp:has(#rng-dashboard-marker) [data-testid="stHeader"] button span,
    .stApp:has(#rng-dashboard-marker) [data-testid="stSidebarCollapsedControl"] button { color: #f1f6ff !important; }
    [data-testid="stMain"]:has(#rng-dashboard-marker) .block-container {
        position: relative; z-index: 2; max-width: none;
        padding: 2.8rem clamp(24px, 3vw, 60px) 3rem;
    }
    .stApp:has(#rng-dashboard-marker), .stApp:has(#rng-dashboard-marker) h1,
    .stApp:has(#rng-dashboard-marker) h2, .stApp:has(#rng-dashboard-marker) h3,
    .stApp:has(#rng-dashboard-marker) [data-testid="stMarkdownContainer"],
    .stApp:has(#rng-dashboard-marker) input, .stApp:has(#rng-dashboard-marker) textarea,
    .stApp:has(#rng-dashboard-marker) label {
        font-family: Arial, Tahoma, sans-serif !important; font-style: normal !important;
    }
    .rng-intro { padding: 10px 0 22px; margin-left: 40%; }
    .rng-eyebrow { display: flex; align-items: center; gap: 10px; color: #a9c8ed; font-size: 11px; font-weight: 700; letter-spacing: 2px; margin-bottom: 14px; }
    .rng-eyebrow::before { content: ""; width: 6px; height: 6px; border-radius: 50%; background: #a8cfff; box-shadow: 0 0 14px #90c3ff88; }
    .rng-intro h1.rng-hero-title { color: #f4f8ff !important; font-size: clamp(34px, 3.5vw, 52px) !important; line-height: 1.12; letter-spacing: -1.5px; font-weight: 750; margin: 0 0 14px; padding: 0; }
    .rng-intro-copy { color: #bed0e6; font-size: 15px; line-height: 1.7; }
    .rng-intro-date { display: inline-flex; align-items: center; gap: 8px; color: #d4e2f4; background: #a5c7f20c; border: 1px solid #bdd6fa26; border-radius: 100px; padding: 8px 12px; font-size: 11px; margin-top: 18px; letter-spacing: .4px; }
    .rng-date-stack { display: flex; flex-direction: column; align-items: flex-start; gap: 7px; }
    .rng-solar-date { color: #d2e1f4; font-size: 13px; font-family: Tahoma, Arial, sans-serif; font-style: normal; padding: 0 12px; }
    .st-key-rng_main_navigation { width: 100% !important; max-width: 840px; margin: 32px auto 0; padding: 0; }
    .rng-navigation-title { color: #d8e5f6; font-size: 16px; font-weight: 600; text-align: center; margin-bottom: 12px; }
    .st-key-rng_main_navigation .stButton button { color: #e6f0ff; background: #a8cfff07; border: 1px solid #aac8ef24; min-height: 132px; border-radius: 12px; display: flex; flex-direction: column; justify-content: center; gap: 12px; padding: 18px 12px; transition: background .2s, border-color .2s; }
    .st-key-rng_main_navigation .stButton button:hover { color: #fff; background: #a8cfff14; border-color: #8eb4e477; }
    .st-key-rng_main_navigation .stButton button:focus-visible { outline: 3px solid #a6c8ff; outline-offset: 2px; }
    .st-key-rng_main_navigation .stButton button:disabled { color: #cdddf3; opacity: 1; background: #a8cfff12; border-color: #aac8ef42; }
    .st-key-rng_main_navigation .stButton button::before {
        content: ""; display: block; flex-shrink: 0; width: 32px; height: 32px;
        background-image: var(--rng-navigation-icon); background-color: transparent;
        background-position: center; background-repeat: no-repeat; background-size: 28px;
        border: 0; border-radius: 0;
    }
    .st-key-rng_main_navigation .stButton button p { font-size: 14px; font-weight: 500; }
    @media (max-width: 700px) {
        .st-key-rng_main_navigation [data-testid="stHorizontalBlock"] { flex-wrap: wrap; }
        .st-key-rng_main_navigation [data-testid="stColumn"] { min-width: 100%; }
        .st-key-rng_main_navigation .stButton button { min-height: 120px; gap: 10px; }
    }
    .st-key-rng_navigation_card_0 { --rng-navigation-icon: url("data:image/svg+xml,%3Csvg%20xmlns%3D%22http%3A%2F%2Fwww.w3.org%2F2000%2Fsvg%22%20viewBox%3D%220%200%2024%2024%22%20fill%3D%22none%22%20stroke%3D%22%23bdd8ff%22%20stroke-width%3D%221.5%22%20stroke-linecap%3D%22round%22%20stroke-linejoin%3D%22round%22%3E%3Crect%20x%3D%223%22%20y%3D%223%22%20width%3D%227%22%20height%3D%227%22%20rx%3D%221.5%22%2F%3E%3Crect%20x%3D%2214%22%20y%3D%223%22%20width%3D%227%22%20height%3D%227%22%20rx%3D%221.5%22%2F%3E%3Crect%20x%3D%223%22%20y%3D%2214%22%20width%3D%227%22%20height%3D%227%22%20rx%3D%221.5%22%2F%3E%3Crect%20x%3D%2214%22%20y%3D%2214%22%20width%3D%227%22%20height%3D%227%22%20rx%3D%221.5%22%2F%3E%3C%2Fsvg%3E"); }
    .st-key-rng_navigation_card_1 { --rng-navigation-icon: url("data:image/svg+xml,%3Csvg%20xmlns%3D%22http%3A%2F%2Fwww.w3.org%2F2000%2Fsvg%22%20viewBox%3D%220%200%2024%2024%22%20fill%3D%22none%22%20stroke%3D%22%23bdd8ff%22%20stroke-width%3D%221.5%22%20stroke-linecap%3D%22round%22%20stroke-linejoin%3D%22round%22%3E%3Cpath%20d%3D%22m3%206%202%202%204-4m3%202h9M3%2013l2%202%204-4m3%202h9M3%2020h18%22%2F%3E%3C%2Fsvg%3E"); }
    .st-key-rng_navigation_card_2 { --rng-navigation-icon: url("data:image/svg+xml,%3Csvg%20xmlns%3D%22http%3A%2F%2Fwww.w3.org%2F2000%2Fsvg%22%20viewBox%3D%220%200%2024%2024%22%20fill%3D%22none%22%20stroke%3D%22%23bdd8ff%22%20stroke-width%3D%221.5%22%20stroke-linecap%3D%22round%22%20stroke-linejoin%3D%22round%22%3E%3Ccircle%20cx%3D%229%22%20cy%3D%228%22%20r%3D%223%22%2F%3E%3Cpath%20d%3D%22M3%2021v-3a6%206%200%200%201%2012%200v3M16%205a3%203%200%200%201%200%206m2%204a5%205%200%200%201%203%205v1%22%2F%3E%3C%2Fsvg%3E"); }
    .st-key-rng_navigation_card_3 { --rng-navigation-icon: url("data:image/svg+xml,%3Csvg%20xmlns%3D%22http%3A%2F%2Fwww.w3.org%2F2000%2Fsvg%22%20viewBox%3D%220%200%2024%2024%22%20fill%3D%22none%22%20stroke%3D%22%23bdd8ff%22%20stroke-width%3D%221.5%22%20stroke-linecap%3D%22round%22%20stroke-linejoin%3D%22round%22%3E%3Cpath%20d%3D%22m3%207%209-4%209%204v11l-9%204-9-4Zm0%200%209%204%209-4m-9%204v11M7%205l10%204v5%22%2F%3E%3C%2Fsvg%3E"); }
    .st-key-rng_navigation_card_4 { --rng-navigation-icon: url("data:image/svg+xml,%3Csvg%20xmlns%3D%22http%3A%2F%2Fwww.w3.org%2F2000%2Fsvg%22%20viewBox%3D%220%200%2024%2024%22%20fill%3D%22none%22%20stroke%3D%22%23bdd8ff%22%20stroke-width%3D%221.5%22%20stroke-linecap%3D%22round%22%20stroke-linejoin%3D%22round%22%3E%3Cpath%20d%3D%22M14%203H5a2%202%200%200%200-2%202v14a2%202%200%200%200%202%202h14a2%202%200%200%200%202-2V8Zm0%200v5h7M7%2013h10m-10%204h6%22%2F%3E%3C%2Fsvg%3E"); }
    .st-key-rng_navigation_card_5 { --rng-navigation-icon: url("data:image/svg+xml,%3Csvg%20xmlns%3D%22http%3A%2F%2Fwww.w3.org%2F2000%2Fsvg%22%20viewBox%3D%220%200%2024%2024%22%20fill%3D%22none%22%20stroke%3D%22%23bdd8ff%22%20stroke-width%3D%221.5%22%20stroke-linecap%3D%22round%22%20stroke-linejoin%3D%22round%22%3E%3Cpath%20d%3D%22M4%2020V10m7%2010V4m7%2016v-7M2%2020h20M15%204h6v6m-1-5-6%206%22%2F%3E%3C%2Fsvg%3E"); }
    .st-key-rng_navigation_card_6 { --rng-navigation-icon: url("data:image/svg+xml,%3Csvg%20xmlns%3D%22http%3A%2F%2Fwww.w3.org%2F2000%2Fsvg%22%20viewBox%3D%220%200%2024%2024%22%20fill%3D%22none%22%20stroke%3D%22%23bdd8ff%22%20stroke-width%3D%221.5%22%20stroke-linecap%3D%22round%22%20stroke-linejoin%3D%22round%22%3E%3Cpath%20d%3D%22m12%203%2010%205-10%205L2%208Zm-10%2010%2010%205%2010-5M2%2018l10%205%2010-5%22%2F%3E%3C%2Fsvg%3E"); }
    .st-key-rng_navigation_card_7 { --rng-navigation-icon: url("data:image/svg+xml,%3Csvg%20xmlns%3D%22http%3A%2F%2Fwww.w3.org%2F2000%2Fsvg%22%20viewBox%3D%220%200%2024%2024%22%20fill%3D%22none%22%20stroke%3D%22%23bdd8ff%22%20stroke-width%3D%221.5%22%20stroke-linecap%3D%22round%22%20stroke-linejoin%3D%22round%22%3E%3Crect%20x%3D%223%22%20y%3D%227%22%20width%3D%2218%22%20height%3D%2214%22%20rx%3D%222%22%2F%3E%3Cpath%20d%3D%22M8%207V4h8v3M3%2012l9%204%209-4m-9%202v4%22%2F%3E%3C%2Fsvg%3E"); }
    .st-key-rng_navigation_card_8 { --rng-navigation-icon: url("data:image/svg+xml,%3Csvg%20xmlns%3D%22http%3A%2F%2Fwww.w3.org%2F2000%2Fsvg%22%20viewBox%3D%220%200%2024%2024%22%20fill%3D%22none%22%20stroke%3D%22%23bdd8ff%22%20stroke-width%3D%221.5%22%20stroke-linecap%3D%22round%22%20stroke-linejoin%3D%22round%22%3E%3Cpath%20d%3D%22M3%2021h18M5%2018V10m5%208V10m4%208V10m5%208V10M3%207l9-4%209%204H3Z%22%2F%3E%3C%2Fsvg%3E"); }
    .st-key-rng_navigation_card_9 { --rng-navigation-icon: url("data:image/svg+xml,%3Csvg%20xmlns%3D%22http%3A%2F%2Fwww.w3.org%2F2000%2Fsvg%22%20viewBox%3D%220%200%2024%2024%22%20fill%3D%22none%22%20stroke%3D%22%23edc2cb%22%20stroke-width%3D%221.5%22%20stroke-linecap%3D%22round%22%20stroke-linejoin%3D%22round%22%3E%3Cpath%20d%3D%22M9%205H5a2%202%200%200%200-2%202v12a2%202%200%200%200%202%202h4M15%205l6%206-6%206m-6-6h12%22%2F%3E%3C%2Fsvg%3E"); }
    .rng-intro-date svg { width: 13px; height: 13px; color: #a9c9ef; }
    .st-key-rng_glass_notes {
        position: relative; overflow: hidden; background: #f5f9fff0; border: 1px solid #ffffffd9; border-top: 2px solid #ffffff; border-radius: 22px;
        width: 60% !important; margin-left: 40%; padding: 28px; margin-top: 10px; box-shadow: 0 24px 70px #030e2338, inset 0 1px 0 #ffffff;
        -webkit-backdrop-filter: blur(18px); backdrop-filter: blur(18px);
    }
    .rng-notes-heading { display: flex; align-items: center; justify-content: space-between; gap: 12px; }
    .rng-notes-title { display: flex; align-items: center; gap: 12px; }
    .rng-notes-icon { width: 36px; height: 36px; border-radius: 11px; color: #365a88; background: #e0eaf8; border: 1px solid #d1dff1; display: flex; align-items: center; justify-content: center; }
    .rng-notes-icon svg { width: 18px; height: 18px; }
    .rng-notes-heading .section-title { margin: 0; color: #1c3455; font-size: 22px; letter-spacing: -.6px; }
    .rng-note-count { color: #365d8a; background: #dfeaf8; border: 1px solid #ccdcef; border-radius: 100px; padding: 5px 12px; font-size: 11px; white-space: nowrap; }
    .rng-section-caption { color: #526d8e; font-size: 13px; margin-top: 5px; }
    .rng-empty-notes { border: 1px dashed #b9cbe2; border-radius: 12px; padding: 26px 22px; color: #526d8e; line-height: 1.9; background: repeating-linear-gradient(#f9fbff, #f9fbff 31px, #e9eff9 32px); margin-top: 20px; }
    .rng-empty-notes strong { display: block; color: #2d4c74; font-size: 20px; font-weight: 600; }
    .st-key-rng_glass_notes .note-box { background: #fffdf4; color: #30435b; border: 1px solid #eae3ca; border-left: 3px solid #a6bce0; border-radius: 10px; padding: 14px 16px; white-space: pre-wrap; overflow-wrap: anywhere; min-height: 48px; }
    .st-key-rng_glass_notes .stButton button { color: #ffffff; background: #203c60; border-color: #33537b; min-height: 44px; }
    .st-key-rng_glass_notes .stButton button:hover { color: #ffffff; background: #2b4a73; border-color: #6a90c0; }
    .st-key-rng_add_note .stButton button { background: linear-gradient(110deg, #1a3558, #2a5180); color: #fff; border-color: #254b79; min-height: 48px; box-shadow: 0 5px 12px #16345b20; }
    .st-key-rng_glass_notes input, .st-key-rng_glass_notes textarea { color: #243f63; background: #f8fbff; }
    .st-key-rng_glass_notes button:focus-visible { outline: 3px solid #7faeff; outline-offset: 3px; }
    .rng-dashboard-signature { display: flex; justify-content: space-between; gap: 12px; color: #a5bdd9; font-size: 11px; padding-top: 22px; }
    @container rng-main (max-width: 900px) {
        .block-container:has(#rng-dashboard-marker) { padding: 2rem 24px 3rem !important; }
        .rng-intro { min-height: 340px; margin-left: 43%; }
        .st-key-rng_glass_notes { width: 100% !important; margin-left: 0; }
        .rng-intro h1.rng-hero-title { font-size: 34px !important; }
    }
    @media (max-width: 1200px) and (min-width: 701px) {
        [data-testid="stMain"]:has(#rng-dashboard-marker)::before { width: 62%; height: 450px; background-size: auto 350px; background-position: -20px bottom; }
    }
    @media (max-width: 700px) {
        [data-testid="stMain"]:has(#rng-dashboard-marker)::before {
            width: 100%; height: 680px; background-size: auto 360px; background-position: -15px bottom;
            -webkit-mask-image: linear-gradient(90deg, #000 0%, #000 48%, transparent 100%), linear-gradient(180deg, #000 90%, transparent 100%);
            mask-image: linear-gradient(90deg, #000 0%, #000 48%, transparent 100%), linear-gradient(180deg, #000 90%, transparent 100%);
        }
        .block-container:has(#rng-dashboard-marker) { padding: 1.3rem 1rem 3rem !important; }
        .rng-intro { min-height: 610px; margin-left: 0; }
        .rng-intro h1.rng-hero-title { font-size: 36px !important; }
        .rng-intro-copy { font-size: 14px; }
        .rng-intro-date { margin-top: 12px; }
        .rng-notes-title { gap: 8px; }
        .rng-notes-heading .section-title { font-size: 20px; }
        .st-key-rng_glass_notes { padding: 18px; }
        .st-key-rng_main_navigation { max-width: 320px; padding: 0; }
        .rng-dashboard-signature { flex-wrap: wrap; }
    }
    @media (prefers-reduced-transparency: reduce) {
        .st-key-rng_glass_notes { background: #f5f9ff; backdrop-filter: none; }
    }
    </style>
    """
    st.markdown(scene_css.replace("__PORTRAIT_URL__", portrait_url), unsafe_allow_html=True)
    st.markdown(
        '<div class="rng-intro"><div class="rng-eyebrow">PERSONAL WORKSPACE</div>'
        '<h1 class="rng-hero-title">My Dashboard</h1>'
        '<div class="rng-intro-copy">A clear space for your plans, projects, and ideas.</div>'
        '<div class="rng-date-stack"><div class="rng-intro-date"><svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.7" aria-hidden="true"><rect x="3" y="5" width="18" height="16" rx="3"/><path d="M7 3v4m10-4v4M3 11h18"/></svg>'
        f'{escape(today_label)}</div>'
        f'<div class="rng-solar-date" lang="fa" dir="rtl">{escape(solar_label)} · هجری شمسی</div></div></div>',
        unsafe_allow_html=True
    )
    if not portrait_url:
        st.warning("My Pic.png not found beside app.py.")

    # =====================================================
    # GLASSY STICKY NOTES PANEL
    # =====================================================

    with st.container(key="rng_glass_notes"):

        note_count = len(st.session_state.sticky_notes)
        st.markdown(
            '<div class="rng-notes-heading"><div class="rng-notes-title"><span class="rng-notes-icon">'
            '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.6" stroke-linecap="round" aria-hidden="true"><path d="M14 3H5a2 2 0 0 0-2 2v14a2 2 0 0 0 2 2h14a2 2 0 0 0 2-2v-9M8 16l1-4L19 2l3 3-10 10-4 1Z"/></svg>'
            '</span><div class="section-title">Sticky Notes</div></div>'
            f'<span class="rng-note-count">{note_count} {"note" if note_count == 1 else "notes"}</span></div>'
            '<div class="rng-section-caption">The little things worth keeping.</div>',
            unsafe_allow_html=True
        )
        if not note_count:
            st.markdown(
                '<div class="rng-empty-notes"><strong>Your next idea starts here.</strong>A reminder, a spark, a thought for later. Make yourself a little space.</div>',
                unsafe_allow_html=True
            )

        # =================================================
        # EXISTING STICKY NOTES
        # =================================================

        for i, note in enumerate(
            st.session_state.sticky_notes
        ):

            if st.session_state.editing_note == i:

                edited_note = st.text_area(
                    "Edit note",
                    value=note,
                    key=f"edit_text_{i}",
                    height=100
                )

                edit_col1, edit_col2 = st.columns(2)

                with edit_col1:

                    if st.button(
                        "Save",
                        key=f"save_{i}",
                        use_container_width=True
                    ):

                        if edited_note.strip():

                            st.session_state.sticky_notes[i] = (
                                edited_note.strip()
                            )

                        st.session_state.editing_note = None
                        rng_rerun()

                with edit_col2:

                    if st.button(
                        "Cancel",
                        key=f"cancel_{i}",
                        use_container_width=True
                    ):

                        st.session_state.editing_note = None
                        rng_rerun()

            else:

                note_col, edit_col, delete_col = st.columns(
                    [7, 1.5, 1.5],
                    gap="small"
                )

                with note_col:

                    st.markdown(
                        f"""
                        <div class="note-box">
                            {escape(str(note))}
                        </div>
                        """,
                        unsafe_allow_html=True
                    )

                with edit_col:

                    if st.button(
                        "Edit",
                        key=f"edit_{i}",
                        use_container_width=True
                    ):

                        st.session_state.editing_note = i
                        rng_rerun()

                with delete_col:

                    if st.button(
                        "Delete",
                        key=f"delete_{i}",
                        use_container_width=True
                    ):

                        st.session_state.sticky_notes.pop(i)
                        rng_rerun()

        # =================================================
        # ADD NEW STICKY NOTE
        # =================================================

        st.markdown(
            '<div style="height: 12px;"></div>',
            unsafe_allow_html=True
        )

        new_note = st.text_input(
            "New sticky note",
            placeholder="Type a new note here...",
            label_visibility="collapsed",
            key="rng_new_sticky_note"
        )

        with st.container(key="rng_add_note"):

            if st.button(
                "+  Add Sticky Note",
                key="rng_add_sticky_note_button",
                use_container_width=True
            ):

                if new_note.strip():

                    st.session_state.sticky_notes.append(
                        new_note.strip()
                    )

                    rng_rerun()

    with st.container(key="rng_main_navigation"):
        st.markdown(
            '<div class="rng-navigation-title">Your workspace</div>',
            unsafe_allow_html=True
        )
        navigation_items = pages + ["Log out"]
        for offset in range(0, len(navigation_items), 3):
            row_items = navigation_items[offset:offset + 3]
            row_columns = st.columns(3, gap="medium")
            if len(row_items) == 1:
                row_columns = [row_columns[1]]
            for index, (column, destination) in enumerate(
                zip(row_columns, row_items), start=offset
            ):
                with column:
                    with st.container(key=f"rng_navigation_card_{index}"):
                        button_key = "rng_logout_button" if destination == "Log out" else f"rng_main_nav_{destination}"
                        if st.button(
                            destination, key=button_key, disabled=destination == page,
                            use_container_width=True
                        ):
                            if destination == "Log out":
                                rng_logout()
                            else:
                                st.session_state.page = destination
                                rng_rerun()

    st.markdown(
        '<div class="rng-dashboard-signature"><span>A little clarity goes a long way.</span><span>YOUR SPACE. YOUR PACE.</span></div>',
        unsafe_allow_html=True
    )

    # =====================================================
    # END DASHBOARD
    # =====================================================

# =========================================================
# WORKSPACE PAGES
# =========================================================

elif page == "To Do List":
    from datetime import datetime
    from html import escape

    # =========================================================
    # TASK DATA
    # =========================================================

    if "tasks" not in st.session_state:
        st.session_state.tasks = []

    if "creating_task" not in st.session_state:
        st.session_state.creating_task = False

    if "editing_task" not in st.session_state:
        st.session_state.editing_task = None

    if "new_subtask_count" not in st.session_state:
        st.session_state.new_subtask_count = 1

    # =========================================================
    # CSS
    # =========================================================

    st.markdown('<div id="rng-todo-marker"></div>', unsafe_allow_html=True)
    st.markdown("""
    <style>
    [data-testid="stMain"]:has(#rng-todo-marker) { background: radial-gradient(ellipse at 100% 0%, #dbe8fb, transparent 50%), #f2f5fa; }
    .stApp:has(#rng-todo-marker), .stApp:has(#rng-todo-marker) h1,
    .stApp:has(#rng-todo-marker) h2, .stApp:has(#rng-todo-marker) h3,
    .stApp:has(#rng-todo-marker) [data-testid="stMarkdownContainer"],
    .stApp:has(#rng-todo-marker) input, .stApp:has(#rng-todo-marker) textarea,
    .stApp:has(#rng-todo-marker) label { font-family: Arial, Tahoma, sans-serif !important; font-style: normal !important; }
    .stApp:has(#rng-todo-marker) .block-container { max-width: 1440px; padding-top: 2.5rem; }
    .rng-todo-heading { background: radial-gradient(ellipse at right top, #305b89, transparent 65%), #142c4b; border: 1px solid #29496e; border-radius: 24px; padding: 32px 34px; margin: 4px 0 8px; position: relative; overflow: hidden; }
    .rng-todo-eyebrow { color: #acc8ec; font-size: 11px; font-weight: 700; letter-spacing: 2px; margin-bottom: 12px; }
    .rng-todo-heading h1 { color: #f5f9ff !important; font-size: clamp(32px, 3.5vw, 44px) !important; font-weight: 750; letter-spacing: -1.3px; margin: 0 0 12px; padding: 0; }
    .rng-todo-heading p { color: #c5d7ed; font-size: 15px; margin: 0; line-height: 1.7; }
    .rng-todo-summary { display: grid; grid-template-columns: repeat(3, 1fr); gap: 14px; margin: 8px 0 4px; }
    .rng-todo-stat { background: #ffffff; border: 1px solid #dce5f1; border-radius: 16px; padding: 19px 22px; }
    .rng-todo-stat strong { display: block; color: #203d64; font-size: 27px; line-height: 1.2; font-weight: 700; margin-bottom: 6px; }
    .rng-todo-stat span { color: #607897; font-size: 12px; }
    .rng-todo-stat:last-child { background: #eaf1fc; border-color: #d4e1f4; }
    .st-key-rng_todo_create_action { margin: 8px 0 12px; }
    .st-key-rng_todo_create_action .stButton button { background: #214b7f; color: white; border: 1px solid #214b7f; min-height: 48px; border-radius: 12px; }
    .st-key-rng_todo_create_action .stButton button:hover { color: #fff; background: #173b68; }
    .rng-todo-section { display: flex; align-items: center; justify-content: space-between; color: #253f62; font-size: 19px; font-weight: 700; margin: 12px 0 4px; }
    .rng-todo-section span { font-size: 12px; font-weight: 400; color: #69809d; }
    .st-key-rng_todo_create_form { background: #fff; border: 1px solid #d5e2f3; border-radius: 18px; padding: 24px; margin-bottom: 16px; box-shadow: 0 8px 30px #173b6810; }
    .st-key-rng_todo_create_form h3 { font-size: 21px; color: #203d64; }
    [class*="st-key-rng_todo_card_"] { background: #fff; border: 1px solid #dde5f0; border-top: 3px solid #7195c6; border-radius: 18px; padding: 22px; box-shadow: 0 6px 20px #16335406; min-height: 340px; }
    [class*="st-key-rng_todo_card_"]:has(.rng-task-done) { background: #f5f8ff; border-top-color: #a1b6d8; }
    .rng-task-status { display: inline-block; color: #3f6190; background: #edf3fc; border: 1px solid #dce7f7; border-radius: 100px; padding: 5px 10px; font-size: 10px; font-weight: 700; letter-spacing: .6px; text-transform: uppercase; margin-bottom: 12px; }
    .rng-task-done { background: #e4edf9; color: #4d6484; }
    .todo-title { color: #243b5b !important; font-size: 19px; font-weight: 700; line-height: 1.5; margin: 0 0 12px; overflow-wrap: anywhere; }
    .todo-section-title { color: #7a8da8; font-size: 10px; font-weight: 700; letter-spacing: 1.5px; text-transform: uppercase; margin: 14px 0 4px; }
    .todo-divider { height: 1px; background: #e8edf5; margin: 8px 0; }
    .todo-date, .todo-completed { color: #7588a3; font-size: 11px; line-height: 1.7; }
    .todo-completed { color: #496d9c; }
    .rng-subtask-progress { height: 4px; background: #edf2f9; border-radius: 10px; overflow: hidden; margin: 9px 0 6px; }
    .rng-subtask-progress span { display: block; height: 100%; background: #769bcf; border-radius: inherit; }
    .rng-subtask-count { color: #768ba6; font-size: 11px; }
    .st-key-rng_todo_board [data-testid="stCheckbox"] { padding: 0; margin: 2px 0; }
    .st-key-rng_todo_board [data-testid="stCheckbox"] label { min-height: 36px; }
    .st-key-rng_todo_board [data-testid="stCheckbox"] label:has(input:checked) > div:not([data-testid]) { background-color: #527db4; border-color: #527db4; }
    .st-key-rng_todo_board [data-testid="stCheckbox"] label p { color: #435a76; font-size: 13px; overflow-wrap: anywhere; }
    .st-key-rng_todo_board .stButton button, .st-key-rng_todo_create_form .stButton button { color: #355780; background: #f3f6fc; border: 1px solid #dde6f4; border-radius: 9px; min-height: 44px; }
    .st-key-rng_todo_board .stButton button:hover, .st-key-rng_todo_create_form .stButton button:hover { border-color: #8aabd8; color: #224773; background: #eaf1fc; }
    .stApp:has(#rng-todo-marker) button:focus-visible { outline: 3px solid #8aaff0; outline-offset: 2px; }
    .rng-todo-empty { text-align: center; background: #ffffffb8; border: 1px dashed #c3d2e7; border-radius: 20px; padding: 52px 20px; margin-top: 10px; }
    .rng-todo-empty svg { width: 48px; height: 48px; color: #87a5cd; margin-bottom: 18px; }
    .rng-todo-empty strong { display: block; color: #304d73; font-size: 21px; margin-bottom: 10px; }
    .rng-todo-empty p { color: #7185a0; font-size: 14px; line-height: 1.8; margin: 0; }
    @media (max-width: 1100px) and (min-width: 701px) {
        .st-key-rng_todo_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_todo_card_"]) { flex-wrap: wrap; }
        .st-key-rng_todo_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_todo_card_"]) > [data-testid="stColumn"] { min-width: calc(50% - 1rem); }
    }
    @media (max-width: 700px) {
        .rng-todo-heading { padding: 26px 22px; border-radius: 20px; }
        .rng-todo-summary { gap: 8px; }
        .rng-todo-stat { padding: 15px 10px; }
        .rng-todo-stat strong { font-size: 24px; }
        .rng-todo-stat span { font-size: 11px; }
        .rng-todo-section span { display: none; }
        .st-key-rng_todo_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_todo_card_"]) { flex-wrap: wrap; }
        .st-key-rng_todo_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_todo_card_"]) > [data-testid="stColumn"] { min-width: 100%; }
        [class*="st-key-rng_todo_card_"] { padding: 20px; }
        .st-key-rng_todo_create_form { padding: 18px; }
    }
    </style>
    """, unsafe_allow_html=True)

    task_total = len(st.session_state.tasks)
    done_total = sum(bool(task["done"]) for task in st.session_state.tasks)
    st.markdown(
        '<div class="rng-todo-heading"><div class="rng-todo-eyebrow">A LITTLE FOCUS GOES A LONG WAY</div>'
        '<h1>To Do List</h1><p>Make space for what matters. One task at a time.</p></div>'
        '<div class="rng-todo-summary">'
        f'<div class="rng-todo-stat"><strong>{task_total}</strong><span>Total tasks</span></div>'
        f'<div class="rng-todo-stat"><strong>{task_total - done_total}</strong><span>In progress</span></div>'
        f'<div class="rng-todo-stat"><strong>{done_total}</strong><span>Completed</span></div></div>',
        unsafe_allow_html=True
    )
    with st.container(key="rng_todo_create_action"):
        if st.button("Create Task", key="create_task_button", use_container_width=True):
            st.session_state.creating_task = True
            st.session_state.new_subtask_count = 1
            rng_rerun()

    # =========================================================
    # CREATE TASK FORM
    # =========================================================

    if st.session_state.creating_task:
        with st.container(key="rng_todo_create_form"):

            st.markdown("### Create New Task")

            new_title = st.text_input(
                "Task Title",
                placeholder="Enter task title...",
                key="new_task_title"
            )

            st.markdown("Sub Tasks")

            subtask_values = []

            for i in range(st.session_state.new_subtask_count):

                value = st.text_input(
                    f"Sub Task {i + 1}",
                    placeholder="Enter sub task...",
                    key=f"create_subtask_{i}"
                )

                subtask_values.append(value)

            add_col, create_col, cancel_col = st.columns(3)

            with add_col:
                if st.button(
                    "Add Sub Task",
                    key="add_subtask_button",
                    use_container_width=True
                ):
                    st.session_state.new_subtask_count += 1
                    rng_rerun()

            with create_col:
                if st.button(
                    "Save Task",
                    key="save_new_task",
                    use_container_width=True
                ):

                    if new_title.strip():

                        subtasks = []

                        for value in subtask_values:

                            if value.strip():
                                subtasks.append({
                                    "text": value.strip(),
                                    "done": False
                                })

                        st.session_state.tasks.append({
                            "title": new_title.strip(),
                            "subtasks": subtasks,
                            "created": datetime.now().strftime("%Y-%m-%d"),
                            "completed": None,
                            "done": False,
                        })

                        st.session_state.creating_task = False
                        st.session_state.new_subtask_count = 1

                        rng_rerun()

            with cancel_col:
                if st.button(
                    "Cancel",
                    key="cancel_new_task",
                    use_container_width=True
                ):
                    st.session_state.creating_task = False
                    st.session_state.new_subtask_count = 1
                    rng_rerun()

            st.divider()

    # =========================================================
    # SORT TASKS
    # =========================================================

    unfinished = [
        task for task in st.session_state.tasks
        if not task["done"]
    ]

    completed = [
        task for task in st.session_state.tasks
        if task["done"]
    ]

    sorted_tasks = unfinished + completed

    def todo_persist_change():
        try:
            rng_save()
        except Exception as exc:
            st.error("The change could not be saved. Please retry. Details: " + str(exc))
            st.stop()

    def todo_toggle_main(task_index):
        task = st.session_state.tasks[task_index]
        task["done"] = bool(st.session_state[f"main_checkbox_{task_index}"])
        task["completed"] = datetime.now().strftime("%Y-%m-%d") if task["done"] else None
        todo_persist_change()

    def todo_toggle_subtask(task_index, sub_index):
        task = st.session_state.tasks[task_index]
        task["subtasks"][sub_index]["done"] = bool(st.session_state[f"sub_checkbox_{task_index}_{sub_index}"])
        task["done"] = bool(task["subtasks"]) and all(item["done"] for item in task["subtasks"])
        task["completed"] = datetime.now().strftime("%Y-%m-%d") if task["done"] else None
        # Callbacks run before rendering, so the main checkbox shares the new state.
        st.session_state[f"main_checkbox_{task_index}"] = task["done"]
        todo_persist_change()

    # =========================================================
    # TASK CARDS
    # =========================================================

    st.markdown(
        '<div class="rng-todo-section">Your tasks<span>Unfinished first · Progress always visible</span></div>',
        unsafe_allow_html=True
    )
    if not sorted_tasks:
        st.markdown(
            '<div class="rng-todo-empty"><svg viewBox="0 0 48 48" fill="none" stroke="currentColor" stroke-width="1.5" aria-hidden="true"><rect x="10" y="9" width="28" height="33" rx="5"/><rect x="18" y="5" width="12" height="8" rx="3" fill="#f2f5fa"/><path d="m17 26 5 5 10-12"/></svg>'
            '<strong>A clear list. A fresh start.</strong><p>Create your first task and break it into small, achievable steps.</p></div>',
            unsafe_allow_html=True
        )

    with st.container(key="rng_todo_board"):
        for row_start in range(0, len(sorted_tasks), 3):

            row_tasks = sorted_tasks[row_start:row_start + 3]

            columns = st.columns(3)

            for column_number, task in enumerate(row_tasks):

                task_index = st.session_state.tasks.index(task)
                with columns[column_number], st.container(key=f"rng_todo_card_{task_index}"):

                    # =================================================
                    # EDIT MODE
                    # =================================================

                    if st.session_state.editing_task == task_index:

                        st.markdown(
                            '<div class="todo-title">Edit Task</div>',
                            unsafe_allow_html=True
                        )

                        st.markdown(
                            '<div class="todo-divider"></div>',
                            unsafe_allow_html=True
                        )

                        edited_title = st.text_input(
                            "Task Title",
                            value=task["title"],
                            key=f"edit_title_{task_index}"
                        )

                        st.markdown("Sub Tasks")

                        edited_subtasks = []

                        for sub_index, subtask in enumerate(task["subtasks"]):

                            edited_text = st.text_input(
                                f"Sub Task {sub_index + 1}",
                                value=subtask["text"],
                                key=f"edit_subtask_{task_index}_{sub_index}"
                            )

                            edited_subtasks.append({
                                "text": edited_text,
                                "done": subtask["done"]
                            })

                        new_edit_subtask = st.text_input(
                            "Add another sub task",
                            placeholder="Optional",
                            key=f"edit_new_subtask_{task_index}"
                        )

                        save_col, cancel_col = st.columns(2)

                        with save_col:

                            if st.button(
                                "Save Changes",
                                key=f"save_edit_{task_index}",
                                use_container_width=True
                            ):

                                if edited_title.strip():
                                    task["title"] = edited_title.strip()

                                task["subtasks"] = [
                                    subtask
                                    for subtask in edited_subtasks
                                    if subtask["text"].strip()
                                ]

                                if new_edit_subtask.strip():
                                    task["subtasks"].append({
                                        "text": new_edit_subtask.strip(),
                                        "done": False
                                    })

                                if (
                                    len(task["subtasks"]) > 0
                                    and all(
                                        subtask["done"]
                                        for subtask in task["subtasks"]
                                    )
                                ):

                                    task["done"] = True

                                    if not task["completed"]:
                                        task["completed"] = (
                                            datetime.now().strftime("%Y-%m-%d")
                                        )

                                else:

                                    task["done"] = False
                                    task["completed"] = None

                                st.session_state.editing_task = None
                                rng_rerun()

                        with cancel_col:

                            if st.button(
                                "Cancel",
                                key=f"cancel_edit_{task_index}",
                                use_container_width=True
                            ):
                                st.session_state.editing_task = None
                                rng_rerun()

                        continue

                    # =================================================
                    # CARD TITLE
                    # =================================================

                    finished_steps = sum(bool(item["done"]) for item in task["subtasks"])
                    step_total = len(task["subtasks"])
                    status = "Completed" if task["done"] else ("In progress" if finished_steps else "To do")
                    done_class = " rng-task-done" if task["done"] else ""
                    st.markdown(
                        f'<div class="rng-task-status{done_class}">{status}</div>'
                        f'<div class="todo-title">{escape(str(task["title"]))}</div>',
                        unsafe_allow_html=True
                    )
                    if step_total:
                        st.markdown(
                            f'<div class="rng-subtask-progress"><span style="width:{100 * finished_steps / step_total:.1f}%"></span></div>'
                            f'<div class="rng-subtask-count">{finished_steps} of {step_total} subtasks complete</div>',
                            unsafe_allow_html=True
                        )

                    # =================================================
                    # MAIN TASK CHECKBOX
                    # =================================================

                    st.session_state[f"main_checkbox_{task_index}"] = task["done"]
                    st.checkbox(
                        "Task completed", key=f"main_checkbox_{task_index}",
                        on_change=todo_toggle_main, args=(task_index,)
                    )

                    # =================================================
                    # SUBTASKS
                    # =================================================

                    st.markdown(
                        '<div class="todo-divider"></div>',
                        unsafe_allow_html=True
                    )

                    st.markdown(
                        '<div class="todo-section-title">Sub Tasks</div>',
                        unsafe_allow_html=True
                    )

                    for sub_index, subtask in enumerate(task["subtasks"]):

                        st.session_state[f"sub_checkbox_{task_index}_{sub_index}"] = subtask["done"]
                        st.checkbox(
                            subtask["text"], key=f"sub_checkbox_{task_index}_{sub_index}",
                            on_change=todo_toggle_subtask, args=(task_index, sub_index)
                        )

                    # =================================================
                    # DATES
                    # =================================================

                    st.markdown(
                        '<div class="todo-divider"></div>',
                        unsafe_allow_html=True
                    )

                    st.markdown(
                        f"""
                        <div class="todo-date">
                            Created: {task["created"]}
                        </div>
                        """,
                        unsafe_allow_html=True
                    )

                    if task["completed"]:

                        st.markdown(
                            f"""
                            <div class="todo-completed">
                                Completed: {task["completed"]}
                            </div>
                            """,
                            unsafe_allow_html=True
                        )

                    else:

                        st.markdown(
                            """
                            <div class="todo-date">
                                Completed: -
                            </div>
                            """,
                            unsafe_allow_html=True
                        )

                    # =================================================
                    # EDIT BUTTON
                    # =================================================

                    rng_task_actions("tasks", task_index, "editing_task", "")



elif page == "Meeting Minutes":

    from html import escape

    st.markdown('<div id="rng-meeting-marker"></div>', unsafe_allow_html=True)
    st.markdown("""
    <style>
    [data-testid="stMain"]:has(#rng-meeting-marker) { background: radial-gradient(ellipse at right top, #dbe7f8, transparent 50%), #f2f5fa; }
    .stApp:has(#rng-meeting-marker), .stApp:has(#rng-meeting-marker) h1,
    .stApp:has(#rng-meeting-marker) h2, .stApp:has(#rng-meeting-marker) h3,
    .stApp:has(#rng-meeting-marker) [data-testid="stMarkdownContainer"],
    .stApp:has(#rng-meeting-marker) input, .stApp:has(#rng-meeting-marker) textarea,
    .stApp:has(#rng-meeting-marker) label { font-family: Arial, Tahoma, sans-serif !important; font-style: normal !important; }
    .stApp:has(#rng-meeting-marker) .block-container { max-width: 1440px; }
    .rng-meeting-hero { background: radial-gradient(ellipse at right top, #345d88, transparent 65%), #172f4e; border: 1px solid #315174; border-radius: 24px; padding: 30px 34px; margin: 8px 0 14px; }
    .rng-meeting-eyebrow { color: #b2cbed; font-size: 11px; font-weight: 700; letter-spacing: 2px; margin-bottom: 14px; }
    .rng-meeting-hero h1 { color: #f6f9ff !important; font-size: clamp(30px, 3.4vw, 44px) !important; font-weight: 750; letter-spacing: -1.2px; padding: 0; margin: 0 0 12px; overflow-wrap: anywhere; }
    .rng-meeting-hero p { color: #c5d7ec; font-size: 15px; line-height: 1.8; margin: 0; }
    .rng-meeting-total { display: inline-block; color: #dce9fb; background: #a9ccff12; border: 1px solid #a9ccff30; border-radius: 100px; padding: 6px 12px; font-size: 11px; margin-top: 18px; }
    .st-key-rng_meeting_add_action .stButton button, .st-key-rng_meeting_save_action .stButton button { color: #fff; background: #254d7e; border: 1px solid #254d7e; min-height: 46px; border-radius: 11px; }
    .st-key-rng_meeting_add_action .stButton button:hover, .st-key-rng_meeting_save_action .stButton button:hover { color: #fff; background: #18395f; }
    .st-key-rng_add_meeting_form { background: #fff; border: 1px solid #d5e1f2; border-radius: 18px; padding: 22px; margin-bottom: 20px; }
    [class*="st-key-rng_meeting_card_"] { background: #fff; border: 1px solid #dce5f1; border-radius: 18px; padding: 24px; min-height: 270px; box-shadow: 0 8px 25px #15385f07; }
    .rng-meeting-card-top { display: flex; align-items: center; justify-content: space-between; margin-bottom: 19px; }
    .rng-meeting-icon { display: flex; align-items: center; justify-content: center; width: 46px; height: 46px; color: #567ead; background: #eaf1fc; border-radius: 14px; }
    .rng-meeting-icon svg { width: 24px; height: 24px; }
    .rng-meeting-index { color: #8698b0; font-size: 11px; letter-spacing: 1px; }
    .meeting-card-title { color: #243e60; font-size: 19px; font-weight: 700; line-height: 1.5; margin: 0 0 10px; overflow-wrap: anywhere; }
    .meeting-card-text { color: #6f839e; font-size: 12px; line-height: 1.8; margin-bottom: 13px; }
    [class*="st-key-rng_meeting_card_"] [data-testid="stHorizontalBlock"] > [data-testid="stColumn"] { min-width: 0 !important; flex: 1; }
    .rng-meeting-card-footer { color: #8092aa; font-size: 10px; letter-spacing: .6px; border-top: 1px solid #e8edf5; padding-top: 14px; margin-bottom: 6px; }
    .st-key-rng_meeting_board .stButton button { color: #48658b; background: #f5f8fd; border: 1px solid #dfe7f3; border-radius: 9px; min-height: 44px; }
    .st-key-rng_meeting_board .stButton button:hover { background: #eaf1fc; color: #234b7b; border-color: #8faed7; }
    [class*="st-key-rng_meeting_open_action_"] .stButton button { background: #244d7c; border-color: #244d7c; color: #fff; }
    [class*="st-key-rng_meeting_open_action_"] .stButton button:hover { background: #17395f; color: #fff; }
    .delete-warning { color: #8c4050; background: #fff2f4; border: 1px solid #f1d4dc; border-radius: 12px; padding: 15px; font-size: 13px; line-height: 1.8; overflow-wrap: anywhere; }
    .rng-meeting-empty { text-align: center; padding: 50px 24px; border: 1px dashed #bdcee5; border-radius: 18px; background: #fff9; color: #70859f; }
    .rng-meeting-empty strong { display: block; color: #304e75; font-size: 21px; margin-bottom: 10px; }
    .rng-meeting-stats { display: grid; grid-template-columns: repeat(3, 1fr); gap: 12px; margin: 4px 0 10px; }
    .rng-meeting-stat { background: #fff; border: 1px solid #dce5f2; border-radius: 14px; padding: 17px 20px; }
    .rng-meeting-stat strong { display: block; color: #294a74; font-size: 25px; margin-bottom: 6px; }
    .rng-meeting-stat span { color: #6a809f; font-size: 12px; }
    .rng-meeting-section { color: #294569; font-size: 21px; font-weight: 700; margin: 12px 0 8px; }
    .rng-meeting-help { color: #7388a4; font-size: 13px; line-height: 1.8; margin-bottom: 12px; }
    .st-key-rng_meeting_agenda [data-testid="stExpander"] { background: #fff; border: 1px solid #dce5f1; border-radius: 14px; overflow: hidden; margin-bottom: 10px; }
    .st-key-rng_meeting_agenda [data-testid="stExpander"] summary { color: #2e4e78; min-height: 60px; }
    .st-key-rng_meeting_agenda [data-testid="stExpander"] summary p { font-weight: 600; overflow-wrap: anywhere; }
    .st-key-rng_meeting_agenda [data-testid="stExpander"] summary:hover { background: #f2f6fd; }
    .st-key-rng_meeting_agenda input, .st-key-rng_meeting_agenda textarea { direction: rtl; text-align: start; background: #f9fbff; color: #2f486b; min-height: 44px; }
    .st-key-rng_meeting_agenda label p { color: #5b7394; font-size: 12px; }
    .st-key-rng_meeting_exports { background: #e8eff9; border: 1px solid #d9e4f3; padding: 22px; border-radius: 18px; margin-top: 20px; }
    .st-key-rng_meeting_exports .stDownloadButton button { background: #fff; color: #345781; border: 1px solid #c5d6ed; min-height: 46px; border-radius: 10px; }
    .stApp:has(#rng-meeting-marker) button:focus-visible { outline: 3px solid #9ab9e8; outline-offset: 2px; }
    @media (max-width: 1100px) and (min-width: 701px) {
        .st-key-rng_meeting_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_meeting_card_"]) { flex-wrap: wrap; }
        .st-key-rng_meeting_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_meeting_card_"]) > [data-testid="stColumn"] { min-width: calc(50% - 1rem); }
    }
    @media (max-width: 700px) {
        .rng-meeting-hero { padding: 26px 22px; border-radius: 20px; }
        .rng-meeting-stats { gap: 8px; }
        .rng-meeting-stat { padding: 15px 10px; }
        .rng-meeting-stat span { font-size: 11px; }
        .st-key-rng_meeting_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_meeting_card_"]) { flex-wrap: wrap; }
        .st-key-rng_meeting_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_meeting_card_"]) > [data-testid="stColumn"] { min-width: 100%; }
        [class*="st-key-rng_meeting_card_"] { padding: 20px; }
        .st-key-rng_meeting_exports { padding: 18px; }
    }
    </style>
    """, unsafe_allow_html=True)

    def meeting_record_counts(name):
        records = st.session_state.get("meeting_data_" + name, [])
        populated = [row for row in records if any(str(row[col]).strip() for col in (0, 1, 2, 4))]
        return len(populated), sum(row[5] == "Done" for row in populated)

    def meeting_card_heading(name, index):
        count, done = meeting_record_counts(name)
        description = f"{count} agenda items · {done} completed" if count else "Ready for your next conversation."
        st.markdown(
            '<div class="rng-meeting-card-top"><div class="rng-meeting-icon">'
            '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.5" stroke-linecap="round" aria-hidden="true"><circle cx="9" cy="8" r="3"/><path d="M3 21v-3a6 6 0 0 1 12 0v3M16 5a3 3 0 0 1 0 6m2 4a5 5 0 0 1 3 5v1"/></svg>'
            f'</div><span class="rng-meeting-index">MEETING {index + 1:02}</span></div>'
            f'<div class="meeting-card-title">{escape(str(name))}</div>'
            f'<div class="meeting-card-text">{description}</div>'
            '<div class="rng-meeting-card-footer">AGENDA / DECISIONS / FOLLOW-UP</div>',
            unsafe_allow_html=True
        )

    if "meeting_cards" not in st.session_state:
        st.session_state.meeting_cards = []

    if "selected_meeting" not in st.session_state:
        st.session_state.selected_meeting = None

    if "editing_meeting" not in st.session_state:
        st.session_state.editing_meeting = None

    if "deleting_meeting" not in st.session_state:
        st.session_state.deleting_meeting = None

    if "adding_meeting" not in st.session_state:
        st.session_state.adding_meeting = False

    # =========================================================
    # MEETING DETAIL PAGE
    # =========================================================

    if st.session_state.selected_meeting is not None:

        meeting_name = st.session_state.selected_meeting

        if st.button(
            "Back to Meeting Minutes",
            key="back_to_meetings"
        ):
            st.session_state.selected_meeting = None
            rng_rerun()

        st.markdown(
            '<div class="rng-meeting-hero"><div class="rng-meeting-eyebrow">MEETING WORKSPACE</div>'
            f'<h1>{escape(str(meeting_name))}</h1>'
            '<p>Keep the discussion, the decision, and the next step together.</p></div>',
            unsafe_allow_html=True
        )

        # =====================================================
        # MEETING TABLE DATA
        # =====================================================

        meeting_key = "meeting_data_" + meeting_name

        if meeting_key not in st.session_state:
            st.session_state[meeting_key] = [
                ["", "", "", "", "", ""]
                for _ in range(11)
            ]
        else:
            st.session_state[meeting_key] = (
                st.session_state[meeting_key][:11]
            )

            while len(st.session_state[meeting_key]) < 11:
                st.session_state[meeting_key].append(
                    ["", "", "", "", "", ""]
                )

        stats_slot = st.empty()
        st.markdown(
            '<div class="rng-meeting-section">Agenda &amp; action items</div>'
            '<div class="rng-meeting-help">Open an item to edit its discussion and follow-up. Save your changes below. All 11 items are included in the A4 PDF.</div>',
            unsafe_allow_html=True
        )
        with st.container(key="rng_meeting_agenda"):
            for row in range(11):
                values = list(st.session_state[meeting_key][row])
                label_title = st.session_state.get(f"{meeting_key}_cell_{row}_1", values[1])
                item_label = f"{row + 1:02} · {label_title or values[0] or 'Agenda item'}"
                has_content = any(str(values[col]).strip() for col in (0, 1, 2, 4))
                with st.expander(item_label, expanded=row == 0 or has_content):
                    agenda_col, title_col = st.columns([1, 2])
                    with agenda_col:
                        values[0] = st.text_input("Agenda", value=values[0], key=f"{meeting_key}_cell_{row}_0")
                    with title_col:
                        values[1] = st.text_input("Title", value=values[1], key=f"{meeting_key}_cell_{row}_1")
                    owner_col, discussion_col, action_col = st.columns([2, 1, 1])
                    with owner_col:
                        values[2] = st.text_input("Person in charge", value=values[2], key=f"{meeting_key}_cell_{row}_2")
                    with discussion_col:
                        options = ["Raised", "not Raised"]
                        values[3] = st.selectbox("Discussion status", options, index=options.index(values[3]) if values[3] in options else 0, key=f"{meeting_key}_select_{row}_3")
                    with action_col:
                        options = ["Done", "not Started", "In Progress", "Stopped"]
                        values[5] = st.selectbox("Action status", options, index=options.index(values[5]) if values[5] in options else 0, key=f"{meeting_key}_select_{row}_5")
                    values[4] = st.text_area("Minute / decision", value=values[4], key=f"{meeting_key}_cell_{row}_4", height=110)
                st.session_state[meeting_key][row] = values

        record_total, record_done = meeting_record_counts(meeting_name)
        stats_slot.markdown(
            '<div class="rng-meeting-stats">'
            f'<div class="rng-meeting-stat"><strong>{record_total}</strong><span>Agenda items</span></div>'
            f'<div class="rng-meeting-stat"><strong>{record_done}</strong><span>Completed</span></div>'
            f'<div class="rng-meeting-stat"><strong>{record_total - record_done}</strong><span>Remaining</span></div></div>',
            unsafe_allow_html=True
        )
        with st.container(key="rng_meeting_save_action"):
            if st.button("Save meeting", key="save_meeting_records", use_container_width=True):
                try:
                    rng_save()
                    st.success("Meeting records saved.")
                except Exception as exc:
                    st.error("Could not save your meeting. Please retry. " + str(exc))

        # =====================================================
        # PDF GENERATION
        # =====================================================

        import io
        import os
        import re
        import arabic_reshaper

        from bidi.algorithm import get_display

        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        def create_meeting_pdf(meeting_data):

            buffer = io.BytesIO()

            page_width, page_height = landscape(A4)

            pdf = canvas.Canvas(
                buffer,
                pagesize=(page_width, page_height)
            )

            # -------------------------------------------------
            # PDF FONT SUPPORT
            # -------------------------------------------------

            font_name = rng_pdf_font_name()

            # -------------------------------------------------
            # PAGE AND TABLE DIMENSIONS
            # -------------------------------------------------

            margin_x = 7 * mm
            margin_y = 7 * mm

            table_width = page_width - 2 * margin_x
            table_height = page_height - 2 * margin_y

            number_of_rows = 12

            row_height = table_height / number_of_rows

            column_ratios = [
                1.0,
                2.4,
                1.0,
                1.0,
                1.0,
                2.4
            ]

            ratio_total = sum(column_ratios)

            column_widths = [
                table_width * ratio / ratio_total
                for ratio in column_ratios
            ]

            column_order = [5, 4, 3, 2, 1, 0]

            pdf_headers = [
                "Status",
                "Minute",
                "Status",
                "Person in charge",
                "Title",
                "Agenda"
            ]

            # -------------------------------------------------
            # PERSIAN AND ENGLISH TEXT SUPPORT
            # -------------------------------------------------

            def format_text(value):

                value = str(
                    value if value is not None else ""
                )

                return get_display(
                    arabic_reshaper.reshape(value)
                )

            def text_width(value, font_size):

                return rng_pdf_width(format_text(value), font_size)

            # -------------------------------------------------
            # WRAP TEXT INTO CELL
            # -------------------------------------------------

            def wrap_text(value, max_width, font_size):

                value = str(
                    value if value is not None else ""
                ).strip()

                words = re.split(r"\s+", value)

                lines = []
                current_line = ""

                for word in words:

                    if not word:
                        continue

                    candidate = (
                        current_line + " " + word
                        if current_line
                        else word
                    )

                    if text_width(
                        candidate,
                        font_size
                    ) <= max_width:

                        current_line = candidate

                    else:

                        if current_line:
                            lines.append(current_line)

                        # Split oversized words so they cannot
                        # cross the cell's right border.

                        current_line = ""

                        for character in word:

                            candidate = current_line + character

                            if (
                                current_line
                                and text_width(
                                    candidate,
                                    font_size
                                ) > max_width
                            ):

                                lines.append(current_line)
                                current_line = character

                            else:
                                current_line = candidate

                if current_line:
                    lines.append(current_line)

                return lines or [""]

            # -------------------------------------------------
            # DRAW TEXT IN CELL
            # -------------------------------------------------

            def draw_cell_text(
                value,
                x,
                y,
                width,
                height,
                is_header=False
            ):

                padding = 4

                available_width = max(
                    width - 2 * padding,
                    1
                )

                font_size = 9.0

                while font_size > 6.0:

                    lines = wrap_text(
                        value,
                        available_width,
                        font_size
                    )

                    line_height = font_size * 1.3

                    if (
                        len(lines) * line_height
                        <= height - 2 * padding
                    ):
                        break

                    font_size -= 0.5

                lines = wrap_text(
                    value,
                    available_width,
                    font_size
                )

                line_height = font_size * 1.3

                max_lines = max(
                    1,
                    int(
                        (height - 2 * padding)
                        / line_height
                    )
                )

                # A fixed A4 table cannot fit unlimited text.
                # Truncate overflow with an ellipsis.

                if len(lines) > max_lines:

                    lines = lines[:max_lines]

                    last_line = lines[-1]

                    while (
                        last_line
                        and text_width(
                            last_line + "...",
                            font_size
                        ) > available_width
                    ):
                        last_line = last_line[:-1]

                    lines[-1] = last_line + "..."

                pdf.setFont(font_name, font_size)

                if is_header:
                    pdf.setFillColorRGB(1, 1, 1)
                else:
                    pdf.setFillColorRGB(
                        0.12,
                        0.17,
                        0.23
                    )

                text_height = len(lines) * line_height

                baseline = (
                    y
                    + (height + text_height) / 2
                    - line_height
                    + 2
                )

                for line in lines:

                    rendered = format_text(line)

                    rng_pdf_draw(pdf, rendered, x + width / 2, baseline, font_size, "center")

                    baseline -= line_height

            # -------------------------------------------------
            # DRAW TABLE HEADER
            # -------------------------------------------------

            current_y = (
                page_height
                - margin_y
                - row_height
            )

            current_x = margin_x

            for column_index, header in enumerate(
                pdf_headers
            ):

                cell_width = column_widths[column_index]

                pdf.setFillColorRGB(
                    23 / 255,
                    35 / 255,
                    51 / 255
                )

                pdf.rect(
                    current_x,
                    current_y,
                    cell_width,
                    row_height,
                    fill=1,
                    stroke=0
                )

                draw_cell_text(
                    header,
                    current_x,
                    current_y,
                    cell_width,
                    row_height,
                    is_header=True
                )

                current_x += cell_width

            # -------------------------------------------------
            # DRAW 11 EDITABLE DATA ROWS
            # -------------------------------------------------

            for row_index in range(11):

                current_y = (
                    page_height
                    - margin_y
                    - row_height * (row_index + 2)
                )

                current_x = margin_x

                for column_index, data_col in enumerate(
                    column_order
                ):

                    cell_width = column_widths[column_index]

                    value = meeting_data[
                        row_index
                    ][data_col]

                    if data_col == 3 and not value:
                        value = "Raised"

                    if data_col == 5 and not value:
                        value = "Done"

                    pdf.setFillColorRGB(1, 1, 1)

                    pdf.rect(
                        current_x,
                        current_y,
                        cell_width,
                        row_height,
                        fill=1,
                        stroke=0
                    )

                    draw_cell_text(
                        value,
                        current_x,
                        current_y,
                        cell_width,
                        row_height
                    )

                    current_x += cell_width

            # -------------------------------------------------
            # DRAW TABLE BORDERS
            # -------------------------------------------------

            pdf.setStrokeColorRGB(
                0.72,
                0.76,
                0.81
            )

            pdf.setLineWidth(0.45)

            for row_number in range(number_of_rows + 1):

                y = (
                    page_height
                    - margin_y
                    - row_number * row_height
                )

                pdf.line(
                    margin_x,
                    y,
                    margin_x + table_width,
                    y
                )

            x = margin_x

            pdf.line(
                x,
                margin_y,
                x,
                margin_y + table_height
            )

            for cell_width in column_widths:

                x += cell_width

                pdf.line(
                    x,
                    margin_y,
                    x,
                    margin_y + table_height
                )

            # -------------------------------------------------
            # COMPLETE SINGLE-PAGE PDF
            # -------------------------------------------------

            pdf.showPage()
            pdf.save()

            buffer.seek(0)

            return buffer.getvalue()

        # =====================================================
        # PDF DOWNLOAD BUTTON
        # =====================================================

        with st.container(key="rng_meeting_exports"):
            st.markdown('<div class="rng-meeting-section">Take it with you</div><div class="rng-meeting-help">Download or share your meeting as a single A4 landscape PDF.</div>', unsafe_allow_html=True)

            try:

                pdf_bytes = create_meeting_pdf(
                    st.session_state[meeting_key]
                )

                safe_meeting_name = re.sub(
                    r'[^\w.-]+',
                    "_",
                    meeting_name
                )

                meeting_pdf_filename = f"Meeting_Minutes_{safe_meeting_name}.pdf"
                download_col, share_col = st.columns(2)
                with download_col:
                    st.download_button(
                        label="Download Meeting Table (A4 PDF)",
                        data=pdf_bytes,
                        file_name=meeting_pdf_filename,
                        mime="application/pdf",
                        key="download_meeting_pdf",
                        use_container_width=True
                    )
                with share_col:
                    rng_pdf_share_button(pdf_bytes, meeting_pdf_filename)

            except Exception as error:

                st.error(
                    f"Unable to generate the meeting PDF: {error}"
                )

    # =========================================================
    # MEETING MINUTES MAIN PAGE
    # =========================================================

    else:

        meeting_total = len(st.session_state.meeting_cards)
        st.markdown(
            '<div class="rng-meeting-hero"><div class="rng-meeting-eyebrow">DISCUSS. DECIDE. MOVE FORWARD.</div>'
            '<h1>Meeting Minutes</h1><p>Good conversations deserve a clear record.</p>'
            f'<span class="rng-meeting-total">{meeting_total} {"meeting" if meeting_total == 1 else "meetings"} in your library</span></div>',
            unsafe_allow_html=True
        )

        # =====================================================
        # ADD MEETING
        # =====================================================

        if st.session_state.adding_meeting:
            with st.container(key="rng_add_meeting_form"):

                st.markdown(
                    '<div class="section-title">Add New Meeting</div>',
                    unsafe_allow_html=True
                )

                new_meeting_title = st.text_input(
                    "Meeting title",
                    placeholder="Enter meeting title",
                    key="new_meeting_title"
                )

                add_col1, add_col2 = st.columns(2)

                with add_col1:

                    if st.button(
                        "Add Meeting",
                        key="add_meeting_confirm",
                        use_container_width=True
                    ):

                        if new_meeting_title.strip():

                            if new_meeting_title.strip() in st.session_state.meeting_cards:
                                st.error("A meeting with this title already exists.")
                            else:
                                st.session_state.meeting_cards.append(new_meeting_title.strip())
                                st.session_state.adding_meeting = False
                                rng_rerun()

                with add_col2:

                    if st.button(
                        "Cancel",
                        key="add_meeting_cancel",
                        use_container_width=True
                    ):

                        st.session_state.adding_meeting = False
                        rng_rerun()

        else:
            with st.container(key="rng_meeting_add_action"):

                if st.button(
                    "+ Add Meeting",
                    key="add_meeting_button",
                    use_container_width=False
                ):

                    st.session_state.adding_meeting = True
                    rng_rerun()

        st.markdown("<br>", unsafe_allow_html=True)

        # =====================================================
        # MEETING CARDS
        # =====================================================

        if not st.session_state.meeting_cards:
            st.markdown('<div class="rng-meeting-empty"><strong>Your next conversation starts here.</strong>Add a meeting to capture its agenda, decisions, and next steps.</div>', unsafe_allow_html=True)
        with st.container(key="rng_meeting_board"):
            cols = st.columns(3)

            for i, meeting in enumerate(
                st.session_state.meeting_cards
            ):

                with cols[i % 3], st.container(key=f"rng_meeting_card_{i}"):

                    meeting_card_heading(meeting, i)

                    # -------------------------------------------------
                    # DELETE CONFIRMATION
                    # -------------------------------------------------

                    if st.session_state.deleting_meeting == i:

                        st.markdown(
                            f'<div class="delete-warning">Delete "{escape(str(meeting))}"? This cannot be undone.</div>',
                            unsafe_allow_html=True
                        )

                        confirm_col, cancel_col = st.columns(2)

                        with confirm_col:

                            if st.button(
                                "Confirm Delete",
                                key=f"confirm_delete_{i}",
                                use_container_width=True
                            ):

                                st.session_state.meeting_cards.pop(i)

                                st.session_state.deleting_meeting = None

                                rng_rerun()

                        with cancel_col:

                            if st.button(
                                "Cancel",
                                key=f"cancel_delete_{i}",
                                use_container_width=True
                            ):

                                st.session_state.deleting_meeting = None

                                rng_rerun()

                    # -------------------------------------------------
                    # EDIT MODE
                    # -------------------------------------------------

                    elif st.session_state.editing_meeting == i:

                        new_title = st.text_input(
                            "Meeting title",
                            value=meeting,
                            key=f"meeting_edit_input_{i}"
                        )

                        edit_col1, edit_col2 = st.columns(2)

                        with edit_col1:

                            if st.button(
                                "Save",
                                key=f"meeting_save_{i}",
                                use_container_width=True
                            ):

                                if new_title.strip():

                                    renamed = new_title.strip()
                                    if renamed != meeting and renamed in st.session_state.meeting_cards:
                                        st.error("A meeting with this title already exists.")
                                    else:
                                        old_key = "meeting_data_" + meeting
                                        if renamed != meeting and old_key in st.session_state:
                                            st.session_state["meeting_data_" + renamed] = copy.deepcopy(st.session_state[old_key])
                                        st.session_state.meeting_cards[i] = renamed
                                        st.session_state.editing_meeting = None
                                        rng_rerun()

                        with edit_col2:

                            if st.button(
                                "Cancel",
                                key=f"meeting_cancel_{i}",
                                use_container_width=True
                            ):

                                st.session_state.editing_meeting = None

                                rng_rerun()

                    # -------------------------------------------------
                    # NORMAL MODE
                    # -------------------------------------------------

                    else:

                        button_col1, button_col2, button_col3 = st.columns(3)

                        with button_col1, st.container(key=f"rng_meeting_open_action_{i}"):

                            open_button = st.button(
                                "Open",
                                key=f"meeting_open_{i}",
                                use_container_width=True
                            )

                        with button_col2:

                            edit_button = st.button(
                                "Edit",
                                key=f"meeting_edit_{i}",
                                use_container_width=True
                            )

                        with button_col3:

                            delete_button = st.button(
                                "Delete",
                                key=f"meeting_delete_{i}",
                                use_container_width=True
                            )

                        if open_button:

                            st.session_state.selected_meeting = meeting

                            rng_rerun()

                        if edit_button:

                            st.session_state.editing_meeting = i

                            rng_rerun()

                        if delete_button:

                            st.session_state.deleting_meeting = i

                            rng_rerun()






elif page == "Warehouse":

    import io
    import os
    import re
    import zipfile
    import hashlib
    import xml.etree.ElementTree as ET

    from pathlib import Path
    from uuid import uuid4
    from html import escape

    # =========================================================
    # ORIGINAL EXCEL TEMPLATE
    # =========================================================

    WAREHOUSE_TEMPLATE = (
        Path(__file__).resolve().parent
        / "\u0641\u0631\u0645 \u062f\u0631\u062e\u0648\u0627\u0633\u062a \u06a9\u0627\u0644\u0627.xlsx"
    )

    def load_warehouse_template():

        if not WAREHOUSE_TEMPLATE.is_file():
            raise FileNotFoundError(
                "Warehouse Excel template not found: "
                + str(WAREHOUSE_TEMPLATE)
            )

        ns = {
            "x": (
                "http://schemas.openxmlformats.org/"
                "spreadsheetml/2006/main"
            )
        }

        with zipfile.ZipFile(WAREHOUSE_TEMPLATE) as z:

            strings_xml = ET.fromstring(
                z.read("xl/sharedStrings.xml")
            )

            strings = []

            for item in strings_xml.findall("x:si", ns):
                value = "".join(
                    element.text or ""
                    for element in item.findall(".//x:t", ns)
                )
                strings.append(value)

            logo = z.read("xl/media/image1.png")

        if len(strings) < 10:
            raise ValueError(
                "The original Excel form labels are missing."
            )

        date_number_labels = strings[5].splitlines()
        defaults = strings[9].splitlines()
        requester_lines = strings[8].splitlines()

        return {
            "headers": strings[0:5],
            "date_label": date_number_labels[0],
            "number_label": date_number_labels[1],
            "title": strings[6],
            "approver_label": strings[7].split(":")[0] + ":",
            "requester_label": requester_lines[0],
            "signature_label": "\u0627\u0645\u0636\u0627",
            "default_requester": requester_lines[1].strip(),
            "default_date": defaults[0].strip(),
            "default_number": defaults[1].strip(),
            "logo": logo
        }

    warehouse_template = None
    template_error = None

    try:
        warehouse_template = load_warehouse_template()
    except Exception as error:
        template_error = str(error)

    # =========================================================
    # SESSION STATE
    # =========================================================

    if "warehouse_cards" not in st.session_state:
        st.session_state.warehouse_cards = []

    if "selected_warehouse" not in st.session_state:
        st.session_state.selected_warehouse = None

    if "editing_warehouse" not in st.session_state:
        st.session_state.editing_warehouse = None

    if "deleting_warehouse" not in st.session_state:
        st.session_state.deleting_warehouse = None

    if "adding_warehouse" not in st.session_state:
        st.session_state.adding_warehouse = False

    if "warehouse_card_ids" not in st.session_state:
        st.session_state.warehouse_card_ids = [
            uuid4().hex
            for _ in st.session_state.warehouse_cards
        ]

    if "warehouse_forms" not in st.session_state:
        st.session_state.warehouse_forms = {}

    while len(st.session_state.warehouse_card_ids) < len(
        st.session_state.warehouse_cards
    ):
        st.session_state.warehouse_card_ids.append(
            uuid4().hex
        )

    # =========================================================
    # EMPTY WAREHOUSE FORM
    # =========================================================

    def create_empty_warehouse_form():

        return {
            "date": "",
            "number": "",
            "items": [
                ["", "", "", "", ""]
                for _ in range(10)
            ],
            "requester": "",
            "approver": "",
            "requester_signature": "",
            "approver_signature": ""
        }

    # =========================================================
    # WAREHOUSE CSS
    # =========================================================

    st.markdown('<div id="rng-warehouse-marker"></div>', unsafe_allow_html=True)
    st.markdown("""
    <style>
    [data-testid="stMain"]:has(#rng-warehouse-marker) { background: radial-gradient(ellipse at right top, #dce8f9, transparent 55%), #f2f5fa; }
    .stApp:has(#rng-warehouse-marker), .stApp:has(#rng-warehouse-marker) h1,
    .stApp:has(#rng-warehouse-marker) h2, .stApp:has(#rng-warehouse-marker) h3,
    .stApp:has(#rng-warehouse-marker) [data-testid="stMarkdownContainer"],
    .stApp:has(#rng-warehouse-marker) input, .stApp:has(#rng-warehouse-marker) textarea,
    .stApp:has(#rng-warehouse-marker) label { font-family: Arial, Tahoma, sans-serif !important; font-style: normal !important; }
    .stApp:has(#rng-warehouse-marker) .block-container { max-width: 1440px; }
    .rng-warehouse-hero { background: radial-gradient(ellipse at right top, #355e89, transparent 65%), #172f4e; border: 1px solid #315174; border-radius: 24px; padding: 30px 34px; margin: 8px 0 14px; }
    .rng-warehouse-eyebrow { color: #b1cae9; font-size: 11px; font-weight: 700; letter-spacing: 2px; margin-bottom: 14px; }
    .rng-warehouse-hero h1 { color: #f6f9ff !important; font-size: clamp(30px, 3.4vw, 44px) !important; font-weight: 750; letter-spacing: -1.2px; margin: 0 0 12px; padding: 0; overflow-wrap: anywhere; }
    .rng-warehouse-hero p { color: #c6d8ed; font-size: 15px; line-height: 1.8; margin: 0; }
    .rng-warehouse-total { display: inline-block; color: #dce9fb; background: #a9ccff12; border: 1px solid #a9ccff30; border-radius: 100px; padding: 6px 12px; font-size: 11px; margin-top: 18px; }
    .st-key-rng_warehouse_add_action .stButton button, .st-key-rng_warehouse_save_action .stButton button { color: #fff; background: #254d7e; border: 1px solid #254d7e; min-height: 46px; border-radius: 11px; }
    .st-key-rng_warehouse_add_action .stButton button:hover, .st-key-rng_warehouse_save_action .stButton button:hover { background: #18395f; color: #fff; }
    .st-key-rng_warehouse_add_form { background: #fff; border: 1px solid #d7e3f2; padding: 24px; border-radius: 18px; margin-bottom: 20px; }
    [class*="st-key-rng_warehouse_card_"] { background: #fff; border: 1px solid #dce5f1; border-radius: 18px; padding: 24px; min-height: 270px; box-shadow: 0 8px 25px #15385f07; }
    .rng-warehouse-card-top { display: flex; align-items: center; justify-content: space-between; margin-bottom: 19px; }
    .rng-warehouse-icon { display: flex; align-items: center; justify-content: center; width: 46px; height: 46px; color: #5279aa; background: #eaf1fc; border-radius: 14px; }
    .rng-warehouse-icon svg { width: 25px; height: 25px; }
    .rng-warehouse-index { color: #8698b0; font-size: 11px; letter-spacing: 1px; }
    .warehouse-card-title { color: #243e60; font-size: 19px; font-weight: 700; line-height: 1.5; margin: 0 0 10px; overflow-wrap: anywhere; }
    .warehouse-card-text { color: #6f839e; font-size: 12px; line-height: 1.8; margin-bottom: 13px; overflow-wrap: anywhere; }
    .rng-warehouse-card-footer { color: #8092aa; font-size: 10px; letter-spacing: .6px; border-top: 1px solid #e8edf5; padding-top: 14px; margin-bottom: 6px; }
    [class*="st-key-rng_warehouse_card_"] [data-testid="stHorizontalBlock"] > [data-testid="stColumn"] { min-width: 0 !important; flex: 1; }
    .st-key-rng_warehouse_board .stButton button { color: #48658b; background: #f5f8fd; border: 1px solid #dfe7f3; border-radius: 9px; min-height: 44px; }
    .st-key-rng_warehouse_board .stButton button:hover { background: #eaf1fc; color: #234b7b; border-color: #8faed7; }
    [class*="st-key-rng_warehouse_open_"] .stButton button { background: #244d7c; border-color: #244d7c; color: #fff; }
    [class*="st-key-rng_warehouse_open_"] .stButton button:hover { background: #17395f; color: #fff; }
    .warehouse-delete-warning { color: #8c4050; background: #fff2f4; border: 1px solid #f1d4dc; border-radius: 12px; padding: 15px; font-size: 13px; line-height: 1.8; overflow-wrap: anywhere; }
    .rng-warehouse-empty { text-align: center; padding: 50px 24px; border: 1px dashed #bdcee5; border-radius: 18px; background: #fff9; color: #70859f; }
    .rng-warehouse-empty strong { display: block; color: #304e75; font-size: 21px; margin-bottom: 10px; }
    .rng-warehouse-section { color: #294569; font-size: 21px; font-weight: 700; margin: 4px 0 8px; }
    .rng-warehouse-help { color: #7388a4; font-size: 13px; line-height: 1.8; margin-bottom: 16px; }
    .warehouse-form-title { direction: rtl; color: #48688e; font-size: 18px; font-weight: 600; text-align: right; margin: 4px 0 18px; }
    .st-key-rng_warehouse_meta, .st-key-rng_warehouse_approval { background: #fff; border: 1px solid #dce5f2; padding: 24px; border-radius: 18px; margin: 6px 0 12px; }
    [class*="st-key-rng_warehouse_item_"] { background: #fff; border: 1px solid #dce5f1; border-radius: 14px; padding: 18px; margin-bottom: 10px; }
    .rng-warehouse-item-label { color: #6c87aa; font-size: 10px; font-weight: 700; letter-spacing: 1px; margin-bottom: 3px; }
    .st-key-warehouse_detail input, .st-key-warehouse_detail textarea { direction: rtl; text-align: start; unicode-bidi: plaintext; background: #f9fbff; color: #2f486b; min-height: 44px; }
    .st-key-warehouse_detail [data-testid="stWidgetLabel"] { direction: rtl; text-align: right; }
    .st-key-warehouse_detail label p { color: #587193; font-size: 12px; }
    .st-key-rng_warehouse_export { background: #e8eff9; border: 1px solid #d9e4f3; padding: 22px; border-radius: 18px; margin-top: 20px; }
    .st-key-rng_warehouse_export .stDownloadButton button { background: #fff; color: #345781; border: 1px solid #c5d6ed; min-height: 46px; border-radius: 10px; }
    .stApp:has(#rng-warehouse-marker) button:focus-visible { outline: 3px solid #9ab9e8; outline-offset: 2px; }
    @media (max-width: 1100px) and (min-width: 701px) {
        .st-key-rng_warehouse_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_warehouse_card_"]) { flex-wrap: wrap; }
        .st-key-rng_warehouse_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_warehouse_card_"]) > [data-testid="stColumn"] { min-width: calc(50% - 1rem); }
    }
    @media (max-width: 700px) {
        .rng-warehouse-hero { padding: 26px 22px; border-radius: 20px; }
        [class*="st-key-rng_warehouse_item_"] [data-testid="stHorizontalBlock"] { flex-direction: column-reverse; flex-wrap: nowrap; }
        [class*="st-key-rng_warehouse_item_"] [data-testid="stColumn"] { width: 100% !important; min-width: 100%; flex: auto; }
        .st-key-rng_warehouse_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_warehouse_card_"]) { flex-wrap: wrap; }
        .st-key-rng_warehouse_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_warehouse_card_"]) > [data-testid="stColumn"] { min-width: 100%; }
        [class*="st-key-rng_warehouse_card_"] { padding: 20px; }
        .st-key-rng_warehouse_meta, .st-key-rng_warehouse_approval, .st-key-rng_warehouse_export { padding: 18px; }
    }
    </style>
    """, unsafe_allow_html=True)

    def warehouse_item_count(data):
        return sum(any(str(value).strip() for value in row) for row in data.get("items", []))

    def warehouse_card_heading(name, index):
        record_id = st.session_state.warehouse_card_ids[index]
        data = st.session_state.warehouse_forms.get(record_id, {})
        count = warehouse_item_count(data)
        description = f"{count} line {('item' if count == 1 else 'items')}" if count else "Ready for your next request."
        if data.get("number"):
            description += " · Request " + str(data["number"])
        st.markdown(
            '<div class="rng-warehouse-card-top"><div class="rng-warehouse-icon">'
            '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.5" stroke-linejoin="round" aria-hidden="true"><path d="m3 7 9-4 9 4v11l-9 4-9-4Zm0 0 9 4 9-4m-9 4v11M7 5l10 4v5"/></svg>'
            f'</div><span class="rng-warehouse-index">WAREHOUSE {index + 1:02}</span></div>'
            f'<div class="warehouse-card-title">{escape(str(name))}</div>'
            f'<div class="warehouse-card-text">{escape(description)}</div>'
            '<div class="rng-warehouse-card-footer">PURCHASE REQUEST / ITEMS / APPROVALS</div>',
            unsafe_allow_html=True
        )

    # =========================================================
    # A5 LANDSCAPE PDF GENERATOR
    # =========================================================

    def create_warehouse_pdf(form, template):

        import arabic_reshaper
        from bidi.algorithm import get_display

        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import A5, landscape
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.lib.utils import ImageReader

        buffer = io.BytesIO()

        page_width, page_height = landscape(A5)

        pdf = canvas.Canvas(
            buffer,
            pagesize=(page_width, page_height),
            pageCompression=1
        )

        pdf.setTitle(template["title"])

        # =====================================================
        # UNICODE FONT FOR PERSIAN
        # =====================================================

        font_name = rng_pdf_font_name()

        # =====================================================
        # RIGHT-TO-LEFT PERSIAN TEXT
        # =====================================================

        def shape_text(value):

            value = str(
                value if value is not None else ""
            )

            reshaped = arabic_reshaper.reshape(value)

            return get_display(reshaped)

        def text_width(value, size):

            return rng_pdf_width(shape_text(value), size)

        # =====================================================
        # TEXT FITTING INSIDE PDF CELLS
        # =====================================================

        def draw_cell_text(
            value,
            x,
            y,
            width,
            height,
            font_size=9,
            align="right"
        ):

            text = str(
                value if value is not None else ""
            ).strip()

            if not text:
                return

            padding = 3

            available_width = max(
                width - 2 * padding,
                1
            )

            available_height = max(
                height - 2 * padding,
                1
            )

            def make_lines(size):

                lines = []

                for paragraph in text.splitlines():

                    if not paragraph.strip():
                        lines.append("")
                        continue

                    words = paragraph.split()

                    current = ""

                    for word in words:

                        candidate = (
                            current + " " + word
                            if current else word
                        )

                        if text_width(
                            candidate,
                            size
                        ) <= available_width:

                            current = candidate

                        else:

                            if current:
                                lines.append(current)

                            current = ""

                            for character in word:

                                candidate = current + character

                                if (
                                    current
                                    and text_width(
                                        candidate,
                                        size
                                    ) > available_width
                                ):
                                    lines.append(current)
                                    current = character
                                else:
                                    current = candidate

                    if current:
                        lines.append(current)

                return lines or [""]

            size = float(font_size)

            while size > 5:

                lines = make_lines(size)
                line_height = size * 1.3

                if (
                    len(lines) * line_height
                    <= available_height
                ):
                    break

                size -= 0.5

            lines = make_lines(size)
            line_height = size * 1.3

            max_lines = max(
                1,
                int(available_height / line_height)
            )

            if len(lines) > max_lines:

                lines = lines[:max_lines]

                final_line = lines[-1]

                while (
                    final_line
                    and text_width(
                        final_line + "...",
                        size
                    ) > available_width
                ):
                    final_line = final_line[:-1]

                lines[-1] = final_line + "..."

            pdf.saveState()

            clip = pdf.beginPath()

            clip.rect(
                x + 1,
                y + 1,
                max(width - 2, 0),
                max(height - 2, 0)
            )

            pdf.clipPath(
                clip,
                stroke=0,
                fill=0
            )

            pdf.setFillColorRGB(0, 0, 0)
            pdf.setFont(font_name, size)

            total_height = len(lines) * line_height

            baseline = (
                y
                + (height + total_height) / 2
                - line_height
                + 2
            )

            for line in lines:

                rendered = shape_text(line)

                if align == "center":

                    rng_pdf_draw(pdf, rendered, x + width / 2, baseline, size, "center")

                elif align == "left":

                    rng_pdf_draw(pdf, rendered, x + padding, baseline, size, "left")

                else:

                    rng_pdf_draw(pdf, rendered, x + width - padding, baseline, size, "right")

                baseline -= line_height

            pdf.restoreState()

        # =====================================================
        # A5 PAGE AND EXCEL DIMENSIONS
        # =====================================================

        margin_x = 6 * mm
        margin_y = 6 * mm

        left = margin_x
        right = page_width - margin_x

        top = page_height - margin_y
        bottom = margin_y

        table_width = right - left
        table_height = top - bottom

        # Same row proportions as the original Excel file.

        excel_row_heights = [
            60,
            42
        ] + [21.75] * 10 + [66]

        height_scale = (
            table_height / sum(excel_row_heights)
        )

        row_heights = [
            height * height_scale
            for height in excel_row_heights
        ]

        # Original A-G Excel widths.

        excel_col_widths = [
            5.85546875,
            17.42578125,
            27.0,
            10.140625,
            15.5703125,
            6.42578125,
            16.7109375
        ]

        width_scale = (
            table_width / sum(excel_col_widths)
        )

        col_widths = [
            width * width_scale
            for width in excel_col_widths
        ]

        # =====================================================
        # RIGHT-TO-LEFT COLUMN COORDINATES
        # =====================================================

        # Excel column A is at the RIGHT.
        # Excel column G is at the LEFT.
        #
        # x_positions[0] = right edge of column A
        # x_positions[7] = left edge of column G

        x_positions = [right]

        for width in col_widths:
            x_positions.append(
                x_positions[-1] - width
            )

        y_positions = [top]

        for height in row_heights:
            y_positions.append(
                y_positions[-1] - height
            )

        pdf.setStrokeColorRGB(0, 0, 0)
        pdf.setLineWidth(0.65)

        # =====================================================
        # DRAW RIGHT-TO-LEFT MERGED CELL
        # =====================================================

        def draw_cell(
            first_col,
            last_col,
            row_index,
            value="",
            font_size=9,
            align="center"
        ):

            # RTL: last column has the smaller X coordinate.

            x = x_positions[last_col + 1]

            width = (
                x_positions[first_col]
                - x_positions[last_col + 1]
            )

            y = y_positions[row_index + 1]
            height = row_heights[row_index]

            pdf.rect(
                x,
                y,
                width,
                height,
                fill=0,
                stroke=1
            )

            draw_cell_text(
                value,
                x,
                y,
                width,
                height,
                font_size=font_size,
                align=align
            )

        # =====================================================
        # ROW 1: LOGO, TITLE, DATE AND NUMBER
        # =====================================================

        # Original Excel:
        # A1:B1 = logo, now on the RIGHT
        # C1:E1 = title, CENTER
        # F1:G1 = date and number, LEFT

        draw_cell(0, 1, 0)

        draw_cell(
            2,
            4,
            0,
            template["title"],
            font_size=16,
            align="center"
        )

        draw_cell(5, 5, 0)
        draw_cell(6, 6, 0)

        # -----------------------------------------------------
        # ORIGINAL EXCEL LOGO AT UPPER-RIGHT
        # -----------------------------------------------------

        logo_reader = ImageReader(
            io.BytesIO(template["logo"])
        )

        logo_x = x_positions[2]
        logo_y = y_positions[1]

        logo_width = x_positions[0] - x_positions[2]
        logo_height = row_heights[0]

        image_width, image_height = logo_reader.getSize()

        logo_scale = min(
            logo_width * 0.94 / image_width,
            logo_height * 0.94 / image_height
        )

        rendered_width = image_width * logo_scale
        rendered_height = image_height * logo_scale

        pdf.drawImage(
            logo_reader,
            logo_x + (logo_width - rendered_width) / 2,
            logo_y + (logo_height - rendered_height) / 2,
            width=rendered_width,
            height=rendered_height,
            preserveAspectRatio=True,
            mask="auto"
        )

        # -----------------------------------------------------
        # DATE AND NUMBER ON LEFT SIDE
        # -----------------------------------------------------

        f_x = x_positions[6]
        f_width = col_widths[5]

        g_x = x_positions[7]
        g_width = col_widths[6]

        half_height = row_heights[0] / 2

        draw_cell_text(
            template["date_label"],
            f_x,
            logo_y + half_height,
            f_width,
            half_height,
            font_size=9,
            align="right"
        )

        draw_cell_text(
            template["number_label"],
            f_x,
            logo_y,
            f_width,
            half_height,
            font_size=9,
            align="right"
        )

        draw_cell_text(
            form["date"],
            g_x,
            logo_y + half_height,
            g_width,
            half_height,
            font_size=9,
            align="right"
        )

        draw_cell_text(
            form["number"],
            g_x,
            logo_y,
            g_width,
            half_height,
            font_size=9,
            align="right"
        )

        # =====================================================
        # ROW 2: TABLE HEADERS, RIGHT TO LEFT
        # =====================================================

        # RIGHT TO LEFT:
        # A = Item code
        # B:C = Description
        # D = Requested quantity
        # E = Usage
        # F:G = Notes

        draw_cell(
            0, 0, 1,
            template["headers"][0],
            font_size=9
        )

        draw_cell(
            1, 2, 1,
            template["headers"][1],
            font_size=10
        )

        draw_cell(
            3, 3, 1,
            template["headers"][2],
            font_size=9
        )

        draw_cell(
            4, 4, 1,
            template["headers"][3],
            font_size=9
        )

        draw_cell(
            5, 6, 1,
            template["headers"][4],
            font_size=10
        )

        # =====================================================
        # ROWS 3-12: TEN DATA ROWS, RIGHT TO LEFT
        # =====================================================

        merged_columns = [
            (0, 0),
            (1, 2),
            (3, 3),
            (4, 4),
            (5, 6)
        ]

        for item_index in range(10):

            excel_row_index = item_index + 2

            for field_index, col_range in enumerate(
                merged_columns
            ):

                draw_cell(
                    col_range[0],
                    col_range[1],
                    excel_row_index,
                    form["items"][item_index][field_index],
                    font_size=8.5,
                    align="right"
                )

        # =====================================================
        # ROW 13: SIGNATURES, RIGHT TO LEFT
        # =====================================================

        footer_row = 12

        # A13:C13 = Requester, RIGHT side
        # D13:G13 = Approver, LEFT side

        draw_cell(0, 2, footer_row)
        draw_cell(3, 6, footer_row)

        footer_y = y_positions[13]
        footer_height = row_heights[12]

        # Requester on RIGHT.

        requester_x = x_positions[3]
        requester_width = (
            x_positions[0] - x_positions[3]
        )

        # Approver on LEFT.

        approver_x = x_positions[7]
        approver_width = (
            x_positions[3] - x_positions[7]
        )

        draw_cell_text(
            template["requester_label"],
            requester_x,
            footer_y + footer_height * 0.70,
            requester_width,
            footer_height * 0.25,
            font_size=10,
            align="right"
        )

        draw_cell_text(
            form["requester"],
            requester_x,
            footer_y + footer_height * 0.43,
            requester_width,
            footer_height * 0.25,
            font_size=10,
            align="right"
        )

        draw_cell_text(
            template["signature_label"] + ": "
            + form["requester_signature"],
            requester_x,
            footer_y + footer_height * 0.05,
            requester_width,
            footer_height * 0.25,
            font_size=9,
            align="right"
        )

        draw_cell_text(
            template["approver_label"],
            approver_x,
            footer_y + footer_height * 0.70,
            approver_width,
            footer_height * 0.25,
            font_size=10,
            align="right"
        )

        draw_cell_text(
            form["approver"],
            approver_x,
            footer_y + footer_height * 0.43,
            approver_width,
            footer_height * 0.25,
            font_size=10,
            align="right"
        )

        draw_cell_text(
            template["signature_label"] + ": "
            + form["approver_signature"],
            approver_x,
            footer_y + footer_height * 0.05,
            approver_width,
            footer_height * 0.25,
            font_size=9,
            align="right"
        )

        # =====================================================
        # FINALIZE PDF
        # =====================================================

        pdf.showPage()
        pdf.save()

        buffer.seek(0)

        return buffer.getvalue()

    # =========================================================
    # WAREHOUSE DETAIL VIEW
    # =========================================================

    selected = st.session_state.selected_warehouse

    if (
        isinstance(selected, int)
        and 0 <= selected < len(st.session_state.warehouse_cards)
    ):

        warehouse_name = st.session_state.warehouse_cards[selected]
        warehouse_id = st.session_state.warehouse_card_ids[selected]

        if warehouse_id not in st.session_state.warehouse_forms:
            st.session_state.warehouse_forms[
                warehouse_id
            ] = create_empty_warehouse_form()

        form = st.session_state.warehouse_forms[warehouse_id]

        if st.button(
            "Back to Warehouse",
            key="warehouse_back"
        ):
            st.session_state.selected_warehouse = None
            rng_rerun()

        st.markdown(
            '<div class="rng-warehouse-hero"><div class="rng-warehouse-eyebrow">PURCHASE REQUEST WORKSPACE</div>'
            f'<h1>{escape(str(warehouse_name))}</h1>'
            '<p>From the first item to the final approval. Every detail in its place.</p></div>',
            unsafe_allow_html=True
        )

        if warehouse_template is None:

            st.error(
                "Unable to load original Excel template: "
                + str(template_error)
            )

        else:

            with st.container(key="warehouse_detail"):

                with st.container(key="rng_warehouse_meta"):
                    st.markdown('<div class="rng-warehouse-section">Request details</div>', unsafe_allow_html=True)
                    st.markdown('<div class="warehouse-form-title">' + escape(warehouse_template["title"]) + '</div>', unsafe_allow_html=True)
                    number_col, date_col = st.columns(2)
                    with number_col:
                        form["number"] = st.text_input(warehouse_template["number_label"], value=form["number"], key=f"warehouse_number_{warehouse_id}")
                    with date_col:
                        form["date"] = st.text_input(warehouse_template["date_label"], value=form["date"], key=f"warehouse_date_{warehouse_id}")

                st.markdown(
                    '<div class="rng-warehouse-section">Requested items</div>'
                    '<div class="rng-warehouse-help">All item details are visible below. Your PDF keeps the original 10-line request form.</div>',
                    unsafe_allow_html=True
                )
                headers = warehouse_template["headers"]
                with st.container(key="rng_warehouse_items"):
                    for row_index in range(10):
                        values = list(form["items"][row_index])
                        with st.container(key=f"rng_warehouse_item_{row_index}"):
                            st.markdown(f'<div class="rng-warehouse-item-label">ITEM {row_index + 1:02}</div>', unsafe_allow_html=True)
                            visual_order = [4, 3, 2, 1, 0]
                            for column, data_index in zip(st.columns([2, 2, 1, 3, 1]), visual_order):
                                with column:
                                    widget_key = f"warehouse_{warehouse_id}_cell_{row_index}_{data_index}"
                                    if data_index in (0, 2):
                                        values[data_index] = st.text_input(headers[data_index], value=values[data_index], key=widget_key)
                                    else:
                                        values[data_index] = st.text_area(headers[data_index], value=values[data_index], key=widget_key, height=90)
                        form["items"][row_index] = values

                with st.container(key="rng_warehouse_approval"):
                    st.markdown('<div class="rng-warehouse-section">Request &amp; approval</div><div class="rng-warehouse-help">Keep the requester and approval information with the request.</div>', unsafe_allow_html=True)
                    approver_col, requester_col = st.columns(2)
                    with approver_col:
                        form["approver"] = st.text_input(warehouse_template["approver_label"], value=form["approver"], key=f"warehouse_approver_{warehouse_id}")
                        form["approver_signature"] = st.text_input(warehouse_template["signature_label"], value=form["approver_signature"], key=f"warehouse_app_sign_{warehouse_id}")
                    with requester_col:
                        form["requester"] = st.text_input(warehouse_template["requester_label"], value=form["requester"], key=f"warehouse_requester_{warehouse_id}")
                        form["requester_signature"] = st.text_input(warehouse_template["signature_label"], value=form["requester_signature"], key=f"warehouse_req_sign_{warehouse_id}")

                with st.container(key="rng_warehouse_save_action"):
                    if st.button("Save request", key=f"warehouse_save_form_{warehouse_id}", use_container_width=True):
                        try:
                            rng_save()
                            st.success("Warehouse request saved.")
                        except Exception as exc:
                            st.error("Could not save your request. Please retry. " + str(exc))

                # =================================================
                # DOWNLOAD A5 LANDSCAPE PDF
                # =================================================

                with st.container(key="rng_warehouse_export"):
                    st.markdown('<div class="rng-warehouse-section">Ready to send</div><div class="rng-warehouse-help">Download or share the original A5 landscape request, with its logo and Persian labels.</div>', unsafe_allow_html=True)

                    try:

                        pdf_bytes = create_warehouse_pdf(
                            form,
                            warehouse_template
                        )

                        safe_filename = re.sub(
                            r"[^\w.-]+",
                            "_",
                            warehouse_name
                        )

                        warehouse_pdf_filename = f"Purchase_Request_{safe_filename}.pdf"
                        download_col, share_col = st.columns(2)
                        with download_col:
                            st.download_button(
                                label="Download A5 Landscape PDF",
                                data=pdf_bytes,
                                file_name=warehouse_pdf_filename,
                                mime="application/pdf",
                                key=f"warehouse_pdf_{warehouse_id}",
                                use_container_width=True
                            )
                        with share_col:
                            rng_pdf_share_button(pdf_bytes, warehouse_pdf_filename)

                    except Exception as error:

                        st.error(
                            "Unable to generate the PDF: "
                            + str(error)
                        )

    # =========================================================
    # WAREHOUSE MAIN PAGE
    # =========================================================

    else:

        st.session_state.selected_warehouse = None

        warehouse_total = len(st.session_state.warehouse_cards)
        st.markdown(
            '<div class="rng-warehouse-hero"><div class="rng-warehouse-eyebrow">ORGANIZED REQUESTS. CLEAR NEXT STEPS.</div>'
            '<h1>Warehouse</h1><p>Keep every request, item, and approval in one place.</p>'
            f'<span class="rng-warehouse-total">{warehouse_total} {"warehouse" if warehouse_total == 1 else "warehouses"} in your workspace</span></div>',
            unsafe_allow_html=True
        )

        # =====================================================
        # ADD NEW WAREHOUSE
        # =====================================================

        if st.session_state.adding_warehouse:
            with st.container(key="rng_warehouse_add_form"):

                st.markdown(
                    '<div class="section-title">Add New Warehouse</div>',
                    unsafe_allow_html=True
                )

                new_title = st.text_input(
                    "Warehouse title",
                    placeholder="Enter warehouse title",
                    key="new_warehouse_title"
                )

                add_col, cancel_col = st.columns(2)

                with add_col:

                    if st.button(
                        "Add Warehouse",
                        key="warehouse_add_confirm",
                        use_container_width=True
                    ):

                        if new_title.strip():

                            st.session_state.warehouse_cards.append(
                                new_title.strip()
                            )

                            st.session_state.warehouse_card_ids.append(
                                uuid4().hex
                            )

                            st.session_state.adding_warehouse = False

                            rng_rerun()

                        else:

                            st.warning(
                                "Please enter a warehouse title."
                            )

                with cancel_col:

                    if st.button(
                        "Cancel",
                        key="warehouse_add_cancel",
                        use_container_width=True
                    ):

                        st.session_state.adding_warehouse = False

                        rng_rerun()

        else:
            with st.container(key="rng_warehouse_add_action"):

                if st.button(
                    "+ Add Warehouse",
                    key="warehouse_add_button"
                ):

                    st.session_state.adding_warehouse = True

                    rng_rerun()

        st.markdown(
            "<br>",
            unsafe_allow_html=True
        )

        # =====================================================
        # WAREHOUSE CARDS
        # =====================================================

        if not st.session_state.warehouse_cards:
            st.markdown('<div class="rng-warehouse-empty"><strong>Everything starts with a clear request.</strong>Add a warehouse to organize purchase requests and their approvals.</div>', unsafe_allow_html=True)
        with st.container(key="rng_warehouse_board"):
            cols = st.columns(3)

            for i, warehouse in enumerate(
                st.session_state.warehouse_cards
            ):

                with cols[i % 3], st.container(key=f"rng_warehouse_card_{i}"):

                    warehouse_card_heading(warehouse, i)

                    safe_title = escape(warehouse)

                    # =============================================
                    # DELETE CONFIRMATION
                    # =============================================

                    if st.session_state.deleting_warehouse == i:

                        st.markdown(
                            f'<div class="warehouse-delete-warning">Delete "{safe_title}"? This cannot be undone.</div>',
                            unsafe_allow_html=True
                        )

                        confirm_col, cancel_col = st.columns(2)

                        with confirm_col:

                            if st.button(
                                "Confirm Delete",
                                key=f"warehouse_confirm_delete_{i}",
                                use_container_width=True
                            ):

                                deleted_id = (
                                    st.session_state.warehouse_card_ids.pop(i)
                                )

                                st.session_state.warehouse_cards.pop(i)

                                st.session_state.warehouse_forms.pop(
                                    deleted_id,
                                    None
                                )

                                st.session_state.deleting_warehouse = None
                                st.session_state.editing_warehouse = None

                                rng_rerun()

                        with cancel_col:

                            if st.button(
                                "Cancel",
                                key=f"warehouse_cancel_delete_{i}",
                                use_container_width=True
                            ):

                                st.session_state.deleting_warehouse = None

                                rng_rerun()

                    # =============================================
                    # EDIT WAREHOUSE TITLE
                    # =============================================

                    elif st.session_state.editing_warehouse == i:

                        edited_title = st.text_input(
                            "Warehouse title",
                            value=warehouse,
                            key=f"warehouse_edit_input_{i}"
                        )

                        save_col, cancel_col = st.columns(2)

                        with save_col:

                            if st.button(
                                "Save",
                                key=f"warehouse_save_{i}",
                                use_container_width=True
                            ):

                                if edited_title.strip():

                                    st.session_state.warehouse_cards[i] = (
                                        edited_title.strip()
                                    )

                                    st.session_state.editing_warehouse = None

                                    rng_rerun()

                                else:

                                    st.warning(
                                        "Please enter a warehouse title."
                                    )

                        with cancel_col:

                            if st.button(
                                "Cancel",
                                key=f"warehouse_edit_cancel_{i}",
                                use_container_width=True
                            ):

                                st.session_state.editing_warehouse = None

                                rng_rerun()

                    # =============================================
                    # NORMAL WAREHOUSE CARD
                    # =============================================

                    else:

                        button_col1, button_col2, button_col3 = st.columns(3)

                        with button_col1, st.container(key=f"rng_warehouse_open_{i}"):

                            open_button = st.button(
                                "Open",
                                key=f"warehouse_open_{i}",
                                use_container_width=True
                            )

                        with button_col2:

                            edit_button = st.button(
                                "Edit",
                                key=f"warehouse_edit_{i}",
                                use_container_width=True
                            )

                        with button_col3:

                            delete_button = st.button(
                                "Delete",
                                key=f"warehouse_delete_{i}",
                                use_container_width=True
                            )

                        if open_button:

                            st.session_state.selected_warehouse = i

                            rng_rerun()

                        if edit_button:

                            st.session_state.editing_warehouse = i

                            st.session_state.deleting_warehouse = None

                            rng_rerun()

                        if delete_button:

                            st.session_state.deleting_warehouse = i

                            st.session_state.editing_warehouse = None

                            rng_rerun()







elif page == "Mission Report":

    import io
    import re
    import hashlib
    import zipfile

    from pathlib import Path
    from uuid import uuid4
    from html import escape

    # =========================================================
    # ORIGINAL EXCEL TEMPLATE
    # =========================================================

    MISSION_TEMPLATE_PATH = (
        Path(__file__).resolve().parent
        / "\u0641\u0631\u0645 \u06af\u0632\u0627\u0631\u0634 \u0645\u0627\u0645\u0648\u0631\u06cc\u062a.xlsx"
    )

    def load_mission_template():

        from openpyxl import load_workbook

        if not MISSION_TEMPLATE_PATH.is_file():
            raise FileNotFoundError(
                "Excel template not found: "
                + str(MISSION_TEMPLATE_PATH)
            )

        wb = load_workbook(
            MISSION_TEMPLATE_PATH,
            read_only=False,
            data_only=True
        )

        if "RAW" not in wb.sheetnames:
            wb.close()
            raise ValueError("RAW worksheet not found.")

        ws = wb["RAW"]

        cells = {}

        for row in range(1, 14):
            for col in range(1, 9):

                cell = ws.cell(row, col)

                if cell.value is not None:
                    cells[cell.coordinate] = str(cell.value)

        column_widths = [
            ws.column_dimensions[letter].width or 10
            for letter in "ABCDEFGH"
        ]

        row_heights = [
            ws.row_dimensions[row].height or 25
            for row in range(1, 14)
        ]

        wb.close()

        logo = None

        with zipfile.ZipFile(MISSION_TEMPLATE_PATH) as archive:

            images = [
                name for name in archive.namelist()
                if name.startswith("xl/media/")
                and name.lower().endswith(
                    (".png", ".jpg", ".jpeg")
                )
            ]

            if images:
                logo = archive.read(images[0])

        return {
            "cells": cells,
            "column_widths": column_widths,
            "row_heights": row_heights,
            "logo": logo
        }

    mission_template = None
    mission_template_error = None

    try:
        mission_template = load_mission_template()

    except Exception as error:
        mission_template_error = str(error)

    # =========================================================
    # SESSION STATE
    # =========================================================

    if "mission_report_cards" not in st.session_state:
        st.session_state.mission_report_cards = []

    if "selected_mission_report" not in st.session_state:
        st.session_state.selected_mission_report = None

    if "editing_mission_report" not in st.session_state:
        st.session_state.editing_mission_report = None

    if "deleting_mission_report" not in st.session_state:
        st.session_state.deleting_mission_report = None

    if "adding_mission_report" not in st.session_state:
        st.session_state.adding_mission_report = False

    if "mission_report_card_ids" not in st.session_state:
        st.session_state.mission_report_card_ids = [
            uuid4().hex
            for _ in st.session_state.mission_report_cards
        ]

    if "mission_report_forms" not in st.session_state:
        st.session_state.mission_report_forms = {}

    while len(st.session_state.mission_report_card_ids) < len(
        st.session_state.mission_report_cards
    ):
        st.session_state.mission_report_card_ids.append(
            uuid4().hex
        )

    # =========================================================
    # UNICODE-SAFE PERSIAN LABELS
    # =========================================================

    FIELD_LABELS = {
        "report_date": "\u062a\u0627\u0631\u06cc\u062e \u06af\u0632\u0627\u0631\u0634",
        "employee": "\u0646\u0627\u0645 \u0648 \u0646\u0627\u0645 \u062e\u0627\u0646\u0648\u0627\u062f\u06af\u06cc",
        "position": "\u067e\u0633\u062a \u0633\u0627\u0632\u0645\u0627\u0646\u06cc",
        "dispatcher": "\u0627\u0639\u0632\u0627\u0645 \u06a9\u0646\u0646\u062f\u0647",
        "city": "\u0634\u0647\u0631 \u0645\u0627\u0645\u0648\u0631\u06cc\u062a",
        "company": "\u0646\u0627\u0645 \u0634\u0631\u06a9\u062a / \u06a9\u0627\u0631\u062e\u0627\u0646\u0647",
        "companions": "\u0647\u0645\u0631\u0627\u0647\u0627\u0646",
        "duration": "\u0645\u062f\u062a \u0645\u0627\u0645\u0648\u0631\u06cc\u062a",
        "start_date": "\u062a\u0627\u0631\u06cc\u062e \u0634\u0631\u0648\u0639",
        "departure_time": "\u0633\u0627\u0639\u062a \u0631\u0641\u062a",
        "end_date": "\u062a\u0627\u0631\u06cc\u062e \u067e\u0627\u06cc\u0627\u0646",
        "return_time": "\u0633\u0627\u0639\u062a \u0628\u0631\u06af\u0634\u062a",
        "transport": "\u0627\u0633\u062a\u0641\u0627\u062f\u0647 \u0627\u0632 \u0648\u0633\u0627\u06cc\u0644 \u0646\u0642\u0644\u06cc\u0647",
        "distance": "\u0645\u0633\u0627\u0641\u062a \u0637\u06cc \u0634\u062f\u0647 (\u06a9\u06cc\u0644\u0648\u0645\u062a\u0631)",
        "equipment": "\u062a\u062c\u0647\u06cc\u0632\u0627\u062a \u0645\u0648\u0631\u062f \u0627\u0633\u062a\u0641\u0627\u062f\u0647",
        "description": "\u0634\u0631\u062d \u0645\u0627\u0645\u0648\u0631\u06cc\u062a",
        "objective": "\u0647\u062f\u0641 \u0645\u0627\u0645\u0648\u0631\u06cc\u062a",
        "result": "\u0646\u062a\u06cc\u062c\u0647 \u0645\u0627\u0645\u0648\u0631\u06cc\u062a",
        "manager_comment": "\u0646\u0638\u0631 \u0645\u062f\u06cc\u0631 / \u0631\u0626\u06cc\u0633 \u0645\u0631\u0628\u0648\u0637\u0647",
        "manager_signature": "\u0627\u0645\u0636\u0627\u06cc \u0645\u062f\u06cc\u0631",
        "employee_signature": "\u0627\u0645\u0636\u0627\u06cc \u0645\u0627\u0645\u0648\u0631"
    }

    # =========================================================
    # UNICODE-SAFE SECTION HEADINGS
    # =========================================================

    MISSION_HEADINGS = {
        "title": (
            "\u0641\u0631\u0645 \u06af\u0632\u0627\u0631\u0634 "
            "\u0645\u0627\u0645\u0648\u0631\u06cc\u062a"
        ),
        "schedule": (
            "\u0632\u0645\u0627\u0646 \u0628\u0646\u062f\u06cc "
            "\u0645\u0627\u0645\u0648\u0631\u06cc\u062a"
        ),
        "transport": (
            "\u0627\u0637\u0644\u0627\u0639\u0627\u062a "
            "\u062d\u0645\u0644 \u0648 \u0646\u0642\u0644"
        ),
        "description": (
            "\u0634\u0631\u062d \u0627\u062c\u0631\u0627\u06cc "
            "\u0645\u0627\u0645\u0648\u0631\u06cc\u062a "
            "\u0648 \u0647\u062f\u0641 \u0622\u0646"
        ),
        "manager": (
            "\u0627\u06cc\u0646 \u0642\u0633\u0645\u062a "
            "\u062a\u0648\u0633\u0637 \u0645\u062f\u06cc\u0631 "
            "\u0645\u0631\u0628\u0648\u0637\u0647 "
            "\u062a\u06a9\u0645\u06cc\u0644 \u06af\u0631\u062f\u062f."
        )
    }

    # =========================================================
    # FORM DATA
    # =========================================================

    def create_empty_mission_form():

        return {
            field: ""
            for field in FIELD_LABELS
        }

    # =========================================================
    # CSS
    # =========================================================

    st.markdown('<div id="rng-mission-marker"></div>', unsafe_allow_html=True)
    st.markdown("""
    <style>
    [data-testid="stMain"]:has(#rng-mission-marker) { background: radial-gradient(ellipse at right top, #dce8f9, transparent 55%), #f2f5fa; }
    .stApp:has(#rng-mission-marker), .stApp:has(#rng-mission-marker) h1,
    .stApp:has(#rng-mission-marker) h2, .stApp:has(#rng-mission-marker) h3,
    .stApp:has(#rng-mission-marker) [data-testid="stMarkdownContainer"],
    .stApp:has(#rng-mission-marker) input, .stApp:has(#rng-mission-marker) textarea,
    .stApp:has(#rng-mission-marker) label { font-family: Arial, Tahoma, sans-serif !important; font-style: normal !important; }
    .stApp:has(#rng-mission-marker) .block-container { max-width: 1440px; }
    .rng-mission-hero { background: radial-gradient(ellipse at right top, #355e89, transparent 65%), #172f4e; border: 1px solid #315174; border-radius: 24px; padding: 30px 34px; margin: 8px 0 14px; }
    .rng-mission-eyebrow { color: #b1cae9; font-size: 11px; font-weight: 700; letter-spacing: 2px; margin-bottom: 14px; }
    .rng-mission-hero h1 { color: #f6f9ff !important; font-size: clamp(30px, 3.4vw, 44px) !important; font-weight: 750; letter-spacing: -1.2px; margin: 0 0 12px; padding: 0; overflow-wrap: anywhere; }
    .rng-mission-hero p { color: #c6d8ed; font-size: 15px; line-height: 1.8; margin: 0; }
    .rng-mission-total { display: inline-block; color: #dce9fb; background: #a9ccff12; border: 1px solid #a9ccff30; border-radius: 100px; padding: 6px 12px; font-size: 11px; margin-top: 18px; }
    .st-key-rng_mission_add_action .stButton button, .st-key-rng_mission_save_action .stButton button { color: #fff; background: #254d7e; border: 1px solid #254d7e; min-height: 46px; border-radius: 11px; }
    .st-key-rng_mission_add_action .stButton button:hover, .st-key-rng_mission_save_action .stButton button:hover { background: #18395f; color: #fff; }
    .st-key-rng_mission_add_form { background: #fff; border: 1px solid #d7e3f2; padding: 24px; border-radius: 18px; margin-bottom: 20px; }
    [class*="st-key-rng_mission_card_"] { background: #fff; border: 1px solid #dce5f1; border-radius: 18px; padding: 24px; min-height: 270px; box-shadow: 0 8px 25px #15385f07; }
    .rng-mission-card-top { display: flex; align-items: center; justify-content: space-between; margin-bottom: 19px; }
    .rng-mission-icon { display: flex; align-items: center; justify-content: center; width: 46px; height: 46px; color: #5279aa; background: #eaf1fc; border-radius: 14px; }
    .rng-mission-icon svg { width: 25px; height: 25px; }
    .rng-mission-index { color: #8698b0; font-size: 11px; letter-spacing: 1px; }
    .mission-report-card-title { color: #243e60; font-size: 19px; font-weight: 700; line-height: 1.5; margin: 0 0 10px; overflow-wrap: anywhere; }
    .mission-report-card-text { color: #6f839e; font-size: 12px; line-height: 1.8; margin-bottom: 13px; overflow-wrap: anywhere; }
    .rng-mission-card-footer { color: #8092aa; font-size: 10px; letter-spacing: .6px; border-top: 1px solid #e8edf5; padding-top: 14px; margin-bottom: 6px; }
    [class*="st-key-rng_mission_card_"] [data-testid="stHorizontalBlock"] > [data-testid="stColumn"] { min-width: 0 !important; flex: 1; }
    .st-key-rng_mission_board .stButton button { color: #48658b; background: #f5f8fd; border: 1px solid #dfe7f3; border-radius: 9px; min-height: 44px; }
    .st-key-rng_mission_board .stButton button:hover { background: #eaf1fc; color: #234b7b; border-color: #8faed7; }
    [class*="st-key-rng_mission_open_"] .stButton button { background: #244d7c; border-color: #244d7c; color: #fff; }
    [class*="st-key-rng_mission_open_"] .stButton button:hover { background: #17395f; color: #fff; }
    .mission-report-delete-warning { color: #8c4050; background: #fff2f4; border: 1px solid #f1d4dc; border-radius: 12px; padding: 15px; font-size: 13px; line-height: 1.8; overflow-wrap: anywhere; }
    .rng-mission-empty { text-align: center; padding: 50px 24px; border: 1px dashed #bdcee5; border-radius: 18px; background: #fff9; color: #70859f; }
    .rng-mission-empty strong { display: block; color: #304e75; font-size: 21px; margin-bottom: 10px; }
    .rng-mission-section { color: #294569; font-size: 21px; font-weight: 700; margin: 4px 0 8px; }
    .rng-mission-help { color: #7388a4; font-size: 13px; line-height: 1.8; margin-bottom: 16px; }
    .rng-mission-step { display: inline-flex; align-items: center; justify-content: center; width: 28px; height: 28px; margin-right: 10px; border-radius: 8px; color: #547bad; background: #e8f0fc; font-size: 11px; font-weight: 700; }
    .rng-mission-card-date { color: #6d85a6; font-size: 11px; margin: 8px 0 14px; }
    .mission-form-title { direction: rtl; color: #48688e; font-size: 18px; font-weight: 600; text-align: right; margin: 4px 0 18px; }
    [class*="st-key-rng_mission_panel_"] { background: #fff; border: 1px solid #dce5f2; padding: 24px; border-radius: 18px; margin: 6px 0 12px; }
    .st-key-mission_detail input, .st-key-mission_detail textarea { direction: rtl; text-align: start; unicode-bidi: plaintext; background: #f9fbff; color: #2f486b; min-height: 44px; }
    .st-key-mission_detail [data-testid="stWidgetLabel"] { direction: rtl; text-align: right; }
    .st-key-mission_detail label p { color: #587193; font-size: 12px; }
    .st-key-rng_mission_export { background: #e8eff9; border: 1px solid #d9e4f3; padding: 22px; border-radius: 18px; margin-top: 20px; }
    .st-key-rng_mission_export .stDownloadButton button { background: #fff; color: #345781; border: 1px solid #c5d6ed; min-height: 46px; border-radius: 10px; }
    .stApp:has(#rng-mission-marker) button:focus-visible { outline: 3px solid #9ab9e8; outline-offset: 2px; }
    @media (max-width: 1100px) and (min-width: 701px) {
        .st-key-rng_mission_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_mission_card_"]) { flex-wrap: wrap; }
        .st-key-rng_mission_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_mission_card_"]) > [data-testid="stColumn"] { min-width: calc(50% - 1rem); }
    }
    @media (max-width: 700px) {
        .rng-mission-hero { padding: 26px 22px; border-radius: 20px; }
        .st-key-rng_mission_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_mission_card_"]) { flex-wrap: wrap; }
        .st-key-rng_mission_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_mission_card_"]) > [data-testid="stColumn"] { min-width: 100%; }
        [class*="st-key-rng_mission_card_"] { padding: 20px; }
        [class*="st-key-rng_mission_panel_"], .st-key-rng_mission_export { padding: 18px; }
    }
    </style>
    """, unsafe_allow_html=True)

    def mission_card_heading(name, index):
        record_id = st.session_state.mission_report_card_ids[index]
        data = st.session_state.mission_report_forms.get(record_id, {})
        location = " / ".join(str(data.get(field, "")).strip() for field in ("city", "company") if str(data.get(field, "")).strip())
        description = location or "Ready for your next assignment."
        date = str(data.get("report_date", "")).strip()
        st.markdown(
            '<div class="rng-mission-card-top"><div class="rng-mission-icon">'
            '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.5" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><path d="M19 10c0 5-7 11-7 11S5 15 5 10a7 7 0 1 1 14 0Z"/><circle cx="12" cy="10" r="2.5"/></svg>'
            f'</div><span class="rng-mission-index">REPORT {index + 1:02}</span></div>'
            f'<div class="mission-report-card-title">{escape(str(name))}</div>'
            f'<div class="mission-report-card-text">{escape(description)}</div>'
            + f'<div class="rng-mission-card-date">{escape(date or "Date not set")}</div>'
            + '<div class="rng-mission-card-footer">ASSIGNMENT / TRAVEL / OUTCOME</div>',
            unsafe_allow_html=True
        )

    # =========================================================
    # PDF GENERATOR
    # =========================================================

    def create_mission_pdf(form, template):

        import arabic_reshaper

        from bidi.algorithm import get_display

        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.lib.utils import ImageReader

        buffer = io.BytesIO()

        pdf = canvas.Canvas(
            buffer,
            pagesize=A4,
            pageCompression=1
        )

        page_width, page_height = A4

        pdf.setTitle(
            template["cells"].get(
                "C1",
                MISSION_HEADINGS["title"]
            )
        )

        # =====================================================
        # PERSIAN FONT SELECTION
        # =====================================================

        font_name = rng_pdf_font_name()

        def shape_text(value):
            value = str(value if value is not None else "")
            return get_display(arabic_reshaper.reshape(value))

        # =====================================================
        # PDF DIMENSIONS - A4 PORTRAIT
        # =====================================================

        margin_x = 7 * mm
        margin_y = 7 * mm

        left = margin_x
        right = page_width - margin_x

        top = page_height - margin_y
        bottom = margin_y

        usable_width = right - left
        usable_height = top - bottom

        col_source = template["column_widths"]
        row_source = template["row_heights"]

        col_total = sum(col_source)
        row_total = sum(row_source)

        col_widths = [
            usable_width * width / col_total
            for width in col_source
        ]

        row_heights = [
            usable_height * height / row_total
            for height in row_source
        ]

        # -----------------------------------------------------
        # RIGHT-TO-LEFT COORDINATES
        # -----------------------------------------------------

        x_edges = [right]

        for width in col_widths:
            x_edges.append(
                x_edges[-1] - width
            )

        y_edges = [top]

        for height in row_heights:
            y_edges.append(
                y_edges[-1] - height
            )

        # =====================================================
        # PDF TEXT HANDLING
        # =====================================================

        def text_width(value, size):

            return rng_pdf_width(shape_text(value), size)

        def wrap_text(value, max_width, size):

            lines = []

            for paragraph in str(value).splitlines():

                if not paragraph.strip():
                    lines.append("")
                    continue

                current = ""

                for word in paragraph.split():

                    candidate = (
                        current + " " + word
                        if current else word
                    )

                    if text_width(
                        candidate,
                        size
                    ) <= max_width:

                        current = candidate

                    else:

                        if current:
                            lines.append(current)

                        current = ""

                        for character in word:

                            candidate = current + character

                            if (
                                current
                                and text_width(
                                    candidate,
                                    size
                                ) > max_width
                            ):

                                lines.append(current)
                                current = character

                            else:
                                current = candidate

                if current:
                    lines.append(current)

            return lines or [""]

        # =====================================================
        # DRAW PDF TEXT
        # =====================================================

        def draw_text(
            value,
            x,
            y,
            width,
            height,
            size=10,
            align="right"
        ):

            text = str(
                value if value is not None else ""
            ).strip()

            if not text:
                return

            padding = 3

            max_width = max(
                width - 2 * padding,
                1
            )

            max_height = max(
                height - 2 * padding,
                1
            )

            font_size = float(size)

            while font_size > 5:

                lines = wrap_text(
                    text,
                    max_width,
                    font_size
                )

                line_height = font_size * 1.3

                if (
                    len(lines) * line_height
                    <= max_height
                ):
                    break

                font_size -= 0.5

            lines = wrap_text(
                text,
                max_width,
                font_size
            )

            line_height = font_size * 1.3

            max_lines = max(
                1,
                int(max_height / line_height)
            )

            if len(lines) > max_lines:

                lines = lines[:max_lines]

                final_line = lines[-1]

                while (
                    final_line
                    and text_width(
                        final_line + "...",
                        font_size
                    ) > max_width
                ):
                    final_line = final_line[:-1]

                lines[-1] = final_line + "..."

            # -------------------------------------------------
            # PREVENT TEXT OVERFLOW
            # -------------------------------------------------

            pdf.saveState()

            clipping = pdf.beginPath()

            clipping.rect(
                x + 1,
                y + 1,
                max(width - 2, 0),
                max(height - 2, 0)
            )

            pdf.clipPath(
                clipping,
                stroke=0,
                fill=0
            )

            pdf.setFont(
                font_name,
                font_size
            )

            pdf.setFillColorRGB(
                0.08,
                0.12,
                0.18
            )

            total_height = len(lines) * line_height

            baseline = (
                y
                + (height + total_height) / 2
                - line_height
                + 2
            )

            for line in lines:

                rendered = shape_text(line)

                if align == "center":

                    rng_pdf_draw(pdf, rendered, x + width / 2, baseline, font_size, "center")

                else:

                    rng_pdf_draw(pdf, rendered, x + width - padding, baseline, font_size, "right")

                baseline -= line_height

            pdf.restoreState()

        # =====================================================
        # DRAW PDF CELL
        # =====================================================

        def draw_cell(
            first_col,
            last_col,
            row,
            text="",
            font_size=10,
            align="right",
            shaded=False
        ):

            x = x_edges[last_col + 1]

            width = (
                x_edges[first_col]
                - x_edges[last_col + 1]
            )

            y = y_edges[row]

            height = row_heights[row - 1]

            if shaded:

                pdf.setFillColorRGB(
                    0.92,
                    0.93,
                    0.95
                )

                pdf.rect(
                    x,
                    y,
                    width,
                    height,
                    fill=1,
                    stroke=0
                )

            pdf.setStrokeColorRGB(
                0.3,
                0.3,
                0.3
            )

            pdf.setLineWidth(0.55)

            pdf.rect(
                x,
                y,
                width,
                height,
                fill=0,
                stroke=1
            )

            draw_text(
                text,
                x,
                y,
                width,
                height,
                size=font_size,
                align=align
            )

        # =====================================================
        # FIELD TEXT HELPER
        # =====================================================

        def field_text(field):

            return (
                FIELD_LABELS[field]
                + ": "
                + str(form.get(field, ""))
            )

        # =====================================================
        # ROW 1 - LOGO / TITLE / DATE
        # =====================================================

        draw_cell(0, 1, 1)

        draw_cell(
            2,
            5,
            1,
            template["cells"].get(
                "C1",
                MISSION_HEADINGS["title"]
            ),
            font_size=16,
            align="center"
        )

        draw_cell(
            6,
            7,
            1,
            field_text("report_date")
        )

        # -----------------------------------------------------
        # ORIGINAL EXCEL LOGO
        # -----------------------------------------------------

        if template["logo"]:

            image = ImageReader(
                io.BytesIO(template["logo"])
            )

            image_width, image_height = image.getSize()

            logo_x = x_edges[2]
            logo_y = y_edges[1]

            logo_width = x_edges[0] - x_edges[2]
            logo_height = row_heights[0]

            scale = min(
                logo_width * 0.92 / image_width,
                logo_height * 0.92 / image_height
            )

            render_width = image_width * scale
            render_height = image_height * scale

            pdf.drawImage(
                image,
                logo_x + (logo_width - render_width) / 2,
                logo_y + (logo_height - render_height) / 2,
                width=render_width,
                height=render_height,
                preserveAspectRatio=True,
                mask="auto"
            )

        # =====================================================
        # ROW 2 - EMPLOYEE AND POSITION
        # =====================================================

        draw_cell(
            0, 3, 2,
            field_text("employee")
        )

        draw_cell(
            4, 7, 2,
            field_text("position")
        )

        # =====================================================
        # ROW 3 - MISSION DESTINATION
        # =====================================================

        draw_cell(
            0, 2, 3,
            field_text("dispatcher")
        )

        draw_cell(
            3, 4, 3,
            field_text("city")
        )

        draw_cell(
            5, 7, 3,
            field_text("company")
        )

        # =====================================================
        # ROW 4 - COMPANIONS
        # =====================================================

        draw_cell(
            0, 7, 4,
            field_text("companions")
        )

        # =====================================================
        # ROW 5 - DATE AND TIME
        # =====================================================

        draw_cell(
            0, 0, 5,
            field_text("duration"),
            font_size=8
        )

        draw_cell(
            1, 2, 5,
            field_text("start_date"),
            font_size=9
        )

        draw_cell(
            3, 3, 5,
            field_text("departure_time"),
            font_size=8
        )

        draw_cell(
            4, 5, 5,
            field_text("end_date"),
            font_size=9
        )

        draw_cell(
            6, 7, 5,
            field_text("return_time"),
            font_size=9
        )

        # =====================================================
        # ROW 6 - TRANSPORTATION
        # =====================================================

        draw_cell(
            0, 7, 6,
            field_text("transport")
            + "       "
            + field_text("distance")
        )

        # =====================================================
        # ROW 7 - DESCRIPTION HEADING
        # =====================================================

        draw_cell(
            0, 7, 7,
            template["cells"].get(
                "A7",
                MISSION_HEADINGS["description"]
            ),
            font_size=11,
            shaded=True
        )

        # =====================================================
        # ROW 8 - EQUIPMENT
        # =====================================================

        draw_cell(
            0, 7, 8,
            field_text("equipment")
        )

        # =====================================================
        # ROW 9 - MISSION DESCRIPTION
        # =====================================================

        draw_cell(
            0, 7, 9,
            FIELD_LABELS["description"]
            + ":\n"
            + str(form.get("description", "")),
            font_size=10
        )

        # =====================================================
        # ROW 10 - OBJECTIVE
        # =====================================================

        draw_cell(
            0, 7, 10,
            field_text("objective")
        )

        # =====================================================
        # ROW 11 - RESULT
        # =====================================================

        draw_cell(
            0, 7, 11,
            field_text("result")
        )

        # =====================================================
        # ROW 12 - MANAGER HEADING
        # =====================================================

        draw_cell(
            0, 7, 12,
            template["cells"].get(
                "A12",
                MISSION_HEADINGS["manager"]
            ),
            font_size=10,
            align="center",
            shaded=True
        )

        # =====================================================
        # ROW 13 - MANAGER REVIEW / SIGNATURES
        # =====================================================

        draw_cell(
            0, 3, 13,
            FIELD_LABELS["manager_comment"]
            + ":\n"
            + str(form.get("manager_comment", ""))
            + "\n"
            + field_text("manager_signature")
        )

        draw_cell(
            4, 7, 13,
            field_text("employee_signature")
        )

        # =====================================================
        # COMPLETE SINGLE-PAGE PDF
        # =====================================================

        pdf.showPage()
        pdf.save()

        buffer.seek(0)

        return buffer.getvalue()

    # =========================================================
    # MISSION REPORT DETAIL VIEW
    # =========================================================

    selected = st.session_state.selected_mission_report

    if (
        isinstance(selected, int)
        and 0 <= selected < len(
            st.session_state.mission_report_cards
        )
    ):

        mission_name = st.session_state.mission_report_cards[
            selected
        ]

        mission_id = st.session_state.mission_report_card_ids[
            selected
        ]

        if mission_id not in st.session_state.mission_report_forms:

            st.session_state.mission_report_forms[
                mission_id
            ] = create_empty_mission_form()

        form = st.session_state.mission_report_forms[
            mission_id
        ]

        if st.button(
            "Back to Mission Reports",
            key="mission_report_back"
        ):

            st.session_state.selected_mission_report = None
            rng_rerun()

        st.markdown(
            '<div class="rng-mission-hero"><div class="rng-mission-eyebrow">MISSION REPORT WORKSPACE</div>'
            f'<h1>{escape(str(mission_name))}</h1>'
            '<p>A clear record of the journey, the work, and the outcome.</p></div>',
            unsafe_allow_html=True
        )

        if mission_template is None:

            st.error(
                "Unable to load the Excel template: "
                + str(mission_template_error)
            )

        else:

            with st.container(key="mission_detail"):

                def mission_input(field):
                    form[field] = st.text_input(
                        label=FIELD_LABELS[field], value=form.get(field, ""),
                        placeholder=FIELD_LABELS[field], key=f"mission_{mission_id}_{field}",
                        label_visibility="visible"
                    )

                def mission_text_area(field, height):
                    form[field] = st.text_area(
                        label=FIELD_LABELS[field], value=form.get(field, ""),
                        placeholder=FIELD_LABELS[field], height=height,
                        key=f"mission_{mission_id}_{field}", label_visibility="visible"
                    )

                def mission_section(number, title):
                    st.markdown(
                        f'<div class="rng-mission-section"><span class="rng-mission-step">{number:02}</span>{escape(title)}</div>',
                        unsafe_allow_html=True
                    )

                with st.container(key="rng_mission_panel_employee"):
                    mission_section(1, "Report & employee")
                    form_title = mission_template["cells"].get("C1", MISSION_HEADINGS["title"])
                    st.markdown('<div class="mission-form-title">' + escape(form_title) + '</div>', unsafe_allow_html=True)
                    mission_input("report_date")
                    position_col, employee_col = st.columns(2)
                    with position_col:
                        mission_input("position")
                    with employee_col:
                        mission_input("employee")

                with st.container(key="rng_mission_panel_destination"):
                    mission_section(2, "Destination & companions")
                    company_col, city_col, dispatcher_col = st.columns(3)
                    with company_col:
                        mission_input("company")
                    with city_col:
                        mission_input("city")
                    with dispatcher_col:
                        mission_input("dispatcher")
                    mission_input("companions")

                with st.container(key="rng_mission_panel_travel"):
                    mission_section(3, "Schedule & travel")
                    duration_col, start_col, departure_col = st.columns(3)
                    with duration_col:
                        mission_input("duration")
                    with start_col:
                        mission_input("start_date")
                    with departure_col:
                        mission_input("departure_time")
                    return_col, end_col = st.columns(2)
                    with return_col:
                        mission_input("return_time")
                    with end_col:
                        mission_input("end_date")
                    distance_col, transport_col = st.columns(2)
                    with distance_col:
                        mission_input("distance")
                    with transport_col:
                        mission_input("transport")

                with st.container(key="rng_mission_panel_notes"):
                    mission_section(4, "Mission notes & outcomes")
                    st.markdown('<div class="rng-mission-help">Capture what you set out to do, what happened, and what comes next.</div>', unsafe_allow_html=True)
                    mission_input("equipment")
                    mission_text_area("description", 180)
                    objective_col, result_col = st.columns(2)
                    with objective_col:
                        mission_text_area("objective", 130)
                    with result_col:
                        mission_text_area("result", 130)

                with st.container(key="rng_mission_panel_review"):
                    mission_section(5, "Review & signatures")
                    mission_text_area("manager_comment", 140)
                    employee_col, manager_col = st.columns(2)
                    with employee_col:
                        mission_input("employee_signature")
                    with manager_col:
                        mission_input("manager_signature")

                with st.container(key="rng_mission_save_action"):
                    if st.button("Save report", key=f"mission_save_form_{mission_id}", use_container_width=True):
                        try:
                            rng_save()
                            st.success("Mission report saved.")
                        except Exception as exc:
                            st.error("Could not save your report. Please retry. " + str(exc))

                # =================================================
                # PDF DOWNLOAD
                # =================================================

                with st.container(key="rng_mission_export"):
                    st.markdown('<div class="rng-mission-section">Ready to send</div><div class="rng-mission-help">Download or share the report in the original A4 portrait format.</div>', unsafe_allow_html=True)

                    try:

                        pdf_bytes = create_mission_pdf(
                            form,
                            mission_template
                        )

                        safe_name = re.sub(
                            r"[^\w.-]+",
                            "_",
                            mission_name
                        )

                        mission_pdf_filename = f"Mission_Report_{safe_name}.pdf"
                        download_col, share_col = st.columns(2)
                        with download_col:
                            st.download_button(
                                label="Download A4 Portrait PDF",
                                data=pdf_bytes,
                                file_name=mission_pdf_filename,
                                mime="application/pdf",
                                key=f"mission_pdf_{mission_id}",
                                use_container_width=True
                            )
                        with share_col:
                            rng_pdf_share_button(pdf_bytes, mission_pdf_filename)

                    except Exception as error:

                        st.error(
                            "PDF generation error: "
                            + str(error)
                        )

    # =========================================================
    # MISSION REPORT MAIN PAGE
    # =========================================================

    else:

        st.session_state.selected_mission_report = None

        report_total = len(st.session_state.mission_report_cards)
        st.markdown(
            '<div class="rng-mission-hero"><div class="rng-mission-eyebrow">EVERY ASSIGNMENT HAS A STORY.</div>'
            '<h1>Mission Report</h1><p>Keep the destination, the details, and the outcome together.</p>'
            f'<span class="rng-mission-total">{report_total} {"report" if report_total == 1 else "reports"} in your workspace</span></div>',
            unsafe_allow_html=True
        )

        # =====================================================
        # ADD MISSION REPORT
        # =====================================================

        if st.session_state.adding_mission_report:
            with st.container(key="rng_mission_add_form"):

                st.markdown(
                    '<div class="section-title">'
                    'Add New Mission Report'
                    '</div>',
                    unsafe_allow_html=True
                )

                new_title = st.text_input(
                    "Mission report title",
                    placeholder="Enter mission report title",
                    key="new_mission_report_title"
                )

                add_col, cancel_col = st.columns(2)

                with add_col:

                    if st.button(
                        "Add Mission Report",
                        key="mission_report_add_confirm",
                        use_container_width=True
                    ):

                        if new_title.strip():

                            st.session_state.mission_report_cards.append(
                                new_title.strip()
                            )

                            st.session_state.mission_report_card_ids.append(
                                uuid4().hex
                            )

                            st.session_state.adding_mission_report = False
                            rng_rerun()

                        else:

                            st.warning(
                                "Please enter a mission report title."
                            )

                with cancel_col:

                    if st.button(
                        "Cancel",
                        key="mission_report_add_cancel",
                        use_container_width=True
                    ):

                        st.session_state.adding_mission_report = False
                        rng_rerun()

        else:
            with st.container(key="rng_mission_add_action"):

                if st.button(
                    "+ Add Mission Report",
                    key="mission_report_add_button"
                ):

                    st.session_state.adding_mission_report = True
                    rng_rerun()

        st.markdown(
            "<br>",
            unsafe_allow_html=True
        )

        # =====================================================
        # MISSION REPORT CARDS
        # =====================================================

        if not st.session_state.mission_report_cards:
            st.markdown('<div class="rng-mission-empty"><strong>A new assignment. A clear record.</strong>Add a mission report to document the journey and its outcome.</div>', unsafe_allow_html=True)
        with st.container(key="rng_mission_board"):
            cols = st.columns(3)

            for i, mission in enumerate(
                st.session_state.mission_report_cards
            ):

                with cols[i % 3], st.container(key=f"rng_mission_card_{i}"):

                    mission_card_heading(mission, i)

                    safe_title = escape(mission)

                    # =================================================
                    # DELETE CONFIRMATION
                    # =================================================

                    if st.session_state.deleting_mission_report == i:

                        st.markdown(
                            f'<div class="mission-report-delete-warning">Delete "{safe_title}"? This cannot be undone.</div>',
                            unsafe_allow_html=True
                        )

                        confirm_col, cancel_col = st.columns(2)

                        with confirm_col:

                            if st.button(
                                "Confirm Delete",
                                key=f"mission_report_confirm_delete_{i}",
                                use_container_width=True
                            ):

                                deleted_id = (
                                    st.session_state.mission_report_card_ids.pop(i)
                                )

                                st.session_state.mission_report_cards.pop(i)

                                st.session_state.mission_report_forms.pop(
                                    deleted_id,
                                    None
                                )

                                st.session_state.deleting_mission_report = None
                                st.session_state.editing_mission_report = None

                                rng_rerun()

                        with cancel_col:

                            if st.button(
                                "Cancel",
                                key=f"mission_report_cancel_delete_{i}",
                                use_container_width=True
                            ):

                                st.session_state.deleting_mission_report = None

                                rng_rerun()

                    # =================================================
                    # EDIT MISSION REPORT TITLE
                    # =================================================

                    elif st.session_state.editing_mission_report == i:

                        edited_title = st.text_input(
                            "Mission report title",
                            value=mission,
                            key=f"mission_report_edit_input_{i}"
                        )

                        save_col, cancel_col = st.columns(2)

                        with save_col:

                            if st.button(
                                "Save",
                                key=f"mission_report_save_{i}",
                                use_container_width=True
                            ):

                                if edited_title.strip():

                                    st.session_state.mission_report_cards[i] = (
                                        edited_title.strip()
                                    )

                                    st.session_state.editing_mission_report = None

                                    rng_rerun()

                                else:

                                    st.warning(
                                        "Please enter a mission report title."
                                    )

                        with cancel_col:

                            if st.button(
                                "Cancel",
                                key=f"mission_report_cancel_{i}",
                                use_container_width=True
                            ):

                                st.session_state.editing_mission_report = None

                                rng_rerun()

                    # =================================================
                    # NORMAL MISSION REPORT CARD
                    # =================================================

                    else:

                        button_col1, button_col2, button_col3 = st.columns(3)

                        with button_col1, st.container(key=f"rng_mission_open_{i}"):

                            open_button = st.button(
                                "Open",
                                key=f"mission_report_open_{i}",
                                use_container_width=True
                            )

                        with button_col2:

                            edit_button = st.button(
                                "Edit",
                                key=f"mission_report_edit_{i}",
                                use_container_width=True
                            )

                        with button_col3:

                            delete_button = st.button(
                                "Delete",
                                key=f"mission_report_delete_{i}",
                                use_container_width=True
                            )

                        if open_button:

                            st.session_state.selected_mission_report = i

                            rng_rerun()

                        if edit_button:

                            st.session_state.editing_mission_report = i
                            st.session_state.deleting_mission_report = None

                            rng_rerun()

                        if delete_button:

                            st.session_state.deleting_mission_report = i
                            st.session_state.editing_mission_report = None

                            rng_rerun()



elif page == "Rah Nia Development":
    from datetime import datetime
    from html import escape

    # =========================================================
    # TASK DATA
    # =========================================================

    if "development_tasks" not in st.session_state:
        st.session_state.development_tasks = []

    if "development_creating_task" not in st.session_state:
        st.session_state.development_creating_task = False

    if "development_editing_task" not in st.session_state:
        st.session_state.development_editing_task = None

    if "development_new_subtask_count" not in st.session_state:
        st.session_state.development_new_subtask_count = 1

    # =========================================================
    # CSS
    # =========================================================

    st.markdown('<div id="rng-development-marker"></div>', unsafe_allow_html=True)
    st.markdown("""
    <style>
    [data-testid="stMain"]:has(#rng-development-marker) { background: radial-gradient(ellipse at 100% 0%, #dbe8fb, transparent 50%), #f2f5fa; }
    .stApp:has(#rng-development-marker), .stApp:has(#rng-development-marker) h1,
    .stApp:has(#rng-development-marker) h2, .stApp:has(#rng-development-marker) h3,
    .stApp:has(#rng-development-marker) [data-testid="stMarkdownContainer"],
    .stApp:has(#rng-development-marker) input, .stApp:has(#rng-development-marker) textarea,
    .stApp:has(#rng-development-marker) label { font-family: Arial, Tahoma, sans-serif !important; font-style: normal !important; }
    .stApp:has(#rng-development-marker) .block-container { max-width: 1440px; padding-top: 2.5rem; }
    .rng-development-heading { background: radial-gradient(ellipse at right top, #305b89, transparent 65%), #142c4b; border: 1px solid #29496e; border-radius: 24px; padding: 32px 34px; margin: 4px 0 8px; position: relative; overflow: hidden; }
    .rng-development-eyebrow { color: #acc8ec; font-size: 11px; font-weight: 700; letter-spacing: 2px; margin-bottom: 12px; }
    .rng-development-heading h1 { color: #f5f9ff !important; font-size: clamp(32px, 3.5vw, 44px) !important; font-weight: 750; letter-spacing: -1.3px; margin: 0 0 12px; padding: 0; }
    .rng-development-heading p { color: #c5d7ed; font-size: 15px; margin: 0; line-height: 1.7; }
    .rng-development-summary { display: grid; grid-template-columns: repeat(3, 1fr); gap: 14px; margin: 8px 0 4px; }
    .rng-development-stat { background: #ffffff; border: 1px solid #dce5f1; border-radius: 16px; padding: 19px 22px; }
    .rng-development-stat strong { display: block; color: #203d64; font-size: 27px; line-height: 1.2; font-weight: 700; margin-bottom: 6px; }
    .rng-development-stat span { color: #607897; font-size: 12px; }
    .rng-development-stat:last-child { background: #eaf1fc; border-color: #d4e1f4; }
    .st-key-rng_development_create_action { margin: 8px 0 12px; }
    .st-key-rng_development_create_action .stButton button { background: #214b7f; color: white; border: 1px solid #214b7f; min-height: 48px; border-radius: 12px; }
    .st-key-rng_development_create_action .stButton button:hover { color: #fff; background: #173b68; }
    .rng-development-section { display: flex; align-items: center; justify-content: space-between; color: #253f62; font-size: 19px; font-weight: 700; margin: 12px 0 4px; }
    .rng-development-section span { font-size: 12px; font-weight: 400; color: #69809d; }
    .st-key-rng_development_create_form { background: #fff; border: 1px solid #d5e2f3; border-radius: 18px; padding: 24px; margin-bottom: 16px; box-shadow: 0 8px 30px #173b6810; }
    .st-key-rng_development_create_form h3 { font-size: 21px; color: #203d64; }
    [class*="st-key-rng_development_card_"] { background: #fff; border: 1px solid #dde5f0; border-top: 3px solid #7195c6; border-radius: 18px; padding: 22px; box-shadow: 0 6px 20px #16335406; min-height: 340px; }
    [class*="st-key-rng_development_card_"]:has(.rng-task-done) { background: #f5f8ff; border-top-color: #a1b6d8; }
    .rng-task-status { display: inline-block; color: #3f6190; background: #edf3fc; border: 1px solid #dce7f7; border-radius: 100px; padding: 5px 10px; font-size: 10px; font-weight: 700; letter-spacing: .6px; text-transform: uppercase; margin-bottom: 12px; }
    .rng-task-done { background: #e4edf9; color: #4d6484; }
    .todo-title { color: #243b5b !important; font-size: 19px; font-weight: 700; line-height: 1.5; margin: 0 0 12px; overflow-wrap: anywhere; }
    .todo-section-title { color: #7a8da8; font-size: 10px; font-weight: 700; letter-spacing: 1.5px; text-transform: uppercase; margin: 14px 0 4px; }
    .todo-divider { height: 1px; background: #e8edf5; margin: 8px 0; }
    .todo-date, .todo-completed { color: #7588a3; font-size: 11px; line-height: 1.7; }
    .todo-completed { color: #496d9c; }
    .rng-subtask-progress { height: 4px; background: #edf2f9; border-radius: 10px; overflow: hidden; margin: 9px 0 6px; }
    .rng-subtask-progress span { display: block; height: 100%; background: #769bcf; border-radius: inherit; }
    .rng-subtask-count { color: #768ba6; font-size: 11px; }
    .st-key-rng_development_board [data-testid="stCheckbox"] { padding: 0; margin: 2px 0; }
    .st-key-rng_development_board [data-testid="stCheckbox"] label { min-height: 36px; }
    .st-key-rng_development_board [data-testid="stCheckbox"] label:has(input:checked) > div:not([data-testid]) { background-color: #527db4; border-color: #527db4; }
    .st-key-rng_development_board [data-testid="stCheckbox"] label p { color: #435a76; font-size: 13px; overflow-wrap: anywhere; }
    .st-key-rng_development_board .stButton button, .st-key-rng_development_create_form .stButton button { color: #355780; background: #f3f6fc; border: 1px solid #dde6f4; border-radius: 9px; min-height: 44px; }
    .st-key-rng_development_board .stButton button:hover, .st-key-rng_development_create_form .stButton button:hover { border-color: #8aabd8; color: #224773; background: #eaf1fc; }
    .stApp:has(#rng-development-marker) button:focus-visible { outline: 3px solid #8aaff0; outline-offset: 2px; }
    .rng-development-empty { text-align: center; background: #ffffffb8; border: 1px dashed #c3d2e7; border-radius: 20px; padding: 52px 20px; margin-top: 10px; }
    .rng-development-empty svg { width: 48px; height: 48px; color: #87a5cd; margin-bottom: 18px; }
    .rng-development-empty strong { display: block; color: #304d73; font-size: 21px; margin-bottom: 10px; }
    .rng-development-empty p { color: #7185a0; font-size: 14px; line-height: 1.8; margin: 0; }
    @media (max-width: 1100px) and (min-width: 701px) {
        .st-key-rng_development_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_development_card_"]) { flex-wrap: wrap; }
        .st-key-rng_development_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_development_card_"]) > [data-testid="stColumn"] { min-width: calc(50% - 1rem); }
    }
    @media (max-width: 700px) {
        .rng-development-heading { padding: 26px 22px; border-radius: 20px; }
        .rng-development-summary { gap: 8px; }
        .rng-development-stat { padding: 15px 10px; }
        .rng-development-stat strong { font-size: 24px; }
        .rng-development-stat span { font-size: 11px; }
        .rng-development-section span { display: none; }
        .st-key-rng_development_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_development_card_"]) { flex-wrap: wrap; }
        .st-key-rng_development_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_development_card_"]) > [data-testid="stColumn"] { min-width: 100%; }
        [class*="st-key-rng_development_card_"] { padding: 20px; }
        .st-key-rng_development_create_form { padding: 18px; }
    }
    </style>
    """, unsafe_allow_html=True)

    task_total = len(st.session_state.development_tasks)
    done_total = sum(bool(task["done"]) for task in st.session_state.development_tasks)
    st.markdown(
        '<div class="rng-development-heading"><div class="rng-development-eyebrow">BUILDING WHAT COMES NEXT</div>'
        '<h1>Rah Nia Development</h1><p>Move Rah Nia forward. One task at a time.</p></div>'
        '<div class="rng-development-summary">'
        f'<div class="rng-development-stat"><strong>{task_total}</strong><span>Total tasks</span></div>'
        f'<div class="rng-development-stat"><strong>{task_total - done_total}</strong><span>In progress</span></div>'
        f'<div class="rng-development-stat"><strong>{done_total}</strong><span>Completed</span></div></div>',
        unsafe_allow_html=True
    )
    with st.container(key="rng_development_create_action"):
        if st.button("Create Task", key="development_create_task_button", use_container_width=True):
            st.session_state.development_creating_task = True
            st.session_state.development_new_subtask_count = 1
            rng_rerun()

    # =========================================================
    # CREATE TASK FORM
    # =========================================================

    if st.session_state.development_creating_task:
        with st.container(key="rng_development_create_form"):

            st.markdown("### Create New Task")

            new_title = st.text_input(
                "Task Title",
                placeholder="Enter task title...",
                key="development_new_task_title"
            )

            st.markdown("Sub Tasks")

            subtask_values = []

            for i in range(st.session_state.development_new_subtask_count):

                value = st.text_input(
                    f"Sub Task {i + 1}",
                    placeholder="Enter sub task...",
                    key=f"development_create_subtask_{i}"
                )

                subtask_values.append(value)

            add_col, create_col, cancel_col = st.columns(3)

            with add_col:
                if st.button(
                    "Add Sub Task",
                    key="development_add_subtask_button",
                    use_container_width=True
                ):
                    st.session_state.development_new_subtask_count += 1
                    rng_rerun()

            with create_col:
                if st.button(
                    "Save Task",
                    key="development_save_new_task",
                    use_container_width=True
                ):

                    if new_title.strip():

                        subtasks = []

                        for value in subtask_values:

                            if value.strip():
                                subtasks.append({
                                    "text": value.strip(),
                                    "done": False
                                })

                        st.session_state.development_tasks.append({
                            "title": new_title.strip(),
                            "subtasks": subtasks,
                            "created": datetime.now().strftime("%Y-%m-%d"),
                            "completed": None,
                            "done": False,
                        })

                        st.session_state.development_creating_task = False
                        st.session_state.development_new_subtask_count = 1

                        rng_rerun()

            with cancel_col:
                if st.button(
                    "Cancel",
                    key="development_cancel_new_task",
                    use_container_width=True
                ):
                    st.session_state.development_creating_task = False
                    st.session_state.development_new_subtask_count = 1
                    rng_rerun()

            st.divider()

    # =========================================================
    # SORT TASKS
    # =========================================================

    unfinished = [
        task for task in st.session_state.development_tasks
        if not task["done"]
    ]

    completed = [
        task for task in st.session_state.development_tasks
        if task["done"]
    ]

    sorted_tasks = unfinished + completed

    def development_persist_change():
        try:
            rng_save()
        except Exception as exc:
            st.error("The change could not be saved. Please retry. Details: " + str(exc))
            st.stop()

    def development_toggle_main(task_index):
        task = st.session_state.development_tasks[task_index]
        task["done"] = bool(st.session_state[f"development_main_checkbox_{task_index}"])
        task["completed"] = datetime.now().strftime("%Y-%m-%d") if task["done"] else None
        development_persist_change()

    def development_toggle_subtask(task_index, sub_index):
        task = st.session_state.development_tasks[task_index]
        task["subtasks"][sub_index]["done"] = bool(st.session_state[f"development_sub_checkbox_{task_index}_{sub_index}"])
        task["done"] = bool(task["subtasks"]) and all(item["done"] for item in task["subtasks"])
        task["completed"] = datetime.now().strftime("%Y-%m-%d") if task["done"] else None
        # Callbacks run before rendering, so the main checkbox shares the new state.
        st.session_state[f"development_main_checkbox_{task_index}"] = task["done"]
        development_persist_change()

    # =========================================================
    # TASK CARDS
    # =========================================================

    st.markdown(
        '<div class="rng-development-section">Your tasks<span>Unfinished first · Progress always visible</span></div>',
        unsafe_allow_html=True
    )
    if not sorted_tasks:
        st.markdown(
            '<div class="rng-development-empty"><svg viewBox="0 0 48 48" fill="none" stroke="currentColor" stroke-width="1.5" aria-hidden="true"><rect x="10" y="9" width="28" height="33" rx="5"/><rect x="18" y="5" width="12" height="8" rx="3" fill="#f2f5fa"/><path d="m17 26 5 5 10-12"/></svg>'
            '<strong>A clear list. A fresh start.</strong><p>Create your first task and break it into small, achievable steps.</p></div>',
            unsafe_allow_html=True
        )

    with st.container(key="rng_development_board"):
        for row_start in range(0, len(sorted_tasks), 3):

            row_tasks = sorted_tasks[row_start:row_start + 3]

            columns = st.columns(3)

            for column_number, task in enumerate(row_tasks):

                task_index = st.session_state.development_tasks.index(task)
                with columns[column_number], st.container(key=f"rng_development_card_{task_index}"):

                    # =================================================
                    # EDIT MODE
                    # =================================================

                    if st.session_state.development_editing_task == task_index:

                        st.markdown(
                            '<div class="todo-title">Edit Task</div>',
                            unsafe_allow_html=True
                        )

                        st.markdown(
                            '<div class="todo-divider"></div>',
                            unsafe_allow_html=True
                        )

                        edited_title = st.text_input(
                            "Task Title",
                            value=task["title"],
                            key=f"development_edit_title_{task_index}"
                        )

                        st.markdown("Sub Tasks")

                        edited_subtasks = []

                        for sub_index, subtask in enumerate(task["subtasks"]):

                            edited_text = st.text_input(
                                f"Sub Task {sub_index + 1}",
                                value=subtask["text"],
                                key=f"development_edit_subtask_{task_index}_{sub_index}"
                            )

                            edited_subtasks.append({
                                "text": edited_text,
                                "done": subtask["done"]
                            })

                        new_edit_subtask = st.text_input(
                            "Add another sub task",
                            placeholder="Optional",
                            key=f"development_edit_new_subtask_{task_index}"
                        )

                        save_col, cancel_col = st.columns(2)

                        with save_col:

                            if st.button(
                                "Save Changes",
                                key=f"development_save_edit_{task_index}",
                                use_container_width=True
                            ):

                                if edited_title.strip():
                                    task["title"] = edited_title.strip()

                                task["subtasks"] = [
                                    subtask
                                    for subtask in edited_subtasks
                                    if subtask["text"].strip()
                                ]

                                if new_edit_subtask.strip():
                                    task["subtasks"].append({
                                        "text": new_edit_subtask.strip(),
                                        "done": False
                                    })

                                if (
                                    len(task["subtasks"]) > 0
                                    and all(
                                        subtask["done"]
                                        for subtask in task["subtasks"]
                                    )
                                ):

                                    task["done"] = True

                                    if not task["completed"]:
                                        task["completed"] = (
                                            datetime.now().strftime("%Y-%m-%d")
                                        )

                                else:

                                    task["done"] = False
                                    task["completed"] = None

                                st.session_state.development_editing_task = None
                                rng_rerun()

                        with cancel_col:

                            if st.button(
                                "Cancel",
                                key=f"development_cancel_edit_{task_index}",
                                use_container_width=True
                            ):
                                st.session_state.development_editing_task = None
                                rng_rerun()

                        continue

                    # =================================================
                    # CARD TITLE
                    # =================================================

                    finished_steps = sum(bool(item["done"]) for item in task["subtasks"])
                    step_total = len(task["subtasks"])
                    status = "Completed" if task["done"] else ("In progress" if finished_steps else "To do")
                    done_class = " rng-task-done" if task["done"] else ""
                    st.markdown(
                        f'<div class="rng-task-status{done_class}">{status}</div>'
                        f'<div class="todo-title">{escape(str(task["title"]))}</div>',
                        unsafe_allow_html=True
                    )
                    if step_total:
                        st.markdown(
                            f'<div class="rng-subtask-progress"><span style="width:{100 * finished_steps / step_total:.1f}%"></span></div>'
                            f'<div class="rng-subtask-count">{finished_steps} of {step_total} subtasks complete</div>',
                            unsafe_allow_html=True
                        )

                    # =================================================
                    # MAIN TASK CHECKBOX
                    # =================================================

                    st.session_state[f"development_main_checkbox_{task_index}"] = task["done"]
                    st.checkbox(
                        "Task completed", key=f"development_main_checkbox_{task_index}",
                        on_change=development_toggle_main, args=(task_index,)
                    )

                    # =================================================
                    # SUBTASKS
                    # =================================================

                    st.markdown(
                        '<div class="todo-divider"></div>',
                        unsafe_allow_html=True
                    )

                    st.markdown(
                        '<div class="todo-section-title">Sub Tasks</div>',
                        unsafe_allow_html=True
                    )

                    for sub_index, subtask in enumerate(task["subtasks"]):

                        st.session_state[f"development_sub_checkbox_{task_index}_{sub_index}"] = subtask["done"]
                        st.checkbox(
                            subtask["text"], key=f"development_sub_checkbox_{task_index}_{sub_index}",
                            on_change=development_toggle_subtask, args=(task_index, sub_index)
                        )

                    # =================================================
                    # DATES
                    # =================================================

                    st.markdown(
                        '<div class="todo-divider"></div>',
                        unsafe_allow_html=True
                    )

                    st.markdown(
                        f"""
                        <div class="todo-date">
                            Created: {task["created"]}
                        </div>
                        """,
                        unsafe_allow_html=True
                    )

                    if task["completed"]:

                        st.markdown(
                            f"""
                            <div class="todo-completed">
                                Completed: {task["completed"]}
                            </div>
                            """,
                            unsafe_allow_html=True
                        )

                    else:

                        st.markdown(
                            """
                            <div class="todo-date">
                                Completed: -
                            </div>
                            """,
                            unsafe_allow_html=True
                        )

                    # =================================================
                    # EDIT BUTTON
                    # =================================================

                    rng_task_actions("development_tasks", task_index, "development_editing_task", "development_")



elif page == "BOM":

    import io
    import re
    import hashlib
    from pathlib import Path
    from uuid import uuid4
    from html import escape

    # =========================================================
    # ORIGINAL BOM EXCEL TEMPLATE
    # =========================================================

    BOM_TEMPLATE_PATH = (
        Path(__file__).resolve().parent / "BOM Raw.xlsx"
    )

    def load_bom_template():

        from openpyxl import load_workbook

        if not BOM_TEMPLATE_PATH.is_file():
            raise FileNotFoundError(
                "BOM Raw.xlsx was not found beside app.py"
            )

        workbook = load_workbook(
            BOM_TEMPLATE_PATH,
            read_only=False,
            data_only=False
        )

        if "RAW" not in workbook.sheetnames:
            workbook.close()
            raise ValueError("RAW worksheet not found.")

        ws = workbook["RAW"]

        labels = {
            "title": str(ws["C1"].value or ""),
            "metadata": str(ws["J1"].value or ""),
            "headers": [
                str(ws[f"{col}2"].value or "")
                for col in "ABDEFGHI"
            ],
            "preparer": str(ws["A18"].value or ""),
            "approver": str(ws["D18"].value or "")
        }

        widths = [
            ws.column_dimensions[c].width or 10
            for c in "ABCDEFGHIJK"
        ]

        heights = [
            ws.row_dimensions[r].height or 25
            for r in range(1, 19)
        ]

        workbook.close()

        return {
            "labels": labels,
            "widths": widths,
            "heights": heights
        }

    bom_template = None
    bom_template_error = None

    try:
        bom_template = load_bom_template()

    except Exception as error:
        bom_template_error = str(error)

    # =========================================================
    # BOM SESSION STATE
    # =========================================================

    if "bom_cards" not in st.session_state:
        st.session_state.bom_cards = []

    if "selected_bom" not in st.session_state:
        st.session_state.selected_bom = None

    if "editing_bom" not in st.session_state:
        st.session_state.editing_bom = None

    if "deleting_bom" not in st.session_state:
        st.session_state.deleting_bom = None

    if "adding_bom" not in st.session_state:
        st.session_state.adding_bom = False

    if "bom_card_ids" not in st.session_state:
        st.session_state.bom_card_ids = [
            uuid4().hex
            for _ in st.session_state.bom_cards
        ]

    if "bom_forms" not in st.session_state:
        st.session_state.bom_forms = {}

    while len(st.session_state.bom_card_ids) < len(
        st.session_state.bom_cards
    ):
        st.session_state.bom_card_ids.append(
            uuid4().hex
        )

    # =========================================================
    # BOM DATA STRUCTURE
    # =========================================================

    def create_empty_bom_table():

        return [
            {
                "description": "",
                "size": "",
                "actual": "",
                "over": "",
                "unit": "",
                "notes": ""
            }
            for _ in range(15)
        ]

    def create_empty_bom():

        return {
            "date": "",
            "number": "",
            "code": "",
            "preparer": "",
            "approver": "",
            "items": create_empty_bom_table(),
            "extra_tables": []
        }

    # =========================================================
    # SUPPORT EXISTING BOM DATA
    # =========================================================

    def prepare_bom_form(form):

        # The original items list remains the first table.
        # This preserves previously entered data.

        if "items" not in form:
            form["items"] = create_empty_bom_table()

        if "extra_tables" not in form:
            form["extra_tables"] = []

        for field in [
            "date",
            "number",
            "code",
            "preparer",
            "approver"
        ]:
            if field not in form:
                form[field] = ""

        # Ensure every existing table has 15 rows.
        all_tables = [form["items"]] + form["extra_tables"]

        for table in all_tables:

            while len(table) < 15:
                table.append({
                    "description": "",
                    "size": "",
                    "actual": "",
                    "over": "",
                    "unit": "",
                    "notes": ""
                })

            for item in table:

                for field in [
                    "description",
                    "size",
                    "actual",
                    "over",
                    "unit",
                    "notes"
                ]:
                    if field not in item:
                        item[field] = ""

    # =========================================================
    # NUMBER CALCULATION
    # =========================================================

    def parse_bom_number(value):

        text = str(value).strip()

        if not text:
            return 0.0

        digit_map = str.maketrans(
            "\u06f0\u06f1\u06f2\u06f3\u06f4\u06f5\u06f6\u06f7\u06f8\u06f9"
            "\u0660\u0661\u0662\u0663\u0664\u0665\u0666\u0667\u0668\u0669",
            "01234567890123456789"
        )

        text = text.translate(digit_map)

        text = text.replace("\u066c", "")
        text = text.replace(",", "")
        text = text.replace("\u066b", ".")

        return float(text)

    def calculate_bom_quantity(item):

        actual_text = str(
            item.get("actual", "")
        ).strip()

        over_text = str(
            item.get("over", "")
        ).strip()

        if not actual_text and not over_text:
            return ""

        try:

            actual = parse_bom_number(actual_text)
            over = parse_bom_number(over_text)

            total = actual + over

            if not (
                float("-inf") < total < float("inf")
            ):
                return "Invalid"

            if total.is_integer():
                return str(int(total))

            return f"{total:.6f}".rstrip("0").rstrip(".")

        except (ValueError, TypeError):
            return "Invalid"

    def update_bom_table(bom_id, table_index, editor_key):
        target = st.session_state.bom_forms[bom_id]
        tables = [target["items"]] + target["extra_tables"]
        allowed_fields = {"description", "size", "actual", "over", "unit", "notes"}
        changes = st.session_state.get(editor_key, {}).get("edited_rows", {})
        for row, values in changes.items():
            row_index = int(row)
            if 0 <= row_index < len(tables[table_index]):
                for field, value in values.items():
                    if field in allowed_fields:
                        tables[table_index][row_index][field] = "" if value is None else str(value)

    # =========================================================
    # BOM CSS
    # =========================================================

    st.markdown('<div id="rng-bom-marker"></div>', unsafe_allow_html=True)
    st.markdown("""
    <style>
    [data-testid="stMain"]:has(#rng-bom-marker) { background: radial-gradient(ellipse at right top, #dce8f9, transparent 55%), #f2f5fa; }
    .stApp:has(#rng-bom-marker), .stApp:has(#rng-bom-marker) h1,
    .stApp:has(#rng-bom-marker) h2, .stApp:has(#rng-bom-marker) h3,
    .stApp:has(#rng-bom-marker) [data-testid="stMarkdownContainer"],
    .stApp:has(#rng-bom-marker) input, .stApp:has(#rng-bom-marker) textarea,
    .stApp:has(#rng-bom-marker) label { font-family: Arial, Tahoma, sans-serif !important; font-style: normal !important; }
    .stApp:has(#rng-bom-marker) .block-container { max-width: 1440px; }
    .rng-bom-hero { background: radial-gradient(ellipse at right top, #355e89, transparent 65%), #172f4e; border: 1px solid #315174; border-radius: 24px; padding: 30px 34px; margin: 8px 0 14px; }
    .rng-bom-eyebrow { color: #b1cae9; font-size: 11px; font-weight: 700; letter-spacing: 2px; margin-bottom: 14px; }
    .rng-bom-hero h1 { color: #f6f9ff !important; font-size: clamp(30px, 3.4vw, 44px) !important; font-weight: 750; letter-spacing: -1.2px; margin: 0 0 12px; padding: 0; overflow-wrap: anywhere; }
    .rng-bom-hero p { color: #c6d8ed; font-size: 15px; line-height: 1.8; margin: 0; }
    .rng-bom-total { display: inline-block; color: #dce9fb; background: #a9ccff12; border: 1px solid #a9ccff30; border-radius: 100px; padding: 6px 12px; font-size: 11px; margin-top: 18px; }
    .st-key-rng_bom_add_action .stButton button, .st-key-rng_bom_save_action .stButton button { color: #fff; background: #254d7e; border: 1px solid #254d7e; min-height: 46px; border-radius: 11px; }
    .st-key-rng_bom_add_action .stButton button:hover, .st-key-rng_bom_save_action .stButton button:hover { background: #18395f; color: #fff; }
    .st-key-rng_bom_add_form { background: #fff; border: 1px solid #d7e3f2; padding: 24px; border-radius: 18px; margin-bottom: 20px; }
    [class*="st-key-rng_bom_card_"] { background: #fff; border: 1px solid #dce5f1; border-radius: 18px; padding: 24px; min-height: 270px; box-shadow: 0 8px 25px #15385f07; }
    .rng-bom-card-top { display: flex; align-items: center; justify-content: space-between; margin-bottom: 19px; }
    .rng-bom-icon { display: flex; align-items: center; justify-content: center; width: 46px; height: 46px; color: #5279aa; background: #eaf1fc; border-radius: 14px; }
    .rng-bom-icon svg { width: 25px; height: 25px; }
    .rng-bom-index { color: #8698b0; font-size: 11px; letter-spacing: 1px; }
    .bom-card-title { color: #243e60; font-size: 19px; font-weight: 700; line-height: 1.5; margin: 0 0 10px; overflow-wrap: anywhere; }
    .bom-card-text { color: #6f839e; font-size: 12px; line-height: 1.8; margin-bottom: 13px; overflow-wrap: anywhere; }
    .rng-bom-card-footer { color: #8092aa; font-size: 10px; letter-spacing: .6px; border-top: 1px solid #e8edf5; padding-top: 14px; margin-bottom: 6px; }
    [class*="st-key-rng_bom_card_"] [data-testid="stHorizontalBlock"] > [data-testid="stColumn"] { min-width: 0 !important; flex: 1; }
    .st-key-rng_bom_board .stButton button { color: #48658b; background: #f5f8fd; border: 1px solid #dfe7f3; border-radius: 9px; min-height: 44px; }
    .st-key-rng_bom_board .stButton button:hover { background: #eaf1fc; color: #234b7b; border-color: #8faed7; }
    [class*="st-key-rng_bom_open_"] .stButton button { background: #244d7c; border-color: #244d7c; color: #fff; }
    [class*="st-key-rng_bom_open_"] .stButton button:hover { background: #17395f; color: #fff; }
    .bom-delete-warning { color: #8c4050; background: #fff2f4; border: 1px solid #f1d4dc; border-radius: 12px; padding: 15px; font-size: 13px; line-height: 1.8; overflow-wrap: anywhere; }
    .rng-bom-empty { text-align: center; padding: 50px 24px; border: 1px dashed #bdcee5; border-radius: 18px; background: #fff9; color: #70859f; }
    .rng-bom-empty strong { display: block; color: #304e75; font-size: 21px; margin-bottom: 10px; }
    .rng-bom-table-title { color: #294569; font-size: 21px; font-weight: 700; margin: 22px 0 7px; }
    .st-key-rng_bom_add_table_action .stButton button { color: #365e90; background: #edf3fc; border: 1px dashed #9bb7dc; border-radius: 12px; min-height: 48px; }
    .rng-bom-section { color: #294569; font-size: 21px; font-weight: 700; margin: 4px 0 8px; }
    .rng-bom-help { color: #7388a4; font-size: 13px; line-height: 1.8; margin-bottom: 16px; }
    .bom-form-title { direction: rtl; color: #48688e; font-size: 18px; font-weight: 600; text-align: right; margin: 4px 0 18px; }
    .st-key-rng_bom_meta, .st-key-rng_bom_approval { background: #fff; border: 1px solid #dce5f2; padding: 24px; border-radius: 18px; margin: 6px 0 12px; }
    .st-key-bom_detail [data-testid="stDataFrame"] { border: 1px solid #cbd9eb; border-radius: 0 0 14px 14px; overflow: hidden; background: #fff; max-width: 100%; box-shadow: 0 8px 24px #193a6010; }
    .rng-bom-grid-heading { display: flex; justify-content: space-between; align-items: center; gap: 16px; direction: rtl; background: linear-gradient(110deg, #264d79, #152f50); padding: 17px 20px; border-radius: 14px 14px 0 0; color: #fff; }
    .rng-bom-grid-heading strong { font: 700 15px Arial, Tahoma, sans-serif; }
    .rng-bom-grid-heading span { color: #c0d2e9; font: 12px Arial, Tahoma, sans-serif; }
    [class*="st-key-bom_form_table_"] { gap: 0 !important; }
    .st-key-bom_detail [data-testid="stDataFrame"] textarea { direction: rtl; text-align: right; font-family: Arial, Tahoma, sans-serif !important; }
    @media (max-width: 600px) { .rng-bom-grid-heading { padding: 14px 12px; } .rng-bom-grid-heading span { font-size: 11px; } }
    .st-key-bom_detail input, .st-key-bom_detail textarea { direction: rtl; text-align: start; unicode-bidi: plaintext; background: #f9fbff; color: #2f486b; min-height: 44px; }
    .st-key-bom_detail [data-testid="stWidgetLabel"] { direction: rtl; text-align: right; }
    .st-key-bom_detail label p { color: #587193; font-size: 12px; }
    .st-key-rng_bom_export { background: #e8eff9; border: 1px solid #d9e4f3; padding: 22px; border-radius: 18px; margin-top: 20px; }
    .st-key-rng_bom_export .stDownloadButton button { background: #fff; color: #345781; border: 1px solid #c5d6ed; min-height: 46px; border-radius: 10px; }
    .stApp:has(#rng-bom-marker) button:focus-visible { outline: 3px solid #9ab9e8; outline-offset: 2px; }
    @media (max-width: 1100px) and (min-width: 701px) {
        .st-key-rng_bom_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_bom_card_"]) { flex-wrap: wrap; }
        .st-key-rng_bom_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_bom_card_"]) > [data-testid="stColumn"] { min-width: calc(50% - 1rem); }
    }
    @media (max-width: 700px) {
        .rng-bom-hero { padding: 26px 22px; border-radius: 20px; }
        .st-key-rng_bom_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_bom_card_"]) { flex-wrap: wrap; }
        .st-key-rng_bom_board [data-testid="stHorizontalBlock"]:has([class*="st-key-rng_bom_card_"]) > [data-testid="stColumn"] { min-width: 100%; }
        [class*="st-key-rng_bom_card_"] { padding: 20px; }
        .st-key-rng_bom_meta, .st-key-rng_bom_approval, .st-key-rng_bom_export { padding: 18px; }
    }
    </style>
    """, unsafe_allow_html=True)

    def bom_card_heading(name, index):
        record_id = st.session_state.bom_card_ids[index]
        data = st.session_state.bom_forms.get(record_id, {})
        tables = [data.get("items", [])] + data.get("extra_tables", [])
        count = sum(any(str(value).strip() for value in item.values()) for table in tables for item in table)
        description = f"{count} material items · {len(tables)} {('table' if len(tables) == 1 else 'tables')}" if count else "Ready for your next material list."
        st.markdown(
            '<div class="rng-bom-card-top"><div class="rng-bom-icon">'
            '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.5" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><path d="m12 3 10 5-10 5L2 8Zm-10 10 10 5 10-5M2 18l10 5 10-5"/></svg>'
            f'</div><span class="rng-bom-index">BOM {index + 1:02}</span></div>'
            f'<div class="bom-card-title">{escape(str(name))}</div>'
            f'<div class="bom-card-text">{escape(description)}</div>'
            '<div class="rng-bom-card-footer">MATERIALS / QUANTITIES / APPROVALS</div>',
            unsafe_allow_html=True
        )

    # =========================================================
    # PDF GENERATOR - MULTIPLE A4 PORTRAIT PAGES
    # =========================================================

    def create_bom_pdf(form, template):

        import arabic_reshaper

        from bidi.algorithm import get_display
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        buffer = io.BytesIO()

        page_width, page_height = A4

        pdf = canvas.Canvas(
            buffer,
            pagesize=A4,
            pageCompression=1
        )

        labels = template["labels"]

        pdf.setTitle(labels["title"])

        # -----------------------------------------------------
        # PERSIAN FONT
        # -----------------------------------------------------

        font_name = rng_pdf_font_name()

        def shape_text(value):
            return get_display(arabic_reshaper.reshape(str(value if value is not None else "")))

        # -----------------------------------------------------
        # A4 GEOMETRY
        # -----------------------------------------------------

        margin_x = 6 * mm
        margin_y = 6 * mm

        left = margin_x
        right = page_width - margin_x

        top = page_height - margin_y
        bottom = margin_y

        usable_width = right - left
        usable_height = top - bottom

        source_widths = template["widths"]
        source_heights = template["heights"]

        width_total = sum(source_widths)
        height_total = sum(source_heights)

        col_widths = [
            usable_width * w / width_total
            for w in source_widths
        ]

        row_heights = [
            usable_height * h / height_total
            for h in source_heights
        ]

        # RTL: Excel column A is on the far right.

        x_edges = [right]

        for width in col_widths:
            x_edges.append(
                x_edges[-1] - width
            )

        y_edges = [top]

        for height in row_heights:
            y_edges.append(
                y_edges[-1] - height
            )

        # -----------------------------------------------------
        # PDF TEXT WRAPPING
        # -----------------------------------------------------

        def text_width(value, size):

            return rng_pdf_width(shape_text(value), size)

        def wrap_text(value, width, size):

            result = []

            for paragraph in str(value).splitlines():

                if not paragraph.strip():
                    result.append("")
                    continue

                current = ""

                for word in paragraph.split():

                    candidate = (
                        current + " " + word
                        if current else word
                    )

                    if text_width(candidate, size) <= width:
                        current = candidate

                    else:

                        if current:
                            result.append(current)

                        current = ""

                        for char in word:

                            candidate = current + char

                            if (
                                current
                                and text_width(candidate, size) > width
                            ):
                                result.append(current)
                                current = char

                            else:
                                current = candidate

                if current:
                    result.append(current)

            return result or [""]

        # -----------------------------------------------------
        # DRAW TEXT
        # -----------------------------------------------------

        def draw_text(
            value,
            x,
            y,
            width,
            height,
            size=8,
            align="center"
        ):

            value = str(
                value if value is not None else ""
            ).strip()

            if not value:
                return

            padding = 2.5

            available_width = max(
                width - 2 * padding,
                1
            )

            available_height = max(
                height - 2 * padding,
                1
            )

            font_size = float(size)

            while font_size > 4.5:

                lines = wrap_text(
                    value,
                    available_width,
                    font_size
                )

                line_height = font_size * 1.25

                if (
                    len(lines) * line_height
                    <= available_height
                ):
                    break

                font_size -= 0.5

            lines = wrap_text(
                value,
                available_width,
                font_size
            )

            line_height = font_size * 1.25

            maximum_lines = max(
                1,
                int(available_height / line_height)
            )

            if len(lines) > maximum_lines:

                lines = lines[:maximum_lines]

                last = lines[-1]

                while (
                    last
                    and text_width(
                        last + "...",
                        font_size
                    ) > available_width
                ):
                    last = last[:-1]

                lines[-1] = last + "..."

            pdf.saveState()

            clipping = pdf.beginPath()

            clipping.rect(
                x + 1,
                y + 1,
                max(width - 2, 0),
                max(height - 2, 0)
            )

            pdf.clipPath(
                clipping,
                stroke=0,
                fill=0
            )

            pdf.setFont(font_name, font_size)
            pdf.setFillColorRGB(0, 0, 0)

            total_height = len(lines) * line_height

            baseline = (
                y
                + (height + total_height) / 2
                - line_height
                + 2
            )

            for line in lines:

                shaped = shape_text(line)

                if align == "right":

                    rng_pdf_draw(pdf, shaped, x + width - padding, baseline, font_size, "right")

                elif align == "left":

                    rng_pdf_draw(pdf, shaped, x + padding, baseline, font_size, "left")

                else:

                    rng_pdf_draw(pdf, shaped, x + width / 2, baseline, font_size, "center")

                baseline -= line_height

            pdf.restoreState()

        # -----------------------------------------------------
        # DRAW EXCEL-STYLE CELL
        # -----------------------------------------------------

        def draw_cell(
            first_col,
            last_col,
            row,
            text="",
            size=8,
            shaded=False,
            align="center"
        ):

            x = x_edges[last_col + 1]

            width = (
                x_edges[first_col]
                - x_edges[last_col + 1]
            )

            y = y_edges[row]
            height = row_heights[row - 1]

            if shaded:

                pdf.setFillColorRGB(
                    0.93,
                    0.94,
                    0.96
                )

                pdf.rect(
                    x,
                    y,
                    width,
                    height,
                    fill=1,
                    stroke=0
                )

            pdf.setStrokeColorRGB(
                0.25,
                0.25,
                0.25
            )

            pdf.setLineWidth(0.5)

            pdf.rect(
                x,
                y,
                width,
                height,
                fill=0,
                stroke=1
            )

            draw_text(
                text,
                x,
                y,
                width,
                height,
                size=size,
                align=align
            )

        # =====================================================
        # DRAW ONE COMPLETE BOM PAGE FOR EACH TABLE
        # =====================================================

        for table_index, table_items in enumerate(all_tables):

            # -------------------------------------------------
            # ROW 1: TITLE AND DOCUMENT INFORMATION
            # -------------------------------------------------

            draw_cell(0, 1, 1)

            draw_cell(
                2,
                8,
                1,
                labels["title"],
                size=17
            )

            metadata_labels = labels["metadata"].splitlines()

            date_label = (
                metadata_labels[0]
                if len(metadata_labels) > 0
                else "Date:"
            )

            number_label = (
                metadata_labels[1]
                if len(metadata_labels) > 1
                else "Number:"
            )

            code_label = (
                metadata_labels[2]
                if len(metadata_labels) > 2
                else "Code:"
            )

            draw_cell(
                9,
                9,
                1,
                "\n".join([
                    date_label,
                    number_label,
                    code_label
                ]),
                size=9,
                align="right"
            )

            draw_cell(
                10,
                10,
                1,
                "\n".join([
                    form["date"],
                    form["number"],
                    form["code"]
                ]),
                size=9
            )

            # -------------------------------------------------
            # ROW 2: TABLE HEADINGS
            # -------------------------------------------------

            headers = labels["headers"]

            merged_ranges = [
                (0, 0),
                (1, 2),
                (3, 3),
                (4, 4),
                (5, 5),
                (6, 6),
                (7, 7),
                (8, 10)
            ]

            for index, (start, end) in enumerate(
                merged_ranges
            ):

                draw_cell(
                    start,
                    end,
                    2,
                    headers[index],
                    size=8,
                    shaded=True
                )

            # -------------------------------------------------
            # ROWS 3-17: FIFTEEN MATERIAL ITEMS
            # -------------------------------------------------

            for item_index in range(15):

                item = table_items[item_index]

                row_number = item_index + 3

                values = [
                    str(item_index + 1),
                    item["description"],
                    item["size"],
                    item["actual"],
                    item["over"],
                    calculate_bom_quantity(item),
                    item["unit"],
                    item["notes"]
                ]

                for column_index, (start, end) in enumerate(
                    merged_ranges
                ):

                    draw_cell(
                        start,
                        end,
                        row_number,
                        values[column_index],
                        size=8
                    )

            # -------------------------------------------------
            # ROW 18: PREPARER AND APPROVER
            # -------------------------------------------------

            draw_cell(
                0,
                2,
                18,
                labels["preparer"]
                + "\n"
                + form["preparer"],
                size=10,
                align="right"
            )

            draw_cell(
                3,
                10,
                18,
                labels["approver"]
                + "\n"
                + form["approver"],
                size=10,
                align="right"
            )

            # -------------------------------------------------
            # END THIS PAGE
            # -------------------------------------------------

            pdf.showPage()

        # =====================================================
        # FINALIZE MULTI-PAGE PDF
        # =====================================================

        pdf.save()

        buffer.seek(0)

        return buffer.getvalue()

    # =========================================================
    # BOM DETAIL PAGE
    # =========================================================

    selected = st.session_state.selected_bom

    if (
        isinstance(selected, int)
        and 0 <= selected < len(st.session_state.bom_cards)
    ):

        bom_name = st.session_state.bom_cards[selected]
        bom_id = st.session_state.bom_card_ids[selected]

        if bom_id not in st.session_state.bom_forms:

            st.session_state.bom_forms[
                bom_id
            ] = create_empty_bom()

        form = st.session_state.bom_forms[bom_id]

        prepare_bom_form(form)

        if st.button(
            "Back to BOM",
            key="bom_back"
        ):
            st.session_state.selected_bom = None
            rng_rerun()

        st.markdown(
            '<div class="rng-bom-hero"><div class="rng-bom-eyebrow">BILL OF MATERIALS WORKSPACE</div>'
            f'<h1>{escape(str(bom_name))}</h1>'
            '<p>A clear material list, from quantities to approval.</p></div>',
            unsafe_allow_html=True
        )

        if bom_template is None:

            st.error(
                "Unable to load BOM Excel template: "
                + str(bom_template_error)
            )

        else:

            with st.container(key="bom_detail"):

                labels = bom_template["labels"]

                with st.container(key="rng_bom_meta"):
                    st.markdown('<div class="rng-bom-section">Document details</div>', unsafe_allow_html=True)
                    st.markdown('<div class="bom-form-title">' + escape(labels["title"]) + '</div>', unsafe_allow_html=True)
                    metadata_labels = labels["metadata"].splitlines()
                    for column, (field, label_index) in zip(st.columns(3), [("code", 2), ("number", 1), ("date", 0)]):
                        with column:
                            label = metadata_labels[label_index] if label_index < len(metadata_labels) else field.capitalize()
                            form[field] = st.text_input(label, value=form[field], key=f"bom_{bom_id}_{field}")

                headers = labels["headers"]
                field_labels = dict(zip(["row", "description", "size", "actual", "over", "quantity", "unit", "notes"], headers))
                all_tables = [form["items"]] + form["extra_tables"]
                for table_index, table_items in enumerate(all_tables):
                    st.markdown(
                        f'<div class="rng-bom-table-title">Materials · Table {table_index + 1:02}</div>'
                        '<div class="rng-bom-help">Read from right to left · Double-click a cell to edit · Quantity = Actual + Over</div>',
                        unsafe_allow_html=True
                    )
                    with st.container(key=f"bom_form_table_{table_index}"):
                        material_count = sum(bool(str(item.get("description", "")).strip()) for item in table_items[:15])
                        st.markdown(
                            f'<div class="rng-bom-grid-heading"><strong>جدول اقلام · {table_index + 1:02}</strong>'
                            f'<span>{material_count} قلم ثبت‌شده · ۱۵ ردیف</span></div>',
                            unsafe_allow_html=True
                        )
                        records = [
                            {"row": row_index + 1, **item, "quantity": calculate_bom_quantity(item)}
                            for row_index, item in enumerate(table_items[:15])
                        ]
                        table_frame = pd.DataFrame(records).style.set_properties(
                            subset=["row"], **{"background-color": "#f1f5fb", "color": "#637c9d"}
                        ).set_properties(
                            subset=["quantity"], **{"background-color": "#eaf1fb", "color": "#214b7f", "font-weight": "bold"}
                        )
                        editor_key = f"bom_{bom_id}_editor_{table_index}"
                        column_config = {
                            "row": st.column_config.NumberColumn(field_labels["row"], width=60, alignment="right", disabled=True, format="%d"),
                            "description": st.column_config.TextColumn(field_labels["description"], width=300, alignment="right"),
                            "size": st.column_config.TextColumn(field_labels["size"], width=90, alignment="right"),
                            "actual": st.column_config.TextColumn(field_labels["actual"], width=100, alignment="right", help="Persian and Arabic digits are supported."),
                            "over": st.column_config.TextColumn(field_labels["over"], width=90, alignment="right"),
                            "quantity": st.column_config.TextColumn(field_labels["quantity"], width=100, alignment="right", disabled=True, help="Actual + Over, calculated automatically."),
                            "unit": st.column_config.TextColumn(field_labels["unit"], width=80, alignment="right"),
                            "notes": st.column_config.TextColumn(field_labels["notes"], width=240, alignment="right"),
                        }
                        st.data_editor(
                            table_frame, key=editor_key, hide_index=True,
                            column_order=["notes", "unit", "quantity", "over", "actual", "size", "description", "row"],
                            column_config=column_config, num_rows="fixed", width="stretch",
                            height=640, row_height=39,
                            on_change=update_bom_table, args=(bom_id, table_index, editor_key)
                        )

                st.components.v1.html("""
                <script>
                const doc = window.parent.document;
                function startOnRight() {
                    doc.querySelectorAll('[class*="st-key-bom_form_table_"] .dvn-scroller').forEach(el => {
                        if (!el.dataset.bomRtlReady && el.clientWidth > 0) {
                            el.scrollLeft = el.scrollWidth;
                            el.dataset.bomRtlReady = "true";
                        }
                    });
                }
                const observer = new MutationObserver(startOnRight);
                observer.observe(doc.body, {childList: true, subtree: true});
                startOnRight();
                window.addEventListener("unload", () => observer.disconnect());
                </script>
                """, height=0)

                # =================================================
                # ADD ANOTHER TABLE BUTTON
                # =================================================

                with st.container(key="rng_bom_add_table_action"):
                    if st.button(
                        "+ Add Another Table",
                        key=f"bom_add_table_{bom_id}",
                        use_container_width=True
                    ):

                        form["extra_tables"].append(
                            create_empty_bom_table()
                        )

                        rng_rerun()

                st.caption(
                    f"Total tables: {1 + len(form['extra_tables'])} "
                    f"| Total material rows: "
                    f"{15 * (1 + len(form['extra_tables']))}"
                )

                # =================================================
                # PREPARER AND APPROVER
                # =================================================

                with st.container(key="rng_bom_approval"):
                    st.markdown('<div class="rng-bom-section">Preparation &amp; approval</div><div class="rng-bom-help">Keep the responsible names with the material list.</div>', unsafe_allow_html=True)
                    approver_col, preparer_col = st.columns(2)
                    with approver_col:
                        form["approver"] = st.text_input(labels["approver"], value=form["approver"], key=f"bom_{bom_id}_approver")
                    with preparer_col:
                        form["preparer"] = st.text_input(labels["preparer"], value=form["preparer"], key=f"bom_{bom_id}_preparer")

                with st.container(key="rng_bom_save_action"):
                    if st.button("Save BOM", key=f"bom_save_form_{bom_id}", use_container_width=True):
                        try:
                            rng_save()
                            st.success("Bill of Materials saved.")
                        except Exception as exc:
                            st.error("Could not save your BOM. Please retry. " + str(exc))

                # =================================================
                # VALIDATION FOR ALL TABLES
                # =================================================

                st.divider()

                invalid_rows = []

                all_tables = (
                    [form["items"]]
                    + form["extra_tables"]
                )

                for table_index, table_items in enumerate(
                    all_tables
                ):

                    for row_index, item in enumerate(
                        table_items
                    ):

                        if calculate_bom_quantity(item) == "Invalid":

                            invalid_rows.append(
                                f"Table {table_index + 1}, "
                                f"Row {row_index + 1}"
                            )

                if invalid_rows:

                    st.warning(
                        "Check Actual and Over values in: "
                        + "; ".join(invalid_rows)
                    )

                # =================================================
                # MULTI-PAGE PDF DOWNLOAD
                # =================================================

                with st.container(key="rng_bom_export"):
                    st.markdown('<div class="rng-bom-section">Ready to send</div><div class="rng-bom-help">Download or share all tables together. Each table keeps its original A4 page.</div>', unsafe_allow_html=True)
                    try:

                        pdf_bytes = create_bom_pdf(
                            form,
                            bom_template
                        )

                        safe_name = re.sub(
                            r"[^\w.-]+",
                            "_",
                            bom_name
                        )

                        page_count = len(all_tables)

                        bom_pdf_filename = f"BOM_{safe_name}.pdf"
                        download_col, share_col = st.columns(2)
                        with download_col:
                            st.download_button(
                                label=(
                                    f"Download Complete BOM "
                                    f"({page_count} A4 Pages)"
                                ),
                                data=pdf_bytes,
                                file_name=bom_pdf_filename,
                                mime="application/pdf",
                                key=f"bom_pdf_{bom_id}",
                                use_container_width=True,
                                disabled=bool(invalid_rows)
                            )
                        with share_col:
                            rng_pdf_share_button(
                                pdf_bytes, bom_pdf_filename,
                                disabled=bool(invalid_rows)
                            )

                    except Exception as error:

                        st.error(
                            "BOM PDF generation error: "
                            + str(error)
                        )

    # =========================================================
    # BOM MAIN PAGE
    # =========================================================

    else:

        st.session_state.selected_bom = None

        bom_total = len(st.session_state.bom_cards)
        st.markdown(
            '<div class="rng-bom-hero"><div class="rng-bom-eyebrow">EVERY MATERIAL. EVERY DETAIL.</div>'
            '<h1>BOM</h1><p>Plan the materials. Keep the quantities clear. Build with confidence.</p>'
            f'<span class="rng-bom-total">{bom_total} {"BOM" if bom_total == 1 else "BOMs"} in your workspace</span></div>',
            unsafe_allow_html=True
        )

        # =====================================================
        # ADD NEW BOM
        # =====================================================

        if st.session_state.adding_bom:
            with st.container(key="rng_bom_add_form"):

                st.markdown(
                    '<div class="section-title">Add New BOM</div>',
                    unsafe_allow_html=True
                )

                new_title = st.text_input(
                    "BOM title",
                    placeholder="Enter BOM title",
                    key="new_bom_title"
                )

                add_col, cancel_col = st.columns(2)

                with add_col:

                    if st.button(
                        "Add BOM",
                        key="bom_add_confirm",
                        use_container_width=True
                    ):

                        if new_title.strip():

                            st.session_state.bom_cards.append(
                                new_title.strip()
                            )

                            st.session_state.bom_card_ids.append(
                                uuid4().hex
                            )

                            st.session_state.adding_bom = False
                            rng_rerun()

                        else:

                            st.warning(
                                "Please enter a BOM title."
                            )

                with cancel_col:

                    if st.button(
                        "Cancel",
                        key="bom_add_cancel",
                        use_container_width=True
                    ):

                        st.session_state.adding_bom = False
                        rng_rerun()

        else:
            with st.container(key="rng_bom_add_action"):

                if st.button(
                    "+ Add BOM",
                    key="bom_add_button"
                ):

                    st.session_state.adding_bom = True
                    rng_rerun()

        st.markdown(
            "<br>",
            unsafe_allow_html=True
        )

        # =====================================================
        # BOM CARDS
        # =====================================================

        if not st.session_state.bom_cards:
            st.markdown('<div class="rng-bom-empty"><strong>A clear list makes a strong start.</strong>Add a Bill of Materials to organize your next build.</div>', unsafe_allow_html=True)
        with st.container(key="rng_bom_board"):
            cols = st.columns(3)

            for i, bom in enumerate(
                st.session_state.bom_cards
            ):

                with cols[i % 3], st.container(key=f"rng_bom_card_{i}"):

                    bom_card_heading(bom, i)

                    safe_title = escape(bom)

                    # =============================================
                    # DELETE CONFIRMATION
                    # =============================================

                    if st.session_state.deleting_bom == i:

                        st.markdown(
                            f'<div class="bom-delete-warning">Delete "{safe_title}"? This cannot be undone.</div>',
                            unsafe_allow_html=True
                        )

                        confirm_col, cancel_col = st.columns(2)

                        with confirm_col:

                            if st.button(
                                "Confirm Delete",
                                key=f"bom_confirm_delete_{i}",
                                use_container_width=True
                            ):

                                deleted_id = (
                                    st.session_state.bom_card_ids.pop(i)
                                )

                                st.session_state.bom_cards.pop(i)

                                st.session_state.bom_forms.pop(
                                    deleted_id,
                                    None
                                )

                                st.session_state.deleting_bom = None
                                st.session_state.editing_bom = None

                                rng_rerun()

                        with cancel_col:

                            if st.button(
                                "Cancel",
                                key=f"bom_cancel_delete_{i}",
                                use_container_width=True
                            ):

                                st.session_state.deleting_bom = None

                                rng_rerun()

                    # =============================================
                    # EDIT BOM TITLE
                    # =============================================

                    elif st.session_state.editing_bom == i:

                        edited_title = st.text_input(
                            "BOM title",
                            value=bom,
                            key=f"bom_edit_input_{i}"
                        )

                        save_col, cancel_col = st.columns(2)

                        with save_col:

                            if st.button(
                                "Save",
                                key=f"bom_save_{i}",
                                use_container_width=True
                            ):

                                if edited_title.strip():

                                    st.session_state.bom_cards[i] = (
                                        edited_title.strip()
                                    )

                                    st.session_state.editing_bom = None

                                    rng_rerun()

                                else:

                                    st.warning(
                                        "Please enter a BOM title."
                                    )

                        with cancel_col:

                            if st.button(
                                "Cancel",
                                key=f"bom_edit_cancel_{i}",
                                use_container_width=True
                            ):

                                st.session_state.editing_bom = None

                                rng_rerun()

                    # =============================================
                    # NORMAL BOM CARD
                    # =============================================

                    else:

                        button_col1, button_col2, button_col3 = st.columns(3)

                        with button_col1, st.container(key=f"rng_bom_open_{i}"):

                            open_button = st.button(
                                "Open",
                                key=f"bom_open_{i}",
                                use_container_width=True
                            )

                        with button_col2:

                            edit_button = st.button(
                                "Edit",
                                key=f"bom_edit_{i}",
                                use_container_width=True
                            )

                        with button_col3:

                            delete_button = st.button(
                                "Delete",
                                key=f"bom_delete_{i}",
                                use_container_width=True
                            )

                        if open_button:

                            st.session_state.selected_bom = i

                            rng_rerun()

                        if edit_button:

                            st.session_state.editing_bom = i
                            st.session_state.deleting_bom = None

                            rng_rerun()

                        if delete_button:

                            st.session_state.deleting_bom = i
                            st.session_state.editing_bom = None

                            rng_rerun()


elif page == "RNG Financial":
    from datetime import datetime
    from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
    from html import escape
    from uuid import uuid4
    from zoneinfo import ZoneInfo

    st.session_state.setdefault("financial_projects", {})
    st.session_state.setdefault("financial_selected", None)
    st.session_state.setdefault("financial_creating", False)
    st.session_state.setdefault("financial_editing", None)
    st.session_state.setdefault("financial_deleting", None)
    projects = st.session_state.financial_projects
    currencies = ["IRR", "IRT", "USD", "EUR", "AED", "GBP"]
    currency_labels = {"IRR": "IRR · Iranian rial", "IRT": "IRT · Iranian toman", "USD": "USD · US dollar", "EUR": "EUR · Euro", "AED": "AED · UAE dirham", "GBP": "GBP · British pound"}

    def financial_amount(value):
        amount = Decimal(str(value))
        if not amount.is_finite() or amount < 0:
            raise ValueError("Amounts must be finite, non-negative numbers.")
        return amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

    def financial_totals(project):
        income = sum((financial_amount(row["amount"]) for row in project["transactions"] if row["type"] == "Income"), Decimal(0))
        expense = sum((financial_amount(row["amount"]) for row in project["transactions"] if row["type"] == "Expense"), Decimal(0))
        return income, expense

    def financial_stats(project):
        income, expense = financial_totals(project)
        stats = [("Budget", financial_amount(project["budget"])), ("Income", income), ("Expenses", expense),
                 ("Cash balance", income - expense), ("Budget remaining", financial_amount(project["budget"]) - expense)]
        st.markdown('<div class="rng-financial-stats">' + ''.join(
            f'<div><span>{label}</span><strong>{value:,.2f}</strong><small>{escape(project["currency"])}</small></div>' for label, value in stats
        ) + '</div>', unsafe_allow_html=True)

    st.markdown('<div id="rng-financial-marker"></div>', unsafe_allow_html=True)
    st.markdown("""<style>
    [data-testid="stMain"]:has(#rng-financial-marker) { background: radial-gradient(ellipse at right top, #dce8fa, transparent 50%), #f2f5fa; }
    .stApp:has(#rng-financial-marker), .stApp:has(#rng-financial-marker) input, .stApp:has(#rng-financial-marker) textarea,
    .stApp:has(#rng-financial-marker) h1, .stApp:has(#rng-financial-marker) h3, .stApp:has(#rng-financial-marker) [data-testid="stMarkdownContainer"] { font-family: Arial, Tahoma, sans-serif !important; font-style: normal !important; }
    .stApp:has(#rng-financial-marker) .block-container { max-width: 1440px; padding-top: 2.5rem; }
    .rng-financial-hero { background: radial-gradient(ellipse at right top, #365f8b, transparent 65%), #142c4b; border: 1px solid #29496e; border-radius: 24px; padding: 30px 34px; margin: 4px 0 14px; }
    .rng-financial-hero small { color: #acc8ec; font-size: 11px; letter-spacing: 2px; font-weight: 700; }
    .rng-financial-hero h1 { color: #f5f9ff !important; font-size: clamp(30px, 3.5vw, 44px) !important; margin: 10px 0; padding: 0; overflow-wrap: anywhere; }
    .rng-financial-hero p { color: #c5d7ed; margin: 0; font-size: 14px; line-height: 1.7; }
    .rng-financial-stats { display: grid; grid-template-columns: repeat(5, 1fr); gap: 12px; margin: 8px 0 24px; }
    .rng-financial-stats > div { background: #fff; border: 1px solid #dce5f1; border-radius: 15px; padding: 18px; min-width: 0; }
    .rng-financial-stats span, .rng-financial-stats small { display: block; color: #607897; font-size: 11px; }
    .rng-financial-stats strong { display: block; color: #203f68; font-size: clamp(16px, 1.7vw, 24px); margin: 8px 0; overflow-wrap: anywhere; }
    [class*="st-key-financial_card_"] { background: #fff; border: 1px solid #dce5f1; border-top: 3px solid #7195c6; border-radius: 18px; padding: 22px; box-shadow: 0 6px 20px #16335408; }
    .rng-financial-icon { display: grid; place-items: center; width: 46px; height: 46px; border-radius: 13px; color: #355f92; background: #eaf1fb; margin-bottom: 16px; }
    .rng-financial-icon svg { width: 26px; height: 26px; }
    .rng-financial-card-title { color: #243b5b; font-size: 20px; font-weight: 700; overflow-wrap: anywhere; margin-bottom: 6px; }
    .rng-financial-card-client { color: #6a809d; font-size: 12px; min-height: 20px; margin-bottom: 20px; overflow-wrap: anywhere; }
    .rng-financial-card-balance { border-top: 1px solid #e7edf5; padding: 16px 0; color: #637c9d; font-size: 12px; }
    .rng-financial-card-balance strong { display: block; color: #284b75; font-size: 21px; margin: 6px 0; overflow-wrap: anywhere; }
    .rng-financial-empty { text-align: center; color: #607897; border: 1px dashed #c3d2e7; border-radius: 18px; background: #ffffffb8; padding: 44px 20px; }
    .rng-financial-empty strong { display: block; color: #304d73; font-size: 22px; margin-bottom: 10px; }
    .stApp:has(#rng-financial-marker) [data-testid="stForm"] { background: #fff; border: 1px solid #d5e2f3; border-radius: 18px; padding: 24px; }
    .stApp:has(#rng-financial-marker) [data-testid="stWidgetLabel"] p { color: #4a6383; }
    .stApp:has(#rng-financial-marker) input, .stApp:has(#rng-financial-marker) textarea { background: #f8faff; color: #243b5b; }
    .stApp:has(#rng-financial-marker) .stButton button { border-radius: 10px; }
    .stApp:has(#rng-financial-marker) button[kind="primary"] { background: #214b7f; border-color: #214b7f; color: #fff; }
    .stApp:has(#rng-financial-marker) button[kind="primary"]:hover { background: #173b68; border-color: #173b68; color: #fff; }
    [class*="st-key-financial_card_"] .stButton button { background: #f2f6fc; color: #355780; border-color: #dce5f1; }
    [class*="st-key-financial_card_"] [data-testid="stHorizontalBlock"] { flex-wrap: nowrap !important; gap: 6px; }
    [class*="st-key-financial_card_"] [data-testid="stHorizontalBlock"] > [data-testid="stColumn"] { min-width: 0 !important; flex: 1 1 0 !important; }
    @media (max-width: 1000px) { .rng-financial-stats { grid-template-columns: repeat(3, 1fr); } }
    @media (max-width: 700px) {
      .rng-financial-hero { padding: 24px 20px; }
      .rng-financial-stats { grid-template-columns: repeat(2, 1fr); gap: 8px; }
      .rng-financial-stats > div { padding: 14px; }
      .st-key-financial_board [data-testid="stHorizontalBlock"]:has([class*="st-key-financial_card_"]) { flex-wrap: wrap; }
      .st-key-financial_board [data-testid="stHorizontalBlock"]:has([class*="st-key-financial_card_"]) > [data-testid="stColumn"] { min-width: 100%; }
      .stApp:has(#rng-financial-marker) [data-testid="stForm"] { padding: 16px; }
    }
    </style>""", unsafe_allow_html=True)

    selected = st.session_state.financial_selected
    if selected not in projects:
        st.session_state.financial_selected = None
        selected = None
    if selected is not None:
        project = projects[selected]
        if st.button("Back to projects", key="financial_back"):
            st.session_state.financial_selected = None
            rng_rerun()
        st.markdown('<div class="rng-financial-hero"><small>RNG FINANCIAL / PROJECT</small>'
                    f'<h1>{escape(project["title"])}</h1><p>Every payment and expense, together in one place.</p></div>', unsafe_allow_html=True)
        if st.session_state.pop("financial_saved", False):
            st.success("Project financial records saved.")
        financial_stats(project)
        with st.form(key=f"financial_detail_{selected}"):
            st.subheader("Project details")
            client_col, currency_col, budget_col = st.columns(3)
            with client_col:
                client = st.text_input("Client / employer", value=project["client"], key=f"financial_client_{selected}")
            with currency_col:
                currency = st.selectbox("Currency", currencies, index=currencies.index(project["currency"]), format_func=currency_labels.get, key=f"financial_currency_{selected}")
            with budget_col:
                budget = st.number_input("Project budget", min_value=0.0, value=float(project["budget"]), step=0.01, format="%.2f", key=f"financial_budget_{selected}")
            st.caption("All amounts use the selected currency. Once transactions are saved, the currency is fixed; use a separate project for another currency.")
            st.subheader("Income & expenses")
            st.caption("Add a row for each transaction. Choose Income or Expense and enter a positive amount. Select rows to remove them. Save changes to update totals.")
            frame = pd.DataFrame(project["transactions"], columns=["date", "type", "category", "amount", "description"])
            frame["amount"] = pd.to_numeric(frame["amount"]).astype(float)
            ledger = st.data_editor(frame, key=f"financial_ledger_{selected}", num_rows="dynamic", hide_index=True, width="stretch", height=420,
                column_config={
                    "date": st.column_config.TextColumn("Date", width=120, default=datetime.now(ZoneInfo("Asia/Tehran")).strftime("%Y-%m-%d"), help="Enter your preferred calendar date, e.g. YYYY-MM-DD."),
                    "type": st.column_config.SelectboxColumn("Type", options=["Income", "Expense"], default="Expense", required=True, width=110),
                    "category": st.column_config.TextColumn("Category", width=160, help="For example: materials, wages, invoice payment."),
                    "amount": st.column_config.NumberColumn("Amount", min_value=0.0, step=0.01, format="%.2f", width=160, required=True),
                    "description": st.column_config.TextColumn("Description / reference", width=350),
                })
            notes = st.text_area("Project notes", value=project["notes"], key=f"financial_notes_{selected}")
            submitted = st.form_submit_button("Save financial records", type="primary", use_container_width=True)
        if submitted:
            try:
                clean_rows = []
                for row_number, row in enumerate(ledger.to_dict("records"), start=1):
                    def cell(name):
                        value = row.get(name)
                        return "" if value is None or pd.isna(value) else str(value).strip()
                    # Ignore unused new rows, but do not discard partially entered transactions.
                    if not any(cell(name) for name in ["amount", "category", "description"]):
                        continue
                    if not cell("amount") or financial_amount(cell("amount")) <= 0:
                        raise ValueError(f"Row {row_number}: enter an amount greater than zero.")
                    if not cell("date") or cell("type") not in ["Income", "Expense"]:
                        raise ValueError(f"Row {row_number}: enter a date and choose Income or Expense.")
                    clean_rows.append({"date": cell("date"), "type": cell("type"), "category": cell("category"),
                                       "amount": str(financial_amount(cell("amount"))), "description": cell("description")})
                if currency != project["currency"] and project["transactions"]:
                    raise ValueError("This project already has transactions. Keep its currency, or create a separate project for another currency.")
                projects[selected] = {**project, "client": client.strip(), "currency": currency,
                                      "budget": str(financial_amount(budget)), "notes": notes.strip(), "transactions": clean_rows}
                st.session_state.financial_saved = True
                rng_rerun()
            except (ValueError, InvalidOperation) as exc:
                st.error(str(exc) or "Please check the transaction amounts.")
    else:
        st.markdown('<div class="rng-financial-hero"><small>PROJECT FINANCES</small><h1>RNG Financial</h1>'
                    '<p>Your projects. A clearer picture of every income and expense.</p></div>', unsafe_allow_html=True)
        st.caption(f"{len(projects)} project{'s' if len(projects) != 1 else ''} · Separate records and currency for each project")
        if st.button("+ New project", key="financial_add", type="primary"):
            st.session_state.financial_creating = True
        if st.session_state.financial_creating:
            with st.form("financial_create_form", clear_on_submit=False):
                st.subheader("Create a project")
                title = st.text_input("Project name", key="financial_new_title")
                client = st.text_input("Client / employer", key="financial_new_client")
                currency = st.selectbox("Currency", currencies, format_func=currency_labels.get, key="financial_new_currency")
                budget = st.number_input("Project budget", min_value=0.0, step=0.01, format="%.2f", key="financial_new_budget")
                save_col, cancel_col = st.columns(2)
                with save_col:
                    create = st.form_submit_button("Create project", type="primary", use_container_width=True)
                with cancel_col:
                    cancel = st.form_submit_button("Cancel", use_container_width=True)
            if cancel:
                st.session_state.financial_creating = False
                rng_rerun()
            if create:
                if not title.strip():
                    st.error("Enter a project name.")
                else:
                    project_id = uuid4().hex
                    projects[project_id] = {"title": title.strip(), "client": client.strip(), "currency": currency, "budget": str(financial_amount(budget)), "notes": "", "transactions": []}
                    st.session_state.financial_creating = False
                    for key in ["financial_new_title", "financial_new_client", "financial_new_currency", "financial_new_budget"]:
                        del st.session_state[key]
                    rng_rerun()
        if not projects:
            st.markdown('<div class="rng-financial-empty"><strong>Your first project starts here</strong>Create a project to keep its budget, payments and expenses together.</div>', unsafe_allow_html=True)
        with st.container(key="financial_board"):
            project_ids = list(projects)
            for offset in range(0, len(project_ids), 3):
                for column, project_id in zip(st.columns(3), project_ids[offset:offset + 3]):
                    project = projects[project_id]
                    income, expense = financial_totals(project)
                    with column, st.container(key=f"financial_card_{project_id}"):
                        st.markdown('<div class="rng-financial-icon"><svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.5"><path d="M3 21h18M5 18V10m5 8V10m4 8V10m5 8V10M3 7l9-4 9 4H3Z"/></svg></div>'
                                    f'<div class="rng-financial-card-title">{escape(project["title"])}</div>'
                                    f'<div class="rng-financial-card-client">{escape(project["client"] or "Client not set")}</div>'
                                    f'<div class="rng-financial-card-balance">Cash balance<strong>{income - expense:,.2f} {escape(project["currency"])}</strong>'
                                    f'{len(project["transactions"])} transactions · Budget {financial_amount(project["budget"]):,.2f}</div>', unsafe_allow_html=True)
                        if st.session_state.financial_editing == project_id:
                            with st.form(f"financial_rename_form_{project_id}"):
                                new_title = st.text_input("Project name", value=project["title"], key=f"financial_rename_title_{project_id}")
                                save_col, cancel_col = st.columns(2)
                                with save_col:
                                    save = st.form_submit_button("Save name", use_container_width=True)
                                with cancel_col:
                                    cancel = st.form_submit_button("Cancel", use_container_width=True)
                            if cancel:
                                st.session_state.financial_editing = None
                                rng_rerun()
                            if save:
                                if new_title.strip():
                                    project["title"] = new_title.strip()
                                    st.session_state.financial_editing = None
                                    rng_rerun()
                                else:
                                    st.error("Enter a project name.")
                        elif st.session_state.financial_deleting == project_id:
                            rng_delete_warning("Delete this project and all its financial records? This cannot be undone.")
                            yes_col, no_col = st.columns(2)
                            with yes_col:
                                if st.button("Yes, delete", key=f"financial_confirm_delete_{project_id}", type="primary", use_container_width=True):
                                    del projects[project_id]
                                    st.session_state.financial_deleting = None
                                    rng_rerun()
                            with no_col:
                                if st.button("Cancel", key=f"financial_cancel_delete_{project_id}", use_container_width=True):
                                    st.session_state.financial_deleting = None
                                    rng_rerun()
                        else:
                            open_col, edit_col, delete_col = st.columns(3)
                            with open_col:
                                if st.button("Open", key=f"financial_open_{project_id}", use_container_width=True):
                                    st.session_state.financial_selected = project_id
                                    rng_rerun()
                            with edit_col:
                                if st.button("Edit", key=f"financial_edit_{project_id}", use_container_width=True):
                                    st.session_state.financial_editing = project_id
                                    st.session_state.financial_deleting = None
                                    st.session_state.pop(f"financial_rename_title_{project_id}", None)
                                    rng_rerun()
                            with delete_col:
                                if st.button("Delete", key=f"financial_delete_{project_id}", use_container_width=True):
                                    st.session_state.financial_deleting = project_id
                                    st.session_state.financial_editing = None
                                    rng_rerun()


elif page == "Employer Report":

    st.title("Employer Report")
    st.write("This tab will be developed later.")

# Save text fields, checkboxes and other edits that reran the page normally.
try:
    rng_save()
except Exception as exc:
    st.error("Unable to save changes to Supabase: " + str(exc))
