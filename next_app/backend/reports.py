"""Original report layouts extracted from app.py; no Streamlit runtime."""
import io, re, os, zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from pathlib import Path as RNGPath
import arabic_reshaper
from bidi.algorithm import get_display
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
ROOT = Path(__file__).resolve().parents[2]
RNG_FONT_FILE = ROOT / "B-NAZANIN.TTF"
WAREHOUSE_TEMPLATE = ROOT / 'فرم درخواست کالا.xlsx'
MISSION_TEMPLATE_PATH = ROOT / 'فرم گزارش ماموریت.xlsx'
MISSION_HEADINGS = {'title': 'فرم گزارش ماموریت', 'schedule': 'زمان بندی ماموریت', 'transport': 'اطلاعات حمل و نقل', 'description': 'شرح اجرای ماموریت و هدف آن', 'manager': 'این قسمت توسط مدیر مربوطه تکمیل گردد.'}
BOM_TEMPLATE_PATH = ROOT / 'BOM Raw.xlsx'

def rng_pdf_font_name():
    """Register B Nazanin for PDF Persian text; do not need Windows fonts."""
    if not RNG_FONT_FILE.is_file():
        raise FileNotFoundError('B-NAZANIN.TTF not found. Upload it beside app.py (if licensed).')
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    name = 'RNGBNazanin'
    if name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(name, str(RNG_FONT_FILE)))
    return name

def rng_pdf_runs(visual_text):
    """B Nazanin lacks Latin letters; use Helvetica for Latin runs."""
    from reportlab.pdfbase import pdfmetrics
    persian_font = rng_pdf_font_name()
    available = pdfmetrics.getFont(persian_font).face.charWidths
    runs = []
    for ch in str(visual_text):
        code = ord(ch)
        persian = 1536 <= code <= 2303 or 64336 <= code <= 65279
        font = persian_font if code in available and (persian or not ch.isascii()) else 'Helvetica'
        if (ch.isdigit() or ch in ' .,:;/-+()[]%') and code in available:
            font = persian_font
        if runs and runs[-1][0] == font:
            runs[-1] = (font, runs[-1][1] + ch)
        else:
            runs.append((font, ch))
    return runs

def rng_pdf_width(visual_text, size):
    from reportlab.pdfbase import pdfmetrics
    return sum((pdfmetrics.stringWidth(part, name, size) for name, part in rng_pdf_runs(visual_text)))

def rng_pdf_draw(pdf, visual_text, anchor_x, baseline, size, align='left'):
    """Draw shaped visual-order Persian/English without missing Latin letters."""
    from reportlab.pdfbase import pdfmetrics
    runs = rng_pdf_runs(visual_text)
    total_width = sum((pdfmetrics.stringWidth(s, name, size) for name, s in runs))
    x = anchor_x - (total_width / 2 if align == 'center' else total_width if align == 'right' else 0)
    for name, part in runs:
        pdf.setFont(name, size)
        pdf.drawString(x, baseline, part)
        x += pdfmetrics.stringWidth(part, name, size)

def create_meeting_pdf(meeting_data):
    buffer = io.BytesIO()
    page_width, page_height = landscape(A4)
    pdf = canvas.Canvas(buffer, pagesize=(page_width, page_height))
    font_name = rng_pdf_font_name()
    margin_x = 7 * mm
    margin_y = 7 * mm
    table_width = page_width - 2 * margin_x
    table_height = page_height - 2 * margin_y
    number_of_rows = 12
    row_height = table_height / number_of_rows
    column_ratios = [1.0, 2.4, 1.0, 1.0, 1.0, 2.4]
    ratio_total = sum(column_ratios)
    column_widths = [table_width * ratio / ratio_total for ratio in column_ratios]
    column_order = [5, 4, 3, 2, 1, 0]
    pdf_headers = ['Status', 'Minute', 'Status', 'Person in charge', 'Title', 'Agenda']

    def format_text(value):
        value = str(value if value is not None else '')
        return get_display(arabic_reshaper.reshape(value))

    def text_width(value, font_size):
        return rng_pdf_width(format_text(value), font_size)

    def wrap_text(value, max_width, font_size):
        value = str(value if value is not None else '').strip()
        words = re.split('\\s+', value)
        lines = []
        current_line = ''
        for word in words:
            if not word:
                continue
            candidate = current_line + ' ' + word if current_line else word
            if text_width(candidate, font_size) <= max_width:
                current_line = candidate
            else:
                if current_line:
                    lines.append(current_line)
                current_line = ''
                for character in word:
                    candidate = current_line + character
                    if current_line and text_width(candidate, font_size) > max_width:
                        lines.append(current_line)
                        current_line = character
                    else:
                        current_line = candidate
        if current_line:
            lines.append(current_line)
        return lines or ['']

    def draw_cell_text(value, x, y, width, height, is_header=False):
        padding = 4
        available_width = max(width - 2 * padding, 1)
        font_size = 9.0
        while font_size > 6.0:
            lines = wrap_text(value, available_width, font_size)
            line_height = font_size * 1.3
            if len(lines) * line_height <= height - 2 * padding:
                break
            font_size -= 0.5
        lines = wrap_text(value, available_width, font_size)
        line_height = font_size * 1.3
        max_lines = max(1, int((height - 2 * padding) / line_height))
        if len(lines) > max_lines:
            lines = lines[:max_lines]
            last_line = lines[-1]
            while last_line and text_width(last_line + '...', font_size) > available_width:
                last_line = last_line[:-1]
            lines[-1] = last_line + '...'
        pdf.setFont(font_name, font_size)
        if is_header:
            pdf.setFillColorRGB(1, 1, 1)
        else:
            pdf.setFillColorRGB(0.12, 0.17, 0.23)
        text_height = len(lines) * line_height
        baseline = y + (height + text_height) / 2 - line_height + 2
        for line in lines:
            rendered = format_text(line)
            rng_pdf_draw(pdf, rendered, x + width / 2, baseline, font_size, 'center')
            baseline -= line_height
    current_y = page_height - margin_y - row_height
    current_x = margin_x
    for column_index, header in enumerate(pdf_headers):
        cell_width = column_widths[column_index]
        pdf.setFillColorRGB(23 / 255, 35 / 255, 51 / 255)
        pdf.rect(current_x, current_y, cell_width, row_height, fill=1, stroke=0)
        draw_cell_text(header, current_x, current_y, cell_width, row_height, is_header=True)
        current_x += cell_width
    for row_index in range(11):
        current_y = page_height - margin_y - row_height * (row_index + 2)
        current_x = margin_x
        for column_index, data_col in enumerate(column_order):
            cell_width = column_widths[column_index]
            value = meeting_data[row_index][data_col]
            if data_col == 3 and (not value):
                value = 'Raised'
            if data_col == 5 and (not value):
                value = 'Done'
            pdf.setFillColorRGB(1, 1, 1)
            pdf.rect(current_x, current_y, cell_width, row_height, fill=1, stroke=0)
            draw_cell_text(value, current_x, current_y, cell_width, row_height)
            current_x += cell_width
    pdf.setStrokeColorRGB(0.72, 0.76, 0.81)
    pdf.setLineWidth(0.45)
    for row_number in range(number_of_rows + 1):
        y = page_height - margin_y - row_number * row_height
        pdf.line(margin_x, y, margin_x + table_width, y)
    x = margin_x
    pdf.line(x, margin_y, x, margin_y + table_height)
    for cell_width in column_widths:
        x += cell_width
        pdf.line(x, margin_y, x, margin_y + table_height)
    pdf.showPage()
    pdf.save()
    buffer.seek(0)
    return buffer.getvalue()

def load_warehouse_template():
    if not WAREHOUSE_TEMPLATE.is_file():
        raise FileNotFoundError('Warehouse Excel template not found: ' + str(WAREHOUSE_TEMPLATE))
    ns = {'x': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(WAREHOUSE_TEMPLATE) as z:
        strings_xml = ET.fromstring(z.read('xl/sharedStrings.xml'))
        strings = []
        for item in strings_xml.findall('x:si', ns):
            value = ''.join((element.text or '' for element in item.findall('.//x:t', ns)))
            strings.append(value)
        logo = z.read('xl/media/image1.png')
    if len(strings) < 10:
        raise ValueError('The original Excel form labels are missing.')
    date_number_labels = strings[5].splitlines()
    defaults = strings[9].splitlines()
    requester_lines = strings[8].splitlines()
    return {'headers': strings[0:5], 'date_label': date_number_labels[0], 'number_label': date_number_labels[1], 'title': strings[6], 'approver_label': strings[7].split(':')[0] + ':', 'requester_label': requester_lines[0], 'signature_label': 'امضا', 'default_requester': requester_lines[1].strip(), 'default_date': defaults[0].strip(), 'default_number': defaults[1].strip(), 'logo': logo}

def create_empty_warehouse_form():
    return {'date': '', 'number': '', 'items': [['', '', '', '', ''] for _ in range(10)], 'requester': '', 'approver': '', 'requester_signature': '', 'approver_signature': ''}

def create_warehouse_pdf(form, template):
    import arabic_reshaper
    from bidi.algorithm import get_display
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A5, landscape
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.utils import ImageReader
    buffer = io.BytesIO()
    page_width, page_height = landscape(A5)
    pdf = canvas.Canvas(buffer, pagesize=(page_width, page_height), pageCompression=1)
    pdf.setTitle(template['title'])
    font_name = rng_pdf_font_name()

    def shape_text(value):
        value = str(value if value is not None else '')
        reshaped = arabic_reshaper.reshape(value)
        return get_display(reshaped)

    def text_width(value, size):
        return rng_pdf_width(shape_text(value), size)

    def draw_cell_text(value, x, y, width, height, font_size=9, align='right'):
        text = str(value if value is not None else '').strip()
        if not text:
            return
        padding = 3
        available_width = max(width - 2 * padding, 1)
        available_height = max(height - 2 * padding, 1)

        def make_lines(size):
            lines = []
            for paragraph in text.splitlines():
                if not paragraph.strip():
                    lines.append('')
                    continue
                words = paragraph.split()
                current = ''
                for word in words:
                    candidate = current + ' ' + word if current else word
                    if text_width(candidate, size) <= available_width:
                        current = candidate
                    else:
                        if current:
                            lines.append(current)
                        current = ''
                        for character in word:
                            candidate = current + character
                            if current and text_width(candidate, size) > available_width:
                                lines.append(current)
                                current = character
                            else:
                                current = candidate
                if current:
                    lines.append(current)
            return lines or ['']
        size = float(font_size)
        while size > 5:
            lines = make_lines(size)
            line_height = size * 1.3
            if len(lines) * line_height <= available_height:
                break
            size -= 0.5
        lines = make_lines(size)
        line_height = size * 1.3
        max_lines = max(1, int(available_height / line_height))
        if len(lines) > max_lines:
            lines = lines[:max_lines]
            final_line = lines[-1]
            while final_line and text_width(final_line + '...', size) > available_width:
                final_line = final_line[:-1]
            lines[-1] = final_line + '...'
        pdf.saveState()
        clip = pdf.beginPath()
        clip.rect(x + 1, y + 1, max(width - 2, 0), max(height - 2, 0))
        pdf.clipPath(clip, stroke=0, fill=0)
        pdf.setFillColorRGB(0, 0, 0)
        pdf.setFont(font_name, size)
        total_height = len(lines) * line_height
        baseline = y + (height + total_height) / 2 - line_height + 2
        for line in lines:
            rendered = shape_text(line)
            if align == 'center':
                rng_pdf_draw(pdf, rendered, x + width / 2, baseline, size, 'center')
            elif align == 'left':
                rng_pdf_draw(pdf, rendered, x + padding, baseline, size, 'left')
            else:
                rng_pdf_draw(pdf, rendered, x + width - padding, baseline, size, 'right')
            baseline -= line_height
        pdf.restoreState()
    margin_x = 6 * mm
    margin_y = 6 * mm
    left = margin_x
    right = page_width - margin_x
    top = page_height - margin_y
    bottom = margin_y
    table_width = right - left
    table_height = top - bottom
    excel_row_heights = [60, 42] + [21.75] * 10 + [66]
    height_scale = table_height / sum(excel_row_heights)
    row_heights = [height * height_scale for height in excel_row_heights]
    excel_col_widths = [5.85546875, 17.42578125, 27.0, 10.140625, 15.5703125, 6.42578125, 16.7109375]
    width_scale = table_width / sum(excel_col_widths)
    col_widths = [width * width_scale for width in excel_col_widths]
    x_positions = [right]
    for width in col_widths:
        x_positions.append(x_positions[-1] - width)
    y_positions = [top]
    for height in row_heights:
        y_positions.append(y_positions[-1] - height)
    pdf.setStrokeColorRGB(0, 0, 0)
    pdf.setLineWidth(0.65)

    def draw_cell(first_col, last_col, row_index, value='', font_size=9, align='center'):
        x = x_positions[last_col + 1]
        width = x_positions[first_col] - x_positions[last_col + 1]
        y = y_positions[row_index + 1]
        height = row_heights[row_index]
        pdf.rect(x, y, width, height, fill=0, stroke=1)
        draw_cell_text(value, x, y, width, height, font_size=font_size, align=align)
    draw_cell(0, 1, 0)
    draw_cell(2, 4, 0, template['title'], font_size=16, align='center')
    draw_cell(5, 5, 0)
    draw_cell(6, 6, 0)
    logo_reader = ImageReader(io.BytesIO(template['logo']))
    logo_x = x_positions[2]
    logo_y = y_positions[1]
    logo_width = x_positions[0] - x_positions[2]
    logo_height = row_heights[0]
    image_width, image_height = logo_reader.getSize()
    logo_scale = min(logo_width * 0.94 / image_width, logo_height * 0.94 / image_height)
    rendered_width = image_width * logo_scale
    rendered_height = image_height * logo_scale
    pdf.drawImage(logo_reader, logo_x + (logo_width - rendered_width) / 2, logo_y + (logo_height - rendered_height) / 2, width=rendered_width, height=rendered_height, preserveAspectRatio=True, mask='auto')
    f_x = x_positions[6]
    f_width = col_widths[5]
    g_x = x_positions[7]
    g_width = col_widths[6]
    half_height = row_heights[0] / 2
    draw_cell_text(template['date_label'], f_x, logo_y + half_height, f_width, half_height, font_size=9, align='right')
    draw_cell_text(template['number_label'], f_x, logo_y, f_width, half_height, font_size=9, align='right')
    draw_cell_text(form['date'], g_x, logo_y + half_height, g_width, half_height, font_size=9, align='right')
    draw_cell_text(form['number'], g_x, logo_y, g_width, half_height, font_size=9, align='right')
    draw_cell(0, 0, 1, template['headers'][0], font_size=9)
    draw_cell(1, 2, 1, template['headers'][1], font_size=10)
    draw_cell(3, 3, 1, template['headers'][2], font_size=9)
    draw_cell(4, 4, 1, template['headers'][3], font_size=9)
    draw_cell(5, 6, 1, template['headers'][4], font_size=10)
    merged_columns = [(0, 0), (1, 2), (3, 3), (4, 4), (5, 6)]
    for item_index in range(10):
        excel_row_index = item_index + 2
        for field_index, col_range in enumerate(merged_columns):
            draw_cell(col_range[0], col_range[1], excel_row_index, form['items'][item_index][field_index], font_size=8.5, align='right')
    footer_row = 12
    draw_cell(0, 2, footer_row)
    draw_cell(3, 6, footer_row)
    footer_y = y_positions[13]
    footer_height = row_heights[12]
    requester_x = x_positions[3]
    requester_width = x_positions[0] - x_positions[3]
    approver_x = x_positions[7]
    approver_width = x_positions[3] - x_positions[7]
    draw_cell_text(template['requester_label'], requester_x, footer_y + footer_height * 0.7, requester_width, footer_height * 0.25, font_size=10, align='right')
    draw_cell_text(form['requester'], requester_x, footer_y + footer_height * 0.43, requester_width, footer_height * 0.25, font_size=10, align='right')
    draw_cell_text(template['signature_label'] + ': ' + form['requester_signature'], requester_x, footer_y + footer_height * 0.05, requester_width, footer_height * 0.25, font_size=9, align='right')
    draw_cell_text(template['approver_label'], approver_x, footer_y + footer_height * 0.7, approver_width, footer_height * 0.25, font_size=10, align='right')
    draw_cell_text(form['approver'], approver_x, footer_y + footer_height * 0.43, approver_width, footer_height * 0.25, font_size=10, align='right')
    draw_cell_text(template['signature_label'] + ': ' + form['approver_signature'], approver_x, footer_y + footer_height * 0.05, approver_width, footer_height * 0.25, font_size=9, align='right')
    pdf.showPage()
    pdf.save()
    buffer.seek(0)
    return buffer.getvalue()

def load_mission_template():
    from openpyxl import load_workbook
    if not MISSION_TEMPLATE_PATH.is_file():
        raise FileNotFoundError('Excel template not found: ' + str(MISSION_TEMPLATE_PATH))
    wb = load_workbook(MISSION_TEMPLATE_PATH, read_only=False, data_only=True)
    if 'RAW' not in wb.sheetnames:
        wb.close()
        raise ValueError('RAW worksheet not found.')
    ws = wb['RAW']
    cells = {}
    for row in range(1, 14):
        for col in range(1, 9):
            cell = ws.cell(row, col)
            if cell.value is not None:
                cells[cell.coordinate] = str(cell.value)
    column_widths = [ws.column_dimensions[letter].width or 10 for letter in 'ABCDEFGH']
    row_heights = [ws.row_dimensions[row].height or 25 for row in range(1, 14)]
    wb.close()
    logo = None
    with zipfile.ZipFile(MISSION_TEMPLATE_PATH) as archive:
        images = [name for name in archive.namelist() if name.startswith('xl/media/') and name.lower().endswith(('.png', '.jpg', '.jpeg'))]
        if images:
            logo = archive.read(images[0])
    return {'cells': cells, 'column_widths': column_widths, 'row_heights': row_heights, 'logo': logo}

def create_empty_mission_form():
    return {field: '' for field in FIELD_LABELS}

def create_mission_pdf(form, template):
    import arabic_reshaper
    from bidi.algorithm import get_display
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.utils import ImageReader
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4, pageCompression=1)
    page_width, page_height = A4
    pdf.setTitle(template['cells'].get('C1', MISSION_HEADINGS['title']))
    font_name = rng_pdf_font_name()

    def shape_text(value):
        value = str(value if value is not None else '')
        return get_display(arabic_reshaper.reshape(value))
    margin_x = 7 * mm
    margin_y = 7 * mm
    left = margin_x
    right = page_width - margin_x
    top = page_height - margin_y
    bottom = margin_y
    usable_width = right - left
    usable_height = top - bottom
    col_source = template['column_widths']
    row_source = template['row_heights']
    col_total = sum(col_source)
    row_total = sum(row_source)
    col_widths = [usable_width * width / col_total for width in col_source]
    row_heights = [usable_height * height / row_total for height in row_source]
    x_edges = [right]
    for width in col_widths:
        x_edges.append(x_edges[-1] - width)
    y_edges = [top]
    for height in row_heights:
        y_edges.append(y_edges[-1] - height)

    def text_width(value, size):
        return rng_pdf_width(shape_text(value), size)

    def wrap_text(value, max_width, size):
        lines = []
        for paragraph in str(value).splitlines():
            if not paragraph.strip():
                lines.append('')
                continue
            current = ''
            for word in paragraph.split():
                candidate = current + ' ' + word if current else word
                if text_width(candidate, size) <= max_width:
                    current = candidate
                else:
                    if current:
                        lines.append(current)
                    current = ''
                    for character in word:
                        candidate = current + character
                        if current and text_width(candidate, size) > max_width:
                            lines.append(current)
                            current = character
                        else:
                            current = candidate
            if current:
                lines.append(current)
        return lines or ['']

    def draw_text(value, x, y, width, height, size=10, align='right'):
        text = str(value if value is not None else '').strip()
        if not text:
            return
        padding = 3
        max_width = max(width - 2 * padding, 1)
        max_height = max(height - 2 * padding, 1)
        font_size = float(size)
        while font_size > 5:
            lines = wrap_text(text, max_width, font_size)
            line_height = font_size * 1.3
            if len(lines) * line_height <= max_height:
                break
            font_size -= 0.5
        lines = wrap_text(text, max_width, font_size)
        line_height = font_size * 1.3
        max_lines = max(1, int(max_height / line_height))
        if len(lines) > max_lines:
            lines = lines[:max_lines]
            final_line = lines[-1]
            while final_line and text_width(final_line + '...', font_size) > max_width:
                final_line = final_line[:-1]
            lines[-1] = final_line + '...'
        pdf.saveState()
        clipping = pdf.beginPath()
        clipping.rect(x + 1, y + 1, max(width - 2, 0), max(height - 2, 0))
        pdf.clipPath(clipping, stroke=0, fill=0)
        pdf.setFont(font_name, font_size)
        pdf.setFillColorRGB(0.08, 0.12, 0.18)
        total_height = len(lines) * line_height
        baseline = y + (height + total_height) / 2 - line_height + 2
        for line in lines:
            rendered = shape_text(line)
            if align == 'center':
                rng_pdf_draw(pdf, rendered, x + width / 2, baseline, font_size, 'center')
            else:
                rng_pdf_draw(pdf, rendered, x + width - padding, baseline, font_size, 'right')
            baseline -= line_height
        pdf.restoreState()

    def draw_cell(first_col, last_col, row, text='', font_size=10, align='right', shaded=False, signature_path=None):
        x = x_edges[last_col + 1]
        width = x_edges[first_col] - x_edges[last_col + 1]
        y = y_edges[row]
        height = row_heights[row - 1]
        if shaded:
            pdf.setFillColorRGB(0.92, 0.93, 0.95)
            pdf.rect(x, y, width, height, fill=1, stroke=0)
        pdf.setStrokeColorRGB(0.3, 0.3, 0.3)
        pdf.setLineWidth(0.55)
        pdf.rect(x, y, width, height, fill=0, stroke=1)
        if signature_path is not None:
            signature = ImageReader(str(signature_path))
            image_width, image_height = signature.getSize()
            scale = min(54 * mm / image_width, 28 * mm / image_height, (height - 12) / image_height)
            signature_width = image_width * scale
            signature_height = image_height * scale
            text_inset = signature_width + 12
            draw_text(text, x + text_inset, y, width - text_inset, height, size=font_size, align=align)
            pdf.drawImage(signature, x + 5, y + 5, width=signature_width, height=signature_height, preserveAspectRatio=True, mask='auto')
        else:
            draw_text(text, x, y, width, height, size=font_size, align=align)

    def field_text(field):
        return FIELD_LABELS[field] + ': ' + str(form.get(field, ''))
    draw_cell(0, 1, 1)
    draw_cell(2, 5, 1, template['cells'].get('C1', MISSION_HEADINGS['title']), font_size=16, align='center')
    draw_cell(6, 7, 1, field_text('report_date'))
    if template['logo']:
        image = ImageReader(io.BytesIO(template['logo']))
        image_width, image_height = image.getSize()
        logo_x = x_edges[2]
        logo_y = y_edges[1]
        logo_width = x_edges[0] - x_edges[2]
        logo_height = row_heights[0]
        scale = min(logo_width * 0.92 / image_width, logo_height * 0.92 / image_height)
        render_width = image_width * scale
        render_height = image_height * scale
        pdf.drawImage(image, logo_x + (logo_width - render_width) / 2, logo_y + (logo_height - render_height) / 2, width=render_width, height=render_height, preserveAspectRatio=True, mask='auto')
    draw_cell(0, 3, 2, field_text('employee'))
    draw_cell(4, 7, 2, field_text('position'))
    draw_cell(0, 2, 3, field_text('dispatcher'))
    draw_cell(3, 4, 3, field_text('city'))
    draw_cell(5, 7, 3, field_text('company'))
    draw_cell(0, 7, 4, field_text('companions'))
    draw_cell(0, 0, 5, field_text('duration'), font_size=8)
    draw_cell(1, 2, 5, field_text('start_date'), font_size=9)
    draw_cell(3, 3, 5, field_text('departure_time'), font_size=8)
    draw_cell(4, 5, 5, field_text('end_date'), font_size=9)
    draw_cell(6, 7, 5, field_text('return_time'), font_size=9)
    draw_cell(0, 7, 6, field_text('transport') + '       ' + field_text('distance'))
    draw_cell(0, 7, 7, template['cells'].get('A7', MISSION_HEADINGS['description']), font_size=11, shaded=True)
    draw_cell(0, 7, 8, field_text('equipment'))
    draw_cell(0, 7, 9, FIELD_LABELS['description'] + ':\n' + str(form.get('description', '')), font_size=10, signature_path=ROOT / 'assets' / 'Mission signature.png')
    draw_cell(0, 7, 10, field_text('objective'))
    draw_cell(0, 7, 11, field_text('result'))
    draw_cell(0, 7, 12, template['cells'].get('A12', MISSION_HEADINGS['manager']), font_size=10, align='center', shaded=True)
    draw_cell(0, 3, 13, FIELD_LABELS['manager_comment'] + ':\n' + str(form.get('manager_comment', '')) + '\n' + field_text('manager_signature'))
    draw_cell(4, 7, 13, field_text('employee_signature'))
    pdf.showPage()
    pdf.save()
    buffer.seek(0)
    return buffer.getvalue()

def load_bom_template():
    from openpyxl import load_workbook
    if not BOM_TEMPLATE_PATH.is_file():
        raise FileNotFoundError('BOM Raw.xlsx was not found beside app.py')
    workbook = load_workbook(BOM_TEMPLATE_PATH, read_only=False, data_only=False)
    if 'RAW' not in workbook.sheetnames:
        workbook.close()
        raise ValueError('RAW worksheet not found.')
    ws = workbook['RAW']
    labels = {'title': str(ws['C1'].value or ''), 'metadata': str(ws['J1'].value or ''), 'headers': [str(ws[f'{col}2'].value or '') for col in 'ABDEFGHI'], 'preparer': str(ws['A18'].value or ''), 'approver': str(ws['D18'].value or '')}
    widths = [ws.column_dimensions[c].width or 10 for c in 'ABCDEFGHIJK']
    heights = [ws.row_dimensions[r].height or 25 for r in range(1, 19)]
    workbook.close()
    return {'labels': labels, 'widths': widths, 'heights': heights}

def create_empty_bom_table():
    return [{'description': '', 'size': '', 'actual': '', 'over': '', 'unit': '', 'notes': ''} for _ in range(15)]

def create_empty_bom():
    return {'date': '', 'number': '', 'code': '', 'preparer': '', 'approver': '', 'items': create_empty_bom_table(), 'extra_tables': []}

def prepare_bom_form(form):
    if 'items' not in form:
        form['items'] = create_empty_bom_table()
    if 'extra_tables' not in form:
        form['extra_tables'] = []
    for field in ['date', 'number', 'code', 'preparer', 'approver']:
        if field not in form:
            form[field] = ''
    all_tables = [form['items']] + form['extra_tables']
    for table in all_tables:
        while len(table) < 15:
            table.append({'description': '', 'size': '', 'actual': '', 'over': '', 'unit': '', 'notes': ''})
        for item in table:
            for field in ['description', 'size', 'actual', 'over', 'unit', 'notes']:
                if field not in item:
                    item[field] = ''

def parse_bom_number(value):
    text = str(value).strip()
    if not text:
        return 0.0
    digit_map = str.maketrans('۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩', '01234567890123456789')
    text = text.translate(digit_map)
    text = text.replace('٬', '')
    text = text.replace(',', '')
    text = text.replace('٫', '.')
    return float(text)

def calculate_bom_quantity(item):
    actual_text = str(item.get('actual', '')).strip()
    over_text = str(item.get('over', '')).strip()
    if not actual_text and (not over_text):
        return ''
    try:
        actual = parse_bom_number(actual_text)
        over = parse_bom_number(over_text)
        total = actual + over
        if not float('-inf') < total < float('inf'):
            return 'Invalid'
        if total.is_integer():
            return str(int(total))
        return f'{total:.6f}'.rstrip('0').rstrip('.')
    except (ValueError, TypeError):
        return 'Invalid'

def create_bom_pdf(form, template):
    prepare_bom_form(form)
    all_tables = [form["items"]] + form["extra_tables"]
    import arabic_reshaper
    from bidi.algorithm import get_display
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    buffer = io.BytesIO()
    page_width, page_height = A4
    pdf = canvas.Canvas(buffer, pagesize=A4, pageCompression=1)
    labels = template['labels']
    pdf.setTitle(labels['title'])
    font_name = rng_pdf_font_name()

    def shape_text(value):
        return get_display(arabic_reshaper.reshape(str(value if value is not None else '')))
    margin_x = 6 * mm
    margin_y = 6 * mm
    left = margin_x
    right = page_width - margin_x
    top = page_height - margin_y
    bottom = margin_y
    usable_width = right - left
    usable_height = top - bottom
    source_widths = template['widths']
    source_heights = template['heights']
    width_total = sum(source_widths)
    height_total = sum(source_heights)
    col_widths = [usable_width * w / width_total for w in source_widths]
    row_heights = [usable_height * h / height_total for h in source_heights]
    x_edges = [right]
    for width in col_widths:
        x_edges.append(x_edges[-1] - width)
    y_edges = [top]
    for height in row_heights:
        y_edges.append(y_edges[-1] - height)

    def text_width(value, size):
        return rng_pdf_width(shape_text(value), size)

    def wrap_text(value, width, size):
        result = []
        for paragraph in str(value).splitlines():
            if not paragraph.strip():
                result.append('')
                continue
            current = ''
            for word in paragraph.split():
                candidate = current + ' ' + word if current else word
                if text_width(candidate, size) <= width:
                    current = candidate
                else:
                    if current:
                        result.append(current)
                    current = ''
                    for char in word:
                        candidate = current + char
                        if current and text_width(candidate, size) > width:
                            result.append(current)
                            current = char
                        else:
                            current = candidate
            if current:
                result.append(current)
        return result or ['']

    def draw_text(value, x, y, width, height, size=8, align='center'):
        value = str(value if value is not None else '').strip()
        if not value:
            return
        padding = 2.5
        available_width = max(width - 2 * padding, 1)
        available_height = max(height - 2 * padding, 1)
        font_size = float(size)
        while font_size > 4.5:
            lines = wrap_text(value, available_width, font_size)
            line_height = font_size * 1.25
            if len(lines) * line_height <= available_height:
                break
            font_size -= 0.5
        lines = wrap_text(value, available_width, font_size)
        line_height = font_size * 1.25
        maximum_lines = max(1, int(available_height / line_height))
        if len(lines) > maximum_lines:
            lines = lines[:maximum_lines]
            last = lines[-1]
            while last and text_width(last + '...', font_size) > available_width:
                last = last[:-1]
            lines[-1] = last + '...'
        pdf.saveState()
        clipping = pdf.beginPath()
        clipping.rect(x + 1, y + 1, max(width - 2, 0), max(height - 2, 0))
        pdf.clipPath(clipping, stroke=0, fill=0)
        pdf.setFont(font_name, font_size)
        pdf.setFillColorRGB(0, 0, 0)
        total_height = len(lines) * line_height
        baseline = y + (height + total_height) / 2 - line_height + 2
        for line in lines:
            shaped = shape_text(line)
            if align == 'right':
                rng_pdf_draw(pdf, shaped, x + width - padding, baseline, font_size, 'right')
            elif align == 'left':
                rng_pdf_draw(pdf, shaped, x + padding, baseline, font_size, 'left')
            else:
                rng_pdf_draw(pdf, shaped, x + width / 2, baseline, font_size, 'center')
            baseline -= line_height
        pdf.restoreState()

    def draw_cell(first_col, last_col, row, text='', size=8, shaded=False, align='center'):
        x = x_edges[last_col + 1]
        width = x_edges[first_col] - x_edges[last_col + 1]
        y = y_edges[row]
        height = row_heights[row - 1]
        if shaded:
            pdf.setFillColorRGB(0.93, 0.94, 0.96)
            pdf.rect(x, y, width, height, fill=1, stroke=0)
        pdf.setStrokeColorRGB(0.25, 0.25, 0.25)
        pdf.setLineWidth(0.5)
        pdf.rect(x, y, width, height, fill=0, stroke=1)
        draw_text(text, x, y, width, height, size=size, align=align)
    for table_index, table_items in enumerate(all_tables):
        draw_cell(0, 1, 1)
        draw_cell(2, 8, 1, labels['title'], size=17)
        metadata_labels = labels['metadata'].splitlines()
        date_label = metadata_labels[0] if len(metadata_labels) > 0 else 'Date:'
        number_label = metadata_labels[1] if len(metadata_labels) > 1 else 'Number:'
        code_label = metadata_labels[2] if len(metadata_labels) > 2 else 'Code:'
        draw_cell(9, 9, 1, '\n'.join([date_label, number_label, code_label]), size=9, align='right')
        draw_cell(10, 10, 1, '\n'.join([form['date'], form['number'], form['code']]), size=9)
        headers = labels['headers']
        merged_ranges = [(0, 0), (1, 2), (3, 3), (4, 4), (5, 5), (6, 6), (7, 7), (8, 10)]
        for index, (start, end) in enumerate(merged_ranges):
            draw_cell(start, end, 2, headers[index], size=8, shaded=True)
        for item_index in range(15):
            item = table_items[item_index]
            row_number = item_index + 3
            values = [str(item_index + 1), item['description'], item['size'], item['actual'], item['over'], calculate_bom_quantity(item), item['unit'], item['notes']]
            for column_index, (start, end) in enumerate(merged_ranges):
                draw_cell(start, end, row_number, values[column_index], size=8)
        draw_cell(0, 2, 18, labels['preparer'] + '\n' + form['preparer'], size=10, align='right')
        draw_cell(3, 10, 18, labels['approver'] + '\n' + form['approver'], size=10, align='right')
        pdf.showPage()
    pdf.save()
    buffer.seek(0)
    return buffer.getvalue()

FIELD_LABELS = {'report_date': 'تاریخ گزارش', 'employee': 'نام و نام خانوادگی', 'position': 'پست سازمانی', 'dispatcher': 'اعزام کننده', 'city': 'شهر ماموریت', 'company': 'نام شرکت / کارخانه', 'companions': 'همراهان', 'duration': 'مدت ماموریت', 'start_date': 'تاریخ شروع', 'departure_time': 'ساعت رفت', 'end_date': 'تاریخ پایان', 'return_time': 'ساعت برگشت', 'transport': 'استفاده از وسایل نقلیه', 'distance': 'مسافت طی شده (کیلومتر)', 'equipment': 'تجهیزات مورد استفاده', 'description': 'شرح ماموریت', 'objective': 'هدف ماموریت', 'result': 'نتیجه ماموریت', 'manager_comment': 'نظر مدیر / رئیس مربوطه', 'manager_signature': 'امضای مدیر', 'employee_signature': 'امضای مامور'}
